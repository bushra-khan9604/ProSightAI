"""Luna-powered structured column mapping for non-canonical workbooks."""

from __future__ import annotations

import json
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

from ..config import get_settings
from .excel import PROJECT_COLUMNS


class OpenAIColumnMapper:
    """Propose, but never execute, a mapping from arbitrary columns to project fields."""

    def propose(self, path: Path) -> dict:
        """Ask Luna for a strict JSON mapping using headers and small sample values."""
        settings = get_settings()
        if not settings.openai_api_key:
            raise RuntimeError("OPENAI_API_KEY is required for non-canonical Excel mapping")
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        samples = {}
        for sheet in workbook.worksheets:
            # Some valid workbooks omit cached worksheet dimensions, causing
            # openpyxl read-only worksheets to expose max_row=None.
            sample_limit = min(sheet.max_row or 6, 6)
            rows = list(sheet.iter_rows(min_row=1, max_row=sample_limit, values_only=True))
            samples[sheet.title] = rows
        workbook.close()
        schema = {
            "type": "object",
            "properties": {
                "sheet": {"type": "string"},
                "columns": {
                    "type": "object",
                    "properties": {name: {"type": "string"} for name in PROJECT_COLUMNS},
                    "required": PROJECT_COLUMNS,
                    "additionalProperties": False,
                },
                "confidence": {"type": "number"},
                "notes": {"type": "array", "items": {"type": "string"}},
            },
            "required": ["sheet", "columns", "confidence", "notes"],
            "additionalProperties": False,
        }
        body = {
            "model": settings.openai_model,
            "instructions": (
                "Map one workbook sheet to the canonical construction project schema. "
                "Use exact source header strings. Do not invent columns. Return only JSON."
            ),
            "input": json.dumps({"required_columns": PROJECT_COLUMNS, "workbook_samples": samples}, default=str),
            "text": {"format": {"type": "json_schema", "name": "excel_mapping", "strict": True, "schema": schema}},
        }
        request = urllib.request.Request(
            "https://api.openai.com/v1/responses",
            data=json.dumps(body).encode(),
            headers={
                "Authorization": f"Bearer {settings.openai_api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=60) as response:
            result = json.loads(response.read())
        output = "".join(
            part.get("text", "")
            for item in result.get("output", [])
            for part in item.get("content", [])
            if part.get("type") == "output_text"
        )
        mapping = json.loads(output)
        if mapping["confidence"] < 0.6:
            mapping["notes"].append("Low-confidence mapping requires careful review")
        return mapping
