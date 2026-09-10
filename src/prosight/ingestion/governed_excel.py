"""Deterministic, tenant-bound Excel profiling and normalization.

Workbook content is untrusted data.  This module never evaluates formulas,
never calls a model, and never publishes operational records.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook


TRANSFORMATION_VERSION = "excel-normalize-v1"
ISO_CURRENCY_CODES = frozenset(
    "AED AFN ALL AMD ANG AOA ARS AUD AWG AZN BAM BBD BDT BGN BHD BIF BMD BND "
    "BOB BOV BRL BSD BTN BWP BYN BZD CAD CDF CHE CHF CHW CLF CLP CNY COP COU "
    "CRC CUC CUP CVE CZK DJF DKK DOP DZD EGP ERN ETB EUR FJD FKP GBP GEL GHS "
    "GIP GMD GNF GTQ GYD HKD HNL HRK HTG HUF IDR ILS INR IQD IRR ISK JMD JOD "
    "JPY KES KGS KHR KMF KPW KRW KWD KYD KZT LAK LBP LKR LRD LSL LYD MAD MDL "
    "MGA MKD MMK MNT MOP MRU MUR MVR MWK MXN MXV MYR MZN NAD NGN NIO NOK NPR "
    "NZD OMR PAB PEN PGK PHP PKR PLN PYG QAR RON RSD RUB RWF SAR SBD SCR SDG SEK "
    "SGD SHP SLE SLL SOS SRD SSP STN SVC SYP SZL THB TJS TMT TND TOP TRY TTD TWD "
    "TZS UAH UGX USD USN UYI UYU UYW UZS VED VES VND VUV WST XAF XAG XAU XBA "
    "XBB XBC XBD XCD XCG XDR XOF XPD XPF XPT XSU XTS XUA XXX YER ZAR ZMW ZWL".split()
)


def _canonical(value: Any) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()


def checksum(value: Any) -> str:
    return hashlib.sha256(_canonical(value)).hexdigest()


def file_checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@dataclass(frozen=True)
class SourceFileMetadata:
    id: str
    organization_id: str
    source_kind: str
    filename: str
    checksum_sha256: str
    byte_size: int
    uploaded_by: str

    def __post_init__(self) -> None:
        uuid.UUID(self.id)
        uuid.UUID(self.organization_id)
        uuid.UUID(self.uploaded_by)
        if Path(self.filename).name != self.filename or not self.filename:
            raise ValueError("Source filename must be a basename")
        if self.byte_size < 1:
            raise ValueError("Source file must not be empty")
        source_dedupe_key(self.organization_id, self.source_kind, self.checksum_sha256)

    @property
    def dedupe_key(self) -> str:
        return source_dedupe_key(self.organization_id, self.source_kind, self.checksum_sha256)


@dataclass(frozen=True)
class FieldSpec:
    kind: Literal[
        "text", "date", "datetime", "decimal", "integer", "currency", "uuid", "boolean"
    ]
    required: bool = False
    scale: int | None = None


@dataclass(frozen=True)
class SheetMapping:
    entity_type: str
    sheet_name: str
    columns: dict[str, str]
    fields: dict[str, FieldSpec]
    business_key: tuple[str, ...]
    sheet_aliases: tuple[str, ...] = ()
    column_aliases: dict[str, tuple[str, ...]] = field(default_factory=dict)

    def __post_init__(self) -> None:
        if not self.entity_type or not self.sheet_name or not self.business_key:
            raise ValueError("A sheet mapping needs an entity, sheet, and business key")
        if set(self.columns) != set(self.fields):
            raise ValueError("Mapped columns and field specifications must match exactly")
        if not (set(self.business_key) - {"organization_id", "project_id"}) <= set(self.fields):
            raise ValueError("Business-key fields must be mapped")
        if not set(self.column_aliases) <= set(self.columns):
            raise ValueError("Column aliases must belong to mapped canonical fields")
        sources: dict[str, str] = {}
        for target, primary in self.columns.items():
            for name in (primary, *self.column_aliases.get(target, ())):
                normalized = _normalize_header(name)
                if not normalized:
                    raise ValueError("Source column labels cannot be empty")
                previous = sources.get(normalized)
                if previous and previous != target:
                    raise ValueError("One source column cannot map to multiple target fields")
                sources[normalized] = target
        sheet_names = [_normalize_header(name) for name in (self.sheet_name, *self.sheet_aliases)]
        if any(not name for name in sheet_names) or len(sheet_names) != len(set(sheet_names)):
            raise ValueError("Sheet labels must be non-empty and unique")

    def source_labels(self, target: str) -> tuple[str, ...]:
        return (self.columns[target], *self.column_aliases.get(target, ()))

    def serializable(self) -> dict[str, Any]:
        payload = {
            "entity_type": self.entity_type,
            "sheet_name": self.sheet_name,
            "columns": self.columns,
            "fields": {
                name: {"kind": spec.kind, "required": spec.required, "scale": spec.scale}
                for name, spec in self.fields.items()
            },
            "business_key": self.business_key,
        }
        if self.sheet_aliases:
            payload["sheet_aliases"] = self.sheet_aliases
        if any(self.column_aliases.values()):
            payload["column_aliases"] = self.column_aliases
        return payload


@dataclass(frozen=True)
class OrganizationMapping:
    organization_id: str
    mapping_profile_id: str
    mapping_version_id: str
    version_no: int
    sheets: tuple[SheetMapping, ...]
    catalog_version: int | None = None
    mapping_checksum: str = field(init=False)

    def __post_init__(self) -> None:
        uuid.UUID(self.organization_id)
        uuid.UUID(self.mapping_profile_id)
        uuid.UUID(self.mapping_version_id)
        if self.version_no < 1 or not self.sheets:
            raise ValueError("Mapping versions start at 1 and contain at least one sheet")
        names = [mapping.sheet_name.casefold() for mapping in self.sheets]
        if len(names) != len(set(names)):
            raise ValueError("A mapping version cannot map the same sheet twice")
        claimed_sheet_labels: dict[str, int] = {}
        for index, mapping in enumerate(self.sheets):
            for label in (mapping.sheet_name, *mapping.sheet_aliases):
                normalized = _normalize_header(label)
                previous = claimed_sheet_labels.get(normalized)
                if previous is not None and previous != index:
                    raise ValueError("A sheet label cannot map to multiple target entities")
                claimed_sheet_labels[normalized] = index
        payload = {
            "organization_id": self.organization_id,
            "mapping_profile_id": self.mapping_profile_id,
            "mapping_version_id": self.mapping_version_id,
            "version_no": self.version_no,
            "sheets": [mapping.serializable() for mapping in self.sheets],
        }
        if self.catalog_version is not None:
            if self.catalog_version < 1:
                raise ValueError("Catalog versions start at 1")
            payload["catalog_version"] = self.catalog_version
        object.__setattr__(self, "mapping_checksum", checksum(payload))


class MappingRegistry:
    """Small immutable-version registry used by services and offline tests."""

    def __init__(self) -> None:
        self._profiles: dict[tuple[str, str, int], OrganizationMapping] = {}

    def register(self, profile: OrganizationMapping) -> OrganizationMapping:
        key = (profile.organization_id, profile.mapping_profile_id, profile.version_no)
        existing = self._profiles.get(key)
        if existing and existing.mapping_checksum != profile.mapping_checksum:
            raise ValueError("An organization mapping version is immutable")
        self._profiles[key] = profile
        return profile


@dataclass(frozen=True)
class ReferenceCatalog:
    organization_ids: frozenset[str]
    project_organizations: dict[str, str]

    def validate(self, organization_id: str, project_id: str | None) -> str | None:
        if organization_id not in self.organization_ids:
            return "organization_id does not reference an authorized organization"
        if project_id is not None and self.project_organizations.get(project_id) != organization_id:
            return "project_id does not reference the same organization"
        return None


@dataclass(frozen=True)
class ValidationIssue:
    code: str
    message: str
    sheet_name: str | None = None
    row_number: int | None = None
    field: str | None = None


@dataclass(frozen=True)
class NormalizedRow:
    entity_type: str
    values: dict[str, Any]
    lineage: dict[str, Any]


@dataclass(frozen=True)
class ValidationPreview:
    organization_id: str
    import_batch_id: str
    source_file_id: str
    mapping_profile_id: str
    mapping_version_id: str
    mapping_version_no: int
    source_checksum: str
    profile_checksum: str
    input_profile_checksum: str
    normalized_preview_checksum: str
    validation_checksum: str
    rows: tuple[NormalizedRow, ...]
    issues: tuple[ValidationIssue, ...]
    sheet_classifications: dict[str, str]
    formula_cells: tuple[str, ...]

    @property
    def valid(self) -> bool:
        return not self.issues


def _normalize_header(value: Any) -> str:
    """Normalize human labels without treating their text as executable input."""
    return re.sub(r"[^\w]+", " ", str(value or "").strip(), flags=re.UNICODE).strip().casefold()


def matching_sheet_names(sheet_names: list[str], mapping: SheetMapping) -> list[str]:
    accepted = {_normalize_header(name) for name in (mapping.sheet_name, *mapping.sheet_aliases)}
    return [name for name in sheet_names if _normalize_header(name) in accepted]


def _decimal(value: Any, scale: int | None) -> Decimal:
    if isinstance(value, bool):
        raise ValueError("must be an exact numeric value")
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("must be finite")
        value = str(value)
    try:
        result = Decimal(value if isinstance(value, (str, int, Decimal)) else str(value))
    except (InvalidOperation, ValueError):
        raise ValueError("must be an exact numeric value") from None
    if not result.is_finite():
        raise ValueError("must be finite")
    if scale is not None:
        quantum = Decimal(1).scaleb(-scale)
        quantized = result.quantize(quantum)
        if quantized != result:
            raise ValueError(f"must have at most {scale} decimal places")
        result = quantized
    return result


def convert_value(value: Any, spec: FieldSpec) -> Any:
    if value is None or (isinstance(value, str) and not value.strip()):
        if spec.required:
            raise ValueError("is required")
        return None
    if spec.kind == "text":
        result = re.sub(r"\s+", " ", str(value).strip())
        if spec.required and not result:
            raise ValueError("is required")
        return result
    if spec.kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        try:
            return date.fromisoformat(str(value).strip()).isoformat()
        except ValueError:
            raise ValueError("must be an ISO date (YYYY-MM-DD)") from None
    if spec.kind == "datetime":
        if isinstance(value, datetime):
            return value.isoformat()
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time()).isoformat()
        candidate = str(value).strip().replace("Z", "+00:00")
        try:
            return datetime.fromisoformat(candidate).isoformat()
        except ValueError:
            raise ValueError("must be an ISO date-time") from None
    if spec.kind == "decimal":
        return _decimal(value, spec.scale)
    if spec.kind == "integer":
        number = _decimal(value, 0)
        return int(number)
    if spec.kind == "currency":
        code = str(value).strip().upper()
        if code not in ISO_CURRENCY_CODES:
            raise ValueError("must be a current ISO 4217 currency code")
        return code
    if spec.kind == "uuid":
        try:
            return str(uuid.UUID(str(value).strip()))
        except ValueError:
            raise ValueError("must be a UUID") from None
    if spec.kind == "boolean":
        if isinstance(value, bool):
            return value
        normalized = str(value).strip().casefold()
        if normalized in {"true", "yes", "y", "1"}:
            return True
        if normalized in {"false", "no", "n", "0"}:
            return False
        raise ValueError("must be true/false, yes/no, or 1/0")
    raise ValueError(f"Unsupported field type: {spec.kind}")


def _json_ready(value: Any) -> Any:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def validate_workbook(
    path: Path,
    profile: OrganizationMapping,
    references: ReferenceCatalog,
    *,
    source_file_id: str,
    import_batch_id: str,
    max_rows: int = 10_000,
    max_bytes: int = 10 * 1024 * 1024,
    max_uncompressed_bytes: int = 100 * 1024 * 1024,
    expected_totals: dict[tuple[str, str], Decimal] | None = None,
    project_id: str | None = None,
) -> ValidationPreview:
    """Profile and normalize a workbook without evaluating any formula."""
    byte_size = path.stat().st_size
    if path.suffix.lower() != ".xlsx" or byte_size < 1 or byte_size > max_bytes:
        raise ValueError("Only non-empty XLSX workbooks are accepted")
    if not zipfile.is_zipfile(path):
        raise ValueError("Workbook is not a valid XLSX package")
    with zipfile.ZipFile(path) as package:
        members = package.infolist()
        if len(members) > 10_000 or sum(member.file_size for member in members) > max_uncompressed_bytes:
            raise ValueError("Workbook package exceeds safe expansion limits")
        if any(member.flag_bits & 0x1 for member in members):
            raise ValueError("Encrypted workbook members are not supported")
        if any(member.filename.casefold().endswith("vbaproject.bin") for member in members):
            raise ValueError("Macro-enabled workbooks are not supported")
    uuid.UUID(source_file_id)
    uuid.UUID(import_batch_id)
    scoped_project_id = str(uuid.UUID(project_id)) if project_id else None
    source_sha = file_checksum(path)
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    issues: list[ValidationIssue] = []
    normalized: list[NormalizedRow] = []
    formulas: list[str] = []
    classifications = {name: "unmapped" for name in book.sheetnames}
    populated_rows = 0
    seen_business_keys: dict[str, set[tuple[Any, ...]]] = {}
    try:
        for mapping in profile.sheets:
            sheet_matches = matching_sheet_names(book.sheetnames, mapping)
            if not sheet_matches:
                issues.append(ValidationIssue("missing_sheet", "Mapped sheet is missing", mapping.sheet_name))
                continue
            if len(sheet_matches) > 1:
                issues.append(ValidationIssue(
                    "ambiguous_sheet",
                    "Multiple workbook sheets match the same organization mapping: "
                    + ", ".join(sheet_matches),
                    mapping.sheet_name,
                ))
                continue
            source_sheet_name = sheet_matches[0]
            classifications[source_sheet_name] = mapping.entity_type
            sheet = book[source_sheet_name]
            iterator = sheet.iter_rows()
            header_cells = next(iterator, ())
            headers = [_normalize_header(cell.value) for cell in header_cells]
            nonempty_headers = [header for header in headers if header]
            if len(nonempty_headers) != len(set(nonempty_headers)):
                issues.append(ValidationIssue("duplicate_header", "Sheet headers must be unique", source_sheet_name))
                continue
            source_indexes: dict[str, int] = {}
            resolved_columns: dict[str, str] = {}
            for target, source in mapping.columns.items():
                accepted = {_normalize_header(label) for label in mapping.source_labels(target)}
                matches = [index for index, header in enumerate(headers) if header in accepted]
                if not matches:
                    if mapping.fields[target].required:
                        issues.append(ValidationIssue("missing_column", f"Mapped source column {source!r} is missing", source_sheet_name, field=target))
                elif len(matches) > 1:
                    issues.append(ValidationIssue("ambiguous_column", f"Multiple source columns match {target!r}", source_sheet_name, field=target))
                else:
                    source_indexes[target] = matches[0]
                    resolved_columns[target] = str(header_cells[matches[0]].value).strip()
            seen_keys = seen_business_keys.setdefault(mapping.entity_type, set())
            for row_number, cells in enumerate(iterator, start=2):
                if not any(cell.value is not None for cell in cells):
                    continue
                populated_rows += 1
                if populated_rows > max_rows:
                    raise ValueError(f"Workbook exceeds the {max_rows:,}-row limit")
                values: dict[str, Any] = {"organization_id": profile.organization_id}
                if scoped_project_id is not None:
                    values["project_id"] = scoped_project_id
                raw_values: dict[str, Any] = {}
                row_has_error = False
                for target, spec in mapping.fields.items():
                    index = source_indexes.get(target)
                    cell = cells[index] if index is not None and index < len(cells) else None
                    raw = None if cell is None else cell.value
                    raw_values[target] = _json_ready(raw)
                    if cell is not None and cell.data_type == "f":
                        address = f"{source_sheet_name}!{cell.coordinate}"
                        formulas.append(address)
                        issues.append(ValidationIssue("formula", "Formula cells are not accepted as import values", source_sheet_name, row_number, target))
                        row_has_error = True
                        continue
                    try:
                        converted = convert_value(raw, spec)
                        if converted is not None:
                            values[target] = converted
                    except ValueError as error:
                        issues.append(ValidationIssue("invalid_value", str(error), source_sheet_name, row_number, target))
                        row_has_error = True
                reference_error = references.validate(values["organization_id"], values.get("project_id"))
                if reference_error:
                    issues.append(ValidationIssue("reference", reference_error, source_sheet_name, row_number, "project_id"))
                    row_has_error = True
                if row_has_error:
                    continue
                key = tuple(values.get(name) for name in mapping.business_key)
                if any(value is None for value in key):
                    issues.append(ValidationIssue("business_key", "Business key cannot contain null", source_sheet_name, row_number))
                    continue
                if key in seen_keys:
                    issues.append(ValidationIssue("duplicate_business_key", "Duplicate business key in workbook", source_sheet_name, row_number))
                    continue
                seen_keys.add(key)
                normalized_values = {name: _json_ready(value) for name, value in values.items()}
                lineage = {
                    "organization_id": profile.organization_id,
                    "source_file_id": source_file_id,
                    "import_batch_id": import_batch_id,
                    "sheet_name": source_sheet_name,
                    "row_number": row_number,
                    "mapping_profile_id": profile.mapping_profile_id,
                    "mapping_version_id": profile.mapping_version_id,
                    "mapping_version_no": profile.version_no,
                    "mapping_checksum": profile.mapping_checksum,
                    "transformation_version": TRANSFORMATION_VERSION,
                    "raw_row_checksum": checksum(raw_values),
                    "normalized_row_checksum": checksum(normalized_values),
                    "raw_values": raw_values,
                    "column_mappings": resolved_columns,
                }
                normalized.append(NormalizedRow(mapping.entity_type, normalized_values, lineage))
    finally:
        book.close()

    if not normalized and not issues:
        issues.append(ValidationIssue("no_records", "Workbook contains no publishable records"))

    expected_totals = expected_totals or {}
    for (entity_type, field_name), expected in sorted(expected_totals.items()):
        actual = sum(
            (Decimal(str(row.values[field_name])) for row in normalized
             if row.entity_type == entity_type and row.values.get(field_name) is not None),
            Decimal("0"),
        )
        if actual != expected:
            issues.append(ValidationIssue(
                "reconciliation", f"{entity_type}.{field_name} total {actual} does not reconcile to {expected}"
            ))

    rows_payload = [{"entity_type": row.entity_type, "values": row.values, "lineage": row.lineage} for row in normalized]
    issues_payload = [issue.__dict__ for issue in issues]
    input_profile = checksum({
        "source_checksum": source_sha,
        "mapping_profile_id": profile.mapping_profile_id,
        "mapping_version_id": profile.mapping_version_id,
        "mapping_version_no": profile.version_no,
        "profile_checksum": profile.mapping_checksum,
    })
    preview_sha = checksum(rows_payload)
    validation_sha = checksum({"input_profile_checksum": input_profile, "preview_checksum": preview_sha, "issues": issues_payload})
    return ValidationPreview(
        organization_id=profile.organization_id,
        import_batch_id=import_batch_id,
        source_file_id=source_file_id,
        mapping_profile_id=profile.mapping_profile_id,
        mapping_version_id=profile.mapping_version_id,
        mapping_version_no=profile.version_no,
        source_checksum=source_sha,
        profile_checksum=profile.mapping_checksum,
        input_profile_checksum=input_profile,
        normalized_preview_checksum=preview_sha,
        validation_checksum=validation_sha,
        rows=tuple(normalized),
        issues=tuple(issues),
        sheet_classifications=classifications,
        formula_cells=tuple(sorted(set(formulas))),
    )


def source_dedupe_key(organization_id: str, source_kind: str, source_sha256: str) -> str:
    """Return the contract's organization-scoped source-file identity."""
    uuid.UUID(organization_id)
    if not re.fullmatch(r"[0-9a-f]{64}", source_sha256):
        raise ValueError("source_sha256 must be a lowercase SHA-256 digest")
    if not source_kind.strip():
        raise ValueError("source_kind is required")
    return checksum([organization_id, source_kind.strip().casefold(), source_sha256])
