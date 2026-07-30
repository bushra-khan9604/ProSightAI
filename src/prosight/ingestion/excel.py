"""Canonical Excel parsing and reviewed mapping-preview generation."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_COLUMNS = [
    "code", "name", "status", "client", "location", "contract_value_usd",
    "planned_start", "planned_finish", "revised_finish", "reporting_date",
    "baseline_progress", "revised_progress", "actual_progress",
]
CANONICAL_SHEETS = {"Projects", "Contacts", "Activities", "Manpower", "Equipment", "Milestones"}
RELATED_SCHEMAS = {
    "Contacts": (["project_code", "name", "project_role", "email", "mobile"], "contacts"),
    "Activities": (["project_code", "activity"], "activities"),
    "Manpower": (["project_code", "designation", "count"], "manpower"),
    "Equipment": (["project_code", "type", "count"], "equipment"),
    "Milestones": (["project_code", "name", "status"], "milestones"),
}


def _json_value(value: Any) -> Any:
    """Convert workbook values into JSON-safe values."""
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def _merge_related_sheets(workbook: Any, projects: list[dict[str, Any]], source_name: str) -> None:
    """Merge canonical child sheets into their parent project payloads."""
    by_code = {project["code"]: project for project in projects}
    for sheet_name, (required, target) in RELATED_SCHEMAS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        if any(column not in headers for column in required):
            raise ValueError(f"{sheet_name} sheet must contain: {', '.join(required)}")
        for row_number, values in enumerate(rows, start=2):
            record = {header: _json_value(value) for header, value in zip(headers, values)}
            code = record.pop("project_code", None)
            if not code and not any(value is not None for value in values):
                continue
            if code not in by_code:
                raise ValueError(f"{sheet_name} row {row_number} references unknown project {code}")
            if target == "activities":
                by_code[code][target].append(record["activity"])
            else:
                if "count" in record:
                    record["count"] = int(record["count"])
                by_code[code][target].append(record)
            by_code[code]["sources"].append(f"{source_name}, {sheet_name} row {row_number}")


def preview_workbook(path: Path, project_code: str, max_rows: int = 10_000) -> dict[str, Any]:
    """Validate canonical workbooks or return a reviewed mapping requirement."""
    # Normal mode computes worksheet dimensions when an otherwise valid XLSX
    # omits the optional cached <dimension> element. Read-only mode can expose
    # ``max_row=None`` for such files and incorrectly reject them during upload.
    workbook = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    populated_rows = sum(max((sheet.max_row or 1) - 1, 0) for sheet in workbook.worksheets)
    if populated_rows > max_rows:
        workbook.close()
        raise ValueError(f"Workbook exceeds the {max_rows:,}-row limit")
    sheets = set(workbook.sheetnames)
    if "Projects" not in sheets:
        result = {
            "mapping_required": True,
            "available_sheets": workbook.sheetnames,
            "suggested_mapping": {
                sheet.title: [str(cell.value or "").strip() for cell in next(sheet.iter_rows(max_row=1))]
                for sheet in workbook.worksheets if sheet.max_row
            },
            "projects": [],
            "warnings": ["Non-canonical workbook requires Admin mapping review"],
        }
        workbook.close()
        return result
    sheet = workbook["Projects"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    missing = [column for column in PROJECT_COLUMNS if column not in headers]
    if missing:
        workbook.close()
        raise ValueError(f"Projects sheet is missing columns: {', '.join(missing)}")
    projects: list[dict[str, Any]] = []
    errors: list[str] = []
    for row_number, values in enumerate(rows, start=2):
        record = {header: _json_value(value) for header, value in zip(headers, values)}
        if not any(value is not None for value in values):
            continue
        try:
            if record["status"] not in {"active", "completed", "future"}:
                raise ValueError("status must be active, completed, or future")
            for field in ("baseline_progress", "revised_progress", "actual_progress"):
                record[field] = float(record[field])
            record["contract_value_usd"] = float(record["contract_value_usd"])
            record.setdefault("contacts", [])
            record.setdefault("activities", [])
            record.setdefault("manpower", [])
            record.setdefault("equipment", [])
            record.setdefault("milestones", [])
            record.setdefault("total_manhours", 0)
            record.setdefault("sources", [f"{path.name}, Projects row {row_number}"])
            projects.append(record)
        except (TypeError, ValueError) as error:
            errors.append(f"Projects row {row_number}: {error}")
    if errors:
        workbook.close()
        raise ValueError("; ".join(errors[:20]))
    try:
        _merge_related_sheets(workbook, projects, path.name)
    except ValueError:
        workbook.close()
        raise
    result = {
        "mapping_required": False,
        "available_sheets": workbook.sheetnames,
        "projects": projects,
        "project_code": project_code,
        "warnings": sorted(CANONICAL_SHEETS - sheets),
        "row_count": populated_rows,
    }
    workbook.close()
    return result


def preview_mapped_workbook(
    path: Path, project_code: str, mapping: dict[str, Any]
) -> dict[str, Any]:
    """Apply an AI-proposed column mapping and return a human-reviewable preview."""
    # See preview_workbook: normal mode supports valid files without cached
    # worksheet dimensions and remains safe under the enforced upload limits.
    workbook = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    sheet_name = mapping.get("sheet")
    columns = mapping.get("columns", {})
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError("Mapped sheet was not found")
    missing_mappings = [column for column in PROJECT_COLUMNS if column not in columns]
    if missing_mappings:
        workbook.close()
        raise ValueError(f"AI mapping did not map: {', '.join(missing_mappings)}")
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    header_indexes = {header: index for index, header in enumerate(headers)}
    if any(source not in header_indexes for source in columns.values()):
        workbook.close()
        raise ValueError("AI mapping referenced a column that does not exist")
    projects = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None for value in values):
            continue
        record = {
            canonical: _json_value(values[header_indexes[source]])
            for canonical, source in columns.items()
        }
        if record["status"] not in {"active", "completed", "future"}:
            workbook.close()
            raise ValueError(f"Mapped row {row_number} has an invalid status")
        for field in ("contract_value_usd", "baseline_progress", "revised_progress", "actual_progress"):
            record[field] = float(record[field])
        record.update(
            contacts=[], activities=[], manpower=[], equipment=[], milestones=[],
            total_manhours=0,
            sources=[f"{path.name}, {sheet_name} row {row_number} (AI-mapped)"],
        )
        projects.append(record)
    workbook.close()
    return {
        "mapping_required": True,
        "mapping": mapping,
        "projects": projects,
        "project_code": project_code,
        "row_count": len(projects),
        "warnings": ["AI-proposed mapping requires Admin approval"],
    }
