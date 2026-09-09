"""Schema-driven parsing for employee, attendance, and deployment workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from .schemas import normalize_header, semantic_match_headers
from .workforce_mapping import load_workforce_mapping
from ..workforce_models import (
    AttendanceRecord,
    DeploymentForecastRecord,
    DirectAllocationRecord,
    EmployeeRecord,
    PayrollRecord,
    SubcontractCrewRecord,
    TrainingRecord,
)


@dataclass(frozen=True)
class SheetSpec:
    sheet: str
    collection: str
    model: type[BaseModel]
    fields: dict[str, str]
    keys: tuple[str, ...]
    derived: frozenset[str] = frozenset()
    employee_field: str | None = None
    project_field: str | None = None


SPECS = {
    "Employees": SheetSpec(
        "Employees", "employees", EmployeeRecord,
        {
            "Employee ID": "employee_id", "Synthetic name": "synthetic_name",
            "Role": "role", "Department": "department",
            "Home project ID": "home_project_id", "Employment type": "employment_type",
            "Join date": "join_date", "Basic AED": "basic_aed",
            "Allowance AED": "allowance_aed", "Gross monthly AED": "gross_monthly_aed",
            "Status": "status", "Data origin": "data_origin",
        },
        ("employee_id",), frozenset({"gross_monthly_aed"}),
        project_field="home_project_id",
    ),
    "Training": SheetSpec(
        "Training", "training", TrainingRecord,
        {
            "Training ID": "training_id", "Employee ID": "employee_id",
            "Course": "course", "Completed date": "completed_date",
            "Expiry date": "expiry_date", "Status": "status", "Evidence ID": "evidence_id",
        },
        ("training_id",), employee_field="employee_id",
    ),
    "Attendance": SheetSpec(
        "Attendance", "attendance", AttendanceRecord,
        {
            "Timesheet ID": "timesheet_id", "Employee ID": "employee_id",
            "Project ID": "project_id", "Work date": "work_date",
            "Attendance status": "attendance_status", "Regular hours": "regular_hours",
            "OT hours": "ot_hours", "Total hours": "total_hours",
        },
        ("timesheet_id",), frozenset({"total_hours"}), "employee_id", "project_id",
    ),
    "Payroll": SheetSpec(
        "Payroll", "payroll", PayrollRecord,
        {
            "Payroll ID": "payroll_id", "Employee ID": "employee_id",
            "Project ID": "project_id", "Month": "month", "Basic AED": "basic_aed",
            "Allowance AED": "allowance_aed", "OT hours": "ot_hours",
            "OT rate AED": "ot_rate_aed", "OT pay AED": "ot_pay_aed",
            "Gross AED": "gross_aed", "Employer burden AED": "employer_burden_aed",
            "Total cost AED": "total_cost_aed", "Payment date": "payment_date",
            "Days worked": "days_worked", "Paid leave days": "paid_leave_days",
        },
        ("payroll_id",), frozenset({"ot_pay_aed", "gross_aed", "total_cost_aed"}),
        "employee_id", "project_id",
    ),
    "Direct allocation": SheetSpec(
        "Direct allocation", "direct_allocations", DirectAllocationRecord,
        {
            "Employee ID": "employee_id", "Project ID": "project_id", "Month": "month",
            "FTE allocation": "fte_allocation", "Regular hours": "regular_hours",
            "OT hours": "ot_hours", "Payroll cost AED": "payroll_cost_aed", "Basis": "basis",
        },
        ("employee_id", "project_id", "month"), employee_field="employee_id",
        project_field="project_id",
    ),
    "Subcontract crews": SheetSpec(
        "Subcontract crews", "subcontract_crews", SubcontractCrewRecord,
        {
            "Subcontract ID": "subcontract_id", "Project ID": "project_id",
            "Vendor ID": "vendor_id", "Month": "month", "Trade": "trade",
            "Planned workers": "planned_workers", "Actual workers": "actual_workers",
            "Worker variance": "worker_variance",
            "Available labor hours": "available_labor_hours", "Definition": "definition",
        },
        ("subcontract_id", "project_id", "month"),
        frozenset({"worker_variance", "available_labor_hours"}),
        project_field="project_id",
    ),
    "Forecast deployment": SheetSpec(
        "Forecast deployment", "deployment_forecasts", DeploymentForecastRecord,
        {
            "Project ID": "project_id", "Month": "month",
            "Direct headcount": "direct_headcount",
            "Subcontract workers": "subcontract_workers", "Basis": "basis",
        },
        ("project_id", "month"), project_field="project_id",
    ),
}

WORKFORCE_DATASETS = {
    "employees_training": ("Employees", "Training"),
    "attendance_payroll": ("Attendance", "Payroll"),
    "manpower_deployment": ("Direct allocation", "Subcontract crews", "Forecast deployment"),
}

WORKFORCE_TABLES = {
    "employees": ("employees", ("employee_id",)),
    "training": ("employee_training", ("training_id",)),
    "attendance": ("attendance_records", ("timesheet_id",)),
    "payroll": ("payroll_records", ("payroll_id",)),
    "direct_allocations": ("direct_allocations", ("employee_id", "project_id", "month")),
    "subcontract_crews": ("subcontract_crews", ("subcontract_id", "project_id", "month")),
    "deployment_forecasts": ("deployment_forecasts", ("project_id", "month")),
}

# Organizational assignments are valid only for workforce records, not projects.
HEAD_OFFICE_ALIASES = {'ho', 'headoffice'}


def is_head_office(value):
    return normalize_header(value) in HEAD_OFFICE_ALIASES

HEADER_ALIASES = {
    "Employee ID": {"employee code", "emp code", "staff id"},
    "Synthetic name": {"employee name", "staff name", "name"},
    "Home project ID": {"home project", "home project code"},
    "Employment type": {"contract type", "worker type"},
    "Join date": {"joining date", "date joined"},
    "Training ID": {"course record id", "training record id"},
    "Completed date": {"completion date", "date completed"},
    "Expiry date": {"expiration date", "valid until"},
    "Evidence ID": {"certificate id", "document id"},
    "Timesheet ID": {"attendance id", "time entry id"},
    "Project ID": {"project code", "job number"},
    "Work date": {"attendance date", "timesheet date"},
    "Attendance status": {"day status", "presence status"},
    "OT hours": {"overtime hours", "overtime hrs"},
    "OT rate AED": {"overtime rate aed", "ot hourly rate aed"},
    "OT pay AED": {"overtime pay aed"},
    "Employer burden AED": {"employer cost aed", "burden aed"},
    "Paid leave days": {"leave days", "paid leave"},
    "FTE allocation": {"fte", "allocation fte"},
    "Payroll cost AED": {"labor cost aed", "employee cost aed"},
    "Subcontract ID": {"crew id", "subcontract crew id"},
    "Vendor ID": {"vendor code", "subcontractor id"},
    "Available labor hours": {"available labour hours", "crew capacity hours"},
    "Forecast deployment": {"deployment forecast"},
}


def _sheet(workbook, expected: str):
    matches = [name for name in workbook.sheetnames if normalize_header(name) == normalize_header(expected)]
    if len(matches) != 1:
        raise ValueError(f"Workbook must contain one {expected} sheet")
    return workbook[matches[0]]


def _headers(sheet, spec: SheetSpec, rule) -> tuple[list[str], int, dict[str, Any]]:
    expected = list(spec.fields)
    aliases = {label: set(rule.fields[field].aliases) | {field} for label, field in spec.fields.items()}
    best = None
    first = rule.header_row or (rule.data_start_row if rule.data_start_row else 1)
    last = first if rule.header_row or rule.data_start_row else min(sheet.max_row or 1, 30)
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=first, max_row=last, values_only=True), start=first
    ):
        actual = [str(value or "").strip() for value in values]
        if rule.data_start_row and not rule.header_row:
            actual = [''] * len(actual)
        # Explicit positions allow unlabeled columns without guessing their meaning.
        for label, field in spec.fields.items():
            column = rule.fields[field].column
            if column:
                actual.extend([''] * max(0, column - len(actual)))
                actual[column - 1] = label
        mapped, mapping = semantic_match_headers(actual, expected, aliases)
        for label in expected:
            if sum(normalize_header(value) in {normalize_header(label), *(normalize_header(a) for a in aliases[label])}
                   for value in actual if value) > 1:
                raise ValueError(f'{sheet.title}: ambiguous duplicate column for {label}')
        score = len(mapping["matches"])
        if best is None or score > best[0]:
            best = (score, mapped, row_number, mapping, actual)
    if best is None or best[0] == 0:
        raise ValueError(f"{spec.sheet} does not contain a recognizable header row")
    optional = {header for header, field in spec.fields.items()
                if field in spec.derived or not spec.model.model_fields[field].is_required()}
    missing = [header for header in best[3]["missing"] if header not in optional]
    if missing:
        raise ValueError(f"{spec.sheet} requires mapping for: {', '.join(missing)}")
    best[3]["missing"] = missing
    best[3]["mapping_required"] = False
    best[3]["schema"] = spec.model.__name__
    best[3]["source_sheet"] = sheet.title
    best[3]["header_row"] = rule.header_row or (None if rule.data_start_row else best[2])
    best[3]["configured_columns"] = {field: entry.column for field, entry in rule.fields.items() if entry.column}
    matched_sources = {item["source"] for item in best[3]["matches"].values()}
    best[3]["ignored_columns"] = [
        value for value in best[4] if value and value not in matched_sources
    ]
    return best[1], (rule.data_start_row - 1 if rule.data_start_row else best[2]), best[3]


def _value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        return value.strip()
    return value


def _derive(spec: SheetSpec, record: dict[str, Any]) -> None:
    if spec.collection == "employees":
        record["gross_monthly_aed"] = float(record.get("basic_aed") or 0) + float(record.get("allowance_aed") or 0)
    elif spec.collection == "attendance":
        record["total_hours"] = float(record.get("regular_hours") or 0) + float(record.get("ot_hours") or 0)
    elif spec.collection == "payroll":
        record["ot_pay_aed"] = float(record.get("ot_hours") or 0) * float(record.get("ot_rate_aed") or 0)
        record["gross_aed"] = float(record.get("basic_aed") or 0) + float(record.get("allowance_aed") or 0)
        record["total_cost_aed"] = record["ot_pay_aed"] + record["gross_aed"] + float(record.get("employer_burden_aed") or 0)
    elif spec.collection == "subcontract_crews":
        actual = int(record.get("actual_workers") or 0)
        record["worker_variance"] = actual - int(record.get("planned_workers") or 0)
        record["available_labor_hours"] = actual * 208


def parse_workforce_workbook(
    path: Path,
    resolve_project: Callable[[str], str | None],
    resolve_employee: Callable[[str], str | None],
    dataset: str,
) -> dict[str, Any]:
    """Map a workbook to strict records without evaluating workbook formulas."""
    sheet_names = WORKFORCE_DATASETS.get(dataset)
    if not sheet_names:
        raise ValueError("Unsupported workforce import dataset")
    profile = load_workforce_mapping(SPECS)
    formulas = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    values = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    parsed: dict[str, Any] = {SPECS[name].collection: [] for name in sheet_names}
    parsed.update(counts={}, semantic_mappings={}, warnings=[], mapping_profile=profile.profile_id)
    imported_employees: dict[str, str] = {}
    try:
        assigned = set()
        for expected_name in sheet_names:
            spec = SPECS[expected_name]
            rule = profile.sheets[expected_name]
            matches = [sheet for sheet in formulas if normalize_header(sheet.title) in
                       {normalize_header(alias) for alias in rule.aliases}]
            if not matches and not any(field.column for field in rule.fields.values()):
                for sheet in formulas:
                    if sheet.title in assigned:
                        continue
                    try:
                        _headers(sheet, spec, rule)
                        matches.append(sheet)
                    except ValueError:
                        continue
            if not matches:
                continue
            if len(matches) != 1 or matches[0].title in assigned:
                raise ValueError(f'Ambiguous worksheets for {expected_name}; configure unique sheet aliases')
            formula_sheet = matches[0]
            assigned.add(formula_sheet.title)
            value_sheet = values[formula_sheet.title]
            mapped_headers, header_row, mapping = _headers(formula_sheet, spec, rule)
            parsed["semantic_mappings"][expected_name] = mapping
            seen: set[tuple[str, ...]] = set()
            formula_rows = formula_sheet.iter_rows(min_row=header_row + 1)
            value_rows = value_sheet.iter_rows(min_row=header_row + 1)
            for row_number, (source_cells, value_cells) in enumerate(
                zip(formula_rows, value_rows), start=header_row + 1
            ):
                if not any(cell.value not in (None, "") for cell in source_cells):
                    continue
                record: dict[str, Any] = {}
                for header, source_cell, value_cell in zip(mapped_headers, source_cells, value_cells):
                    field = spec.fields.get(header)
                    if not field:
                        continue
                    if source_cell.data_type == "f":
                        if field not in spec.derived:
                            raise ValueError(f"{expected_name} row {row_number}: formulas are not allowed in {header}")
                        continue
                    record[field] = _value(value_cell.value)
                if spec.project_field and is_head_office(record.get(spec.project_field)):
                    record[spec.project_field] = 'HO'
                _derive(spec, record)
                try:
                    clean = spec.model.model_validate(record).model_dump(mode="json")
                except ValidationError as error:
                    details = "; ".join(
                        f"{'.'.join(str(value) for value in issue['loc'])}: {issue['msg']}"
                        for issue in error.errors()
                    )
                    raise ValueError(f"{expected_name} row {row_number}: {details}") from error
                key = tuple(str(clean[field]) for field in spec.keys)
                folded = tuple(value.casefold() for value in key)
                if folded in seen:
                    raise ValueError(f"{expected_name} row {row_number}: duplicate key {' / '.join(key)}")
                seen.add(folded)
                if spec.project_field:
                    source_project = clean[spec.project_field]
                    project_code = 'HO' if source_project == 'HO' else resolve_project(source_project)
                    if not project_code:
                        raise ValueError(f"{expected_name} row {row_number}: unknown project {source_project}")
                    clean[spec.project_field] = project_code
                if spec.employee_field:
                    employee_id = clean[spec.employee_field]
                    canonical_employee = imported_employees.get(employee_id.casefold()) or resolve_employee(employee_id)
                    if not canonical_employee:
                        raise ValueError(f"{expected_name} row {row_number}: unknown employee {employee_id}")
                    clean[spec.employee_field] = canonical_employee
                clean.update(source_sheet=formula_sheet.title, source_row=row_number)
                parsed[spec.collection].append(clean)
                if spec.collection == "employees":
                    imported_employees[clean["employee_id"].casefold()] = clean["employee_id"]
            parsed["counts"][spec.collection] = len(parsed[spec.collection])
        if not any(parsed[SPECS[name].collection] for name in sheet_names):
            raise ValueError('No matching workforce data rows found; check the selected action and mapping profile')
        ignored = [name for name in formulas.sheetnames if name not in assigned]
        if ignored:
            parsed['warnings'].append('Worksheets not imported for this action: ' + ', '.join(ignored))
        if any(row.get(SPECS[name].project_field) == 'HO'
               for name in sheet_names for row in parsed[SPECS[name].collection]):
            parsed['warnings'].append('HO means Head Office: an organizational assignment, not a construction project.')
        derived = sorted({field for name in sheet_names for field in SPECS[name].derived})
        if derived:
            parsed["warnings"].append(
                "Recognized calculated columns were recomputed by ProSight; workbook formulas were not executed: "
                + ", ".join(derived)
            )
        return parsed
    finally:
        formulas.close()
        values.close()


def workforce_schema_manifest() -> dict[str, Any]:
    """Describe workbook-to-record mappings for client-side review."""
    profile = load_workforce_mapping(SPECS)
    return {
        dataset: {
            'mapping_profile': profile.profile_id,
            'organizational_assignments': {'HO': 'Head Office'},
            "sheets": [
                {
                    "sheet": name,
                    "collection": SPECS[name].collection,
                    "keys": list(SPECS[name].keys),
                    "columns": SPECS[name].fields,
                    "mapping": profile.sheets[name].model_dump(),
                    "derived_fields": sorted(SPECS[name].derived),
                }
                for name in sheets
            ]
        }
        for dataset, sheets in WORKFORCE_DATASETS.items()
    }
