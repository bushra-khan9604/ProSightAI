"""Transactional persistence for governed workbook ingestion.

Workbook cells are untrusted data.  They are bound only as SQL values and are
never allowed to choose an organization, project scope, SQL object, approval,
or publication target.  Construction writes are available only through the
database-owned atomic publisher.
"""

from __future__ import annotations

import json
import uuid
from contextlib import closing
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, quote_sheetname

from .governed_excel import (
    TRANSFORMATION_VERSION,
    _normalize_header,
    FieldSpec,
    OrganizationMapping,
    ReferenceCatalog,
    SheetMapping,
    SourceFileMetadata,
    ValidationPreview,
    checksum,
    file_checksum,
    matching_sheet_names,
    validate_workbook,
)
from .publication import ApprovalBinding, AtomicPublicationAdapter, PublicationRequest


class DuplicateSourceFile(ValueError):
    """The organization already registered the same source bytes and kind."""


@dataclass(frozen=True)
class MappingRecord:
    mapping: OrganizationMapping
    name: str
    source_kind: str
    description: str | None = None


@dataclass(frozen=True)
class PreparedImport:
    preview: ValidationPreview
    transformation_run_id: str


# This is the application mirror of the database-owned publication registry.
# It validates mapping definitions early; the private publisher remains the
# authority and repeats the check while holding its transaction lock.
PUBLICATION_BUSINESS_KEYS: dict[str, tuple[str, ...]] = {
    "projects": ("organization_id", "code"),
    "employees": ("organization_id", "employee_number"),
    "activities": ("project_id", "activity_code"),
    "invoices": ("project_id", "direction", "invoice_number"),
    "risks": ("project_id", "risk_number"),
    "rfis": ("project_id", "rfi_number"),
    "nonconformance_reports": ("project_id", "ncr_number"),
    "safety_incidents": ("project_id", "incident_number"),
    "daily_reports": ("project_id", "report_date", "shift_code"),
}


LOAD_MAPPING_SQL = """
select p.id as mapping_profile_id,p.organization_id,p.name,p.source_kind,p.description,
       v.id as mapping_version_id,v.version_no,v.mapping_specification,
       v.mapping_checksum,v.transformation_version
from ingestion.mapping_profiles p
join ingestion.mapping_profile_versions v
  on v.organization_id=p.organization_id and v.mapping_profile_id=p.id
where p.organization_id=%(organization_id)s::uuid
  and p.id=%(mapping_profile_id)s::uuid
  and v.version_no=%(version_no)s::integer
  and p.is_active
"""

LIST_MAPPINGS_SQL = """
select p.id as mapping_profile_id,p.organization_id,p.name,p.source_kind,p.description,
       v.id as mapping_version_id,v.version_no,v.mapping_specification,
       v.mapping_checksum,v.transformation_version
from ingestion.mapping_profiles p
join ingestion.mapping_profile_versions v
  on v.organization_id=p.organization_id and v.mapping_profile_id=p.id
where p.organization_id=%(organization_id)s::uuid and p.is_active
order by v.created_at desc,p.name,v.version_no desc
"""

LIST_IMPORTS_SQL = """
select b.id,b.organization_id,b.project_id,b.status,b.requester_user_id,b.requested_at,
       b.mapping_version_id,b.normalized_preview_checksum,b.validation_checksum,
       f.original_filename,f.checksum_sha256 as source_checksum,
       p.id as mapping_profile_id,p.name as mapping_name,v.version_no,v.mapping_checksum,
       r.id as transformation_run_id,r.row_count,
       coalesce((
         select jsonb_agg(jsonb_build_object(
             'id',staged.id,'entity_type',staged.target_entity_type,
             'row_number',staged.row_number,'values',staged.normalized_record,
             'status',staged.status
         ) order by staged.row_number,staged.id)
         from ingestion.staged_rows staged
         where staged.organization_id=b.organization_id
           and staged.import_batch_id=b.id and staged.transformation_run_id=r.id
       ),'[]'::jsonb) as rows,
       coalesce((
         select jsonb_agg(jsonb_build_object(
             'code',issue.issue_code,'message',issue.message,'severity',issue.severity,
             'details',issue.details
         ) order by issue.created_at,issue.id)
         from ingestion.validation_issues issue
         where issue.organization_id=b.organization_id
           and issue.import_batch_id=b.id and issue.transformation_run_id=r.id
       ),'[]'::jsonb) as issues
from ingestion.import_batches b
join ingestion.source_files f
  on f.organization_id=b.organization_id and f.import_batch_id=b.id
join ingestion.mapping_profile_versions v
  on v.organization_id=b.organization_id and v.id=b.mapping_version_id
join ingestion.mapping_profiles p
  on p.organization_id=v.organization_id and p.id=v.mapping_profile_id
join lateral (
  select candidate.* from ingestion.transformation_runs candidate
  where candidate.organization_id=b.organization_id and candidate.import_batch_id=b.id
  order by candidate.run_number desc limit 1
) r on true
where b.organization_id=%(organization_id)s::uuid
  and (%(batch_id)s::uuid is null or b.id=%(batch_id)s::uuid)
order by b.requested_at desc
limit 100
"""

INSERT_MAPPING_PROFILE_SQL = """
insert into ingestion.mapping_profiles (
  id,organization_id,name,source_kind,description,created_by
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(name)s,%(source_kind)s,
  %(description)s,%(created_by)s::uuid
)
on conflict (organization_id,name) do nothing
returning id
"""

LOAD_MAPPING_PROFILE_SQL = """
select id,organization_id,name,source_kind,description
from ingestion.mapping_profiles
where organization_id=%(organization_id)s::uuid and name=%(name)s
for share
"""

INSERT_MAPPING_VERSION_SQL = """
insert into ingestion.mapping_profile_versions (
  id,organization_id,mapping_profile_id,version_no,mapping_specification,
  mapping_checksum,transformation_version,created_by
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(mapping_profile_id)s::uuid,%(version_no)s,
  %(mapping_specification)s::jsonb,%(mapping_checksum)s,%(transformation_version)s,
  %(created_by)s::uuid
)
on conflict (mapping_profile_id,version_no) do nothing
returning id
"""

INSERT_MAPPING_VERSION_SHEET_SQL = """
insert into ingestion.mapping_version_sheets (
  id,organization_id,mapping_version_id,sheet_classification,sheet_ordinal,
  source_sheet_matcher,target_entity_type,mapping_specification,
  business_key_columns,formula_policy
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(mapping_version_id)s::uuid,
  %(sheet_classification)s,%(sheet_ordinal)s,%(source_sheet_matcher)s::jsonb,
  %(target_entity_type)s,%(mapping_specification)s::jsonb,
  %(business_key_columns)s::text[],'reject'
)
"""

INSERT_BATCH_SQL = """
insert into ingestion.import_batches (
  id,organization_id,project_id,source_kind,idempotency_key,status,requester_user_id
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(project_id)s::uuid,%(source_kind)s,
  %(idempotency_key)s,'uploaded',%(requester_user_id)s::uuid
)
returning id
"""

INSERT_SOURCE_FILE_SQL = """
insert into ingestion.source_files (
  id,organization_id,project_id,import_batch_id,source_kind,original_filename,
  mime_type,byte_size,checksum_sha256,uploaded_by,verified_at
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(project_id)s::uuid,%(import_batch_id)s::uuid,
  %(source_kind)s,%(original_filename)s,%(mime_type)s,%(byte_size)s,
  %(checksum_sha256)s,%(uploaded_by)s::uuid,now()
)
on conflict (organization_id,source_kind,checksum_sha256) do nothing
returning id
"""

INSERT_SOURCE_SHEET_SQL = """
insert into ingestion.source_sheets (
  id,organization_id,project_id,source_file_id,sheet_name,sheet_ordinal,
  row_count,column_count,is_hidden,profile
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(project_id)s::uuid,%(source_file_id)s::uuid,
  %(sheet_name)s,%(sheet_ordinal)s,%(row_count)s,%(column_count)s,%(is_hidden)s,
  %(profile)s::jsonb
)
"""

INSERT_SOURCE_COLUMN_SQL = """
insert into ingestion.source_columns (
  id,organization_id,project_id,source_sheet_id,column_ordinal,column_letter,
  raw_header,normalized_header,inferred_type,null_count,formula_count,profile
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(project_id)s::uuid,%(source_sheet_id)s::uuid,
  %(column_ordinal)s,%(column_letter)s,%(raw_header)s,%(normalized_header)s,
  %(inferred_type)s,%(null_count)s,%(formula_count)s,%(profile)s::jsonb
)
"""

INSERT_TRANSFORMATION_RUN_SQL = """
insert into ingestion.transformation_runs (
  id,organization_id,project_id,import_batch_id,mapping_version_id,run_number,status,
  input_profile_checksum,normalized_preview_checksum,validation_checksum,
  transformation_version,row_count,started_at,created_by
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(project_id)s::uuid,%(import_batch_id)s::uuid,
  %(mapping_version_id)s::uuid,%(run_number)s,'running',%(input_profile_checksum)s,
  %(normalized_preview_checksum)s,%(validation_checksum)s,%(transformation_version)s,
  %(row_count)s,now(),%(created_by)s::uuid
)
"""

INSERT_STAGED_ROW_SQL = """
insert into ingestion.staged_rows (
  id,organization_id,project_id,import_batch_id,source_file_id,source_sheet_id,
  transformation_run_id,mapping_version_sheet_id,row_number,target_entity_type,raw_record,normalized_record,
  business_key,business_key_checksum,raw_row_checksum,normalized_row_checksum,status
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(project_id)s::uuid,%(import_batch_id)s::uuid,
  %(source_file_id)s::uuid,%(source_sheet_id)s::uuid,%(transformation_run_id)s::uuid,
  %(mapping_version_sheet_id)s::uuid,%(row_number)s,%(target_entity_type)s,
  %(raw_record)s::jsonb,%(normalized_record)s::jsonb,
  %(business_key)s::jsonb,%(business_key_checksum)s,%(raw_row_checksum)s,
  %(normalized_row_checksum)s,%(status)s
)
"""

INSERT_STAGED_CELL_SQL = """
insert into ingestion.staged_cells (
  id,organization_id,project_id,staged_row_id,source_column_id,cell_reference,
  target_field,raw_value,normalized_value,formula_text
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(project_id)s::uuid,%(staged_row_id)s::uuid,
  %(source_column_id)s::uuid,%(cell_reference)s,%(target_field)s,
  %(raw_value)s::jsonb,%(normalized_value)s::jsonb,%(formula_text)s
)
"""

INSERT_VALIDATION_ISSUE_SQL = """
insert into ingestion.validation_issues (
  id,organization_id,project_id,import_batch_id,transformation_run_id,
  staged_row_id,staged_cell_id,severity,issue_code,message,issue_fingerprint,details
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(project_id)s::uuid,%(import_batch_id)s::uuid,
  %(transformation_run_id)s::uuid,%(staged_row_id)s::uuid,%(staged_cell_id)s::uuid,
  %(severity)s,%(issue_code)s,%(message)s,%(issue_fingerprint)s,%(details)s::jsonb
)
"""

APPROVAL_CONTEXT_SQL = """
select b.id as import_batch_id,b.organization_id,b.project_id,b.status,
       b.requester_user_id,b.mapping_version_id,b.input_profile_checksum,
       b.normalized_preview_checksum,b.validation_checksum,
       f.id as source_file_id,f.checksum_sha256 as source_checksum,
       v.mapping_profile_id,v.version_no,v.mapping_checksum,
       r.id as transformation_run_id,r.status as run_status,
       r.input_profile_checksum as run_input_profile_checksum,
       exists (
         select 1 from ingestion.validation_issues i
         where i.organization_id=b.organization_id and i.import_batch_id=b.id
           and i.transformation_run_id=r.id and i.severity='error'
       ) as has_errors
from ingestion.import_batches b
join ingestion.source_files f
  on f.organization_id=b.organization_id and f.import_batch_id=b.id
join ingestion.mapping_profile_versions v
  on v.organization_id=b.organization_id and v.id=b.mapping_version_id
join lateral (
  select candidate.*
  from ingestion.transformation_runs candidate
  where candidate.organization_id=b.organization_id
    and candidate.import_batch_id=b.id
    and candidate.mapping_version_id=b.mapping_version_id
  order by candidate.run_number desc
  limit 1
) r on true
where b.organization_id=%(organization_id)s::uuid
  and b.id=%(batch_id)s::uuid
order by r.run_number desc
for update of b
"""

APPROVER_ROLE_SQL = """
select role
from construction.organization_members
where organization_id=%(organization_id)s::uuid
  and user_id=%(user_id)s::uuid and is_active
"""

INSERT_APPROVAL_SQL = """
insert into ingestion.approval_records (
  id,organization_id,project_id,import_batch_id,mapping_version_id,
  transformation_run_id,requester_user_id,approver_user_id,decision,
  validation_checksum,normalized_preview_checksum,decision_reason
) values (
  %(id)s::uuid,%(organization_id)s::uuid,%(project_id)s::uuid,%(import_batch_id)s::uuid,
  %(mapping_version_id)s::uuid,%(transformation_run_id)s::uuid,
  %(requester_user_id)s::uuid,%(approver_user_id)s::uuid,%(decision)s,
  %(validation_checksum)s,%(normalized_preview_checksum)s,%(decision_reason)s
)
returning id
"""

PUBLICATION_CONTEXT_SQL = """
select a.id as approval_id,a.organization_id,a.import_batch_id,a.requester_user_id,
       a.approver_user_id,a.decision,a.validation_checksum,
       a.normalized_preview_checksum,b.input_profile_checksum,b.status as batch_status,
       a.transformation_run_id,
       f.id as source_file_id,f.checksum_sha256 as source_checksum,
       v.mapping_profile_id,v.id as mapping_version_id,v.version_no,v.mapping_checksum,
       r.input_profile_checksum as run_input_profile_checksum,
       r.validation_checksum as run_validation_checksum,
       r.normalized_preview_checksum as run_preview_checksum,
       m.role as approver_role
from ingestion.approval_records a
join ingestion.import_batches b
  on b.organization_id=a.organization_id and b.id=a.import_batch_id
join ingestion.source_files f
  on f.organization_id=b.organization_id and f.import_batch_id=b.id
join ingestion.mapping_profile_versions v
  on v.organization_id=b.organization_id and v.id=b.mapping_version_id
join ingestion.transformation_runs r
  on r.organization_id=a.organization_id and r.id=a.transformation_run_id
join construction.organization_members m
  on m.organization_id=a.organization_id and m.user_id=a.approver_user_id and m.is_active
where a.organization_id=%(organization_id)s::uuid and a.import_batch_id=%(batch_id)s::uuid
  and a.decision='approved'
"""


def _json(value: Any) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), default=str)


def _uuid5(parent: str, value: str) -> str:
    return str(uuid.uuid5(uuid.UUID(parent), value))


def _mapping_specification(mapping: OrganizationMapping) -> dict[str, Any]:
    payload = {"sheets": [sheet.serializable() for sheet in mapping.sheets]}
    if mapping.catalog_version is not None:
        payload["catalog_version"] = mapping.catalog_version
    return payload


def _mapping_sheet_id(mapping_version_id: str, ordinal: int, sheet: SheetMapping) -> str:
    return _uuid5(
        mapping_version_id,
        f"sheet|{ordinal}|{sheet.sheet_name.casefold()}|{sheet.entity_type}",
    )


def _governed_business_key(sheet: SheetMapping) -> tuple[str, ...]:
    expected = PUBLICATION_BUSINESS_KEYS.get(sheet.entity_type)
    if expected is None:
        raise ValueError(f"Unsupported governed publication entity: {sheet.entity_type}")
    supplied = tuple(sheet.business_key)
    normalized = supplied
    if expected[0] in {"organization_id", "project_id"} and expected[0] not in supplied:
        normalized = (expected[0], *supplied)
    if normalized != expected:
        raise ValueError(
            f"Business key for {sheet.entity_type} must be {', '.join(expected)}"
        )
    return expected


def _expected_input_profile_checksum(context: dict[str, Any]) -> str:
    return checksum({
        "source_checksum": context["source_checksum"],
        "mapping_profile_id": str(context["mapping_profile_id"]),
        "mapping_version_id": str(context["mapping_version_id"]),
        "mapping_version_no": int(context["version_no"]),
        "profile_checksum": context["mapping_checksum"],
    })


def _decode_mapping(row: dict[str, Any]) -> MappingRecord:
    specification = row["mapping_specification"]
    if isinstance(specification, str):
        specification = json.loads(specification)
    sheets = []
    for item in specification.get("sheets", []):
        fields = {
            name: FieldSpec(**definition) for name, definition in item["fields"].items()
        }
        sheets.append(SheetMapping(
            entity_type=item["entity_type"],
            sheet_name=item["sheet_name"],
            columns=dict(item["columns"]),
            fields=fields,
            business_key=tuple(item["business_key"]),
            sheet_aliases=tuple(item.get("sheet_aliases", ())),
            column_aliases={
                name: tuple(values) for name, values in item.get("column_aliases", {}).items()
            },
        ))
    mapping = OrganizationMapping(
        organization_id=str(row["organization_id"]),
        mapping_profile_id=str(row["mapping_profile_id"]),
        mapping_version_id=str(row["mapping_version_id"]),
        version_no=int(row["version_no"]),
        sheets=tuple(sheets),
        catalog_version=specification.get("catalog_version"),
    )
    if mapping.mapping_checksum != row["mapping_checksum"]:
        raise ValueError("Persisted mapping checksum does not match its immutable definition")
    if row["transformation_version"] != TRANSFORMATION_VERSION:
        raise ValueError("Mapping transformation version is not supported by this runtime")
    return MappingRecord(mapping, row["name"], row["source_kind"], row.get("description"))


class GovernedIngestionAdapter:
    """Persist each governed lifecycle phase in a distinct transaction."""

    def __init__(self, connection_factory):
        self.connection_factory = connection_factory

    def _transaction(self):
        return closing(self.connection_factory())

    def load_mapping(
        self, organization_id: str, mapping_profile_id: str, version_no: int
    ) -> MappingRecord:
        params = {
            "organization_id": str(uuid.UUID(organization_id)),
            "mapping_profile_id": str(uuid.UUID(mapping_profile_id)),
            "version_no": int(version_no),
        }
        if params["version_no"] < 1:
            raise ValueError("Mapping versions start at 1")
        with self._transaction() as connection:
            with connection:
                row = connection.execute_native(LOAD_MAPPING_SQL, params).fetchone()
        if not row:
            raise LookupError("Active mapping profile version was not found in this organization")
        return _decode_mapping(dict(row))

    def list_mappings(
        self, organization_id: str, *, entity_type: str | None = None
    ) -> list[MappingRecord]:
        organization = str(uuid.UUID(organization_id))
        with self._transaction() as connection:
            with connection:
                rows = connection.execute_native(
                    LIST_MAPPINGS_SQL, {"organization_id": organization}
                ).fetchall()
        records = [_decode_mapping(dict(row)) for row in rows]
        if entity_type:
            records = [
                record for record in records
                if any(sheet.entity_type == entity_type for sheet in record.mapping.sheets)
            ]
        return records

    def list_imports(
        self, organization_id: str, *, batch_id: str | None = None
    ) -> list[dict[str, Any]]:
        params = {
            "organization_id": str(uuid.UUID(organization_id)),
            "batch_id": str(uuid.UUID(batch_id)) if batch_id else None,
        }
        with self._transaction() as connection:
            with connection:
                rows = connection.execute_native(LIST_IMPORTS_SQL, params).fetchall()
        result = []
        for raw in rows:
            row = dict(raw)
            for field_name in ("rows", "issues"):
                if isinstance(row.get(field_name), str):
                    row[field_name] = json.loads(row[field_name])
            result.append(row)
        return result

    def register_mapping(self, record: MappingRecord, *, created_by: str) -> None:
        """Create one immutable mapping profile/version transactionally.

        The current DB migration retains a scalar compatibility label on the
        profile.  Multi-entity authority lives only in the immutable version
        specification; the label is never used to authorize staged targets.
        """
        mapping = record.mapping
        creator = str(uuid.UUID(created_by))
        if not record.name.strip() or record.source_kind != "excel":
            raise ValueError("A named Excel mapping profile is required")
        for sheet in mapping.sheets:
            _governed_business_key(sheet)
        profile_params = {
            "id": mapping.mapping_profile_id,
            "organization_id": mapping.organization_id,
            "name": record.name,
            "source_kind": record.source_kind,
            "description": record.description,
            "created_by": creator,
        }
        version_params = {
            "id": mapping.mapping_version_id,
            "organization_id": mapping.organization_id,
            "mapping_profile_id": mapping.mapping_profile_id,
            "version_no": mapping.version_no,
            "mapping_specification": _json(_mapping_specification(mapping)),
            "mapping_checksum": mapping.mapping_checksum,
            "transformation_version": TRANSFORMATION_VERSION,
            "created_by": creator,
        }
        with self._transaction() as connection:
            with connection:
                profile = connection.execute_native(INSERT_MAPPING_PROFILE_SQL, profile_params).fetchone()
                if not profile:
                    existing = connection.execute_native(
                        LOAD_MAPPING_PROFILE_SQL,
                        {
                            "organization_id": mapping.organization_id,
                            "name": record.name,
                        },
                    ).fetchone()
                    if not existing or str(existing["id"]) != mapping.mapping_profile_id:
                        raise ValueError(
                            "Mapping profile name already belongs to a different immutable profile"
                        )
                    if existing["source_kind"] != record.source_kind:
                        raise ValueError("Mapping profile source kind is immutable")
                version = connection.execute_native(INSERT_MAPPING_VERSION_SQL, version_params).fetchone()
                if not version:
                    raise ValueError(
                        "Mapping profile version already exists; immutable definitions cannot be replaced"
                    )
                for ordinal, sheet in enumerate(mapping.sheets):
                    connection.execute_native(
                        INSERT_MAPPING_VERSION_SHEET_SQL,
                        {
                            "id": _mapping_sheet_id(mapping.mapping_version_id, ordinal, sheet),
                            "organization_id": mapping.organization_id,
                            "mapping_version_id": mapping.mapping_version_id,
                            "sheet_classification": sheet.sheet_name,
                            "sheet_ordinal": ordinal,
                            "source_sheet_matcher": _json({
                                "sheet_name": sheet.sheet_name,
                                "aliases": list(sheet.sheet_aliases),
                            }),
                            "target_entity_type": sheet.entity_type,
                            "mapping_specification": _json(sheet.serializable()),
                            "business_key_columns": list(_governed_business_key(sheet)),
                        },
                    )

    def register_source(
        self,
        metadata: SourceFileMetadata,
        *,
        import_batch_id: str,
        project_id: str | None,
        mime_type: str | None,
        idempotency_key: str,
    ) -> None:
        batch_id = str(uuid.UUID(import_batch_id))
        project = str(uuid.UUID(project_id)) if project_id else None
        if not idempotency_key.strip():
            raise ValueError("Import idempotency key is required")
        batch = {
            "id": batch_id,
            "organization_id": metadata.organization_id,
            "project_id": project,
            "source_kind": metadata.source_kind,
            "idempotency_key": idempotency_key,
            "requester_user_id": metadata.uploaded_by,
        }
        source = {
            "id": metadata.id,
            "organization_id": metadata.organization_id,
            "project_id": project,
            "import_batch_id": batch_id,
            "source_kind": metadata.source_kind,
            "original_filename": metadata.filename,
            "mime_type": mime_type,
            "byte_size": metadata.byte_size,
            "checksum_sha256": metadata.checksum_sha256,
            "uploaded_by": metadata.uploaded_by,
        }
        with self._transaction() as connection:
            with connection:
                inserted_batch = connection.execute_native(INSERT_BATCH_SQL, batch).fetchone()
                if not inserted_batch:
                    raise ValueError("Import idempotency key is already registered")
                inserted = connection.execute_native(INSERT_SOURCE_FILE_SQL, source).fetchone()
                if not inserted:
                    raise DuplicateSourceFile(
                        "This organization has already registered the same source file"
                    )

    def persist_source_profile(
        self,
        *,
        preview: ValidationPreview,
        project_id: str | None,
        sheets: list[dict[str, Any]],
        columns: list[dict[str, Any]],
    ) -> None:
        project = str(uuid.UUID(project_id)) if project_id else None
        sheet_ids = set()
        for sheet in sheets:
            if (
                sheet.get("organization_id") != preview.organization_id
                or sheet.get("project_id") != project
                or sheet.get("source_file_id") != preview.source_file_id
            ):
                raise ValueError("Source sheet scope does not match the validated preview")
            sheet_ids.add(sheet.get("id"))
        for column in columns:
            if (
                column.get("organization_id") != preview.organization_id
                or column.get("project_id") != project
                or column.get("source_sheet_id") not in sheet_ids
            ):
                raise ValueError("Source column scope does not match its source sheet")
        params = {"organization_id": preview.organization_id, "batch_id": preview.import_batch_id}
        with self._transaction() as connection:
            with connection:
                changed = connection.execute_native(
                    """update ingestion.import_batches set status='profiling',status_changed_at=now(),updated_at=now()
                       where organization_id=%(organization_id)s::uuid and id=%(batch_id)s::uuid
                         and status='uploaded' returning id""",
                    params,
                ).fetchone()
                if not changed:
                    raise ValueError("Import batch is not in the uploaded state")
                for sheet in sheets:
                    connection.execute_native(INSERT_SOURCE_SHEET_SQL, sheet)
                for column in columns:
                    connection.execute_native(INSERT_SOURCE_COLUMN_SQL, column)
                mapped = connection.execute_native(
                    """update ingestion.import_batches set status='mapping',status_changed_at=now(),updated_at=now()
                       where organization_id=%(organization_id)s::uuid and id=%(batch_id)s::uuid
                         and status='profiling' returning id""",
                    params,
                ).fetchone()
                if not mapped:
                    raise RuntimeError("Import batch could not enter the mapping state")

    def persist_validation(
        self,
        *,
        preview: ValidationPreview,
        project_id: str | None,
        transformation_run_id: str,
        actor_id: str,
        staged_rows: list[dict[str, Any]],
        staged_cells: list[dict[str, Any]],
        issues: list[dict[str, Any]],
    ) -> None:
        project = str(uuid.UUID(project_id)) if project_id else None
        run_id = str(uuid.UUID(transformation_run_id))
        actor = str(uuid.UUID(actor_id))
        for row in staged_rows:
            if (
                row.get("organization_id") != preview.organization_id
                or row.get("project_id") != project
                or row.get("import_batch_id") != preview.import_batch_id
                or row.get("source_file_id") != preview.source_file_id
                or row.get("transformation_run_id") != run_id
                or not row.get("mapping_version_sheet_id")
            ):
                raise ValueError("Staged row scope does not match the validated transformation")
        row_ids = {row["id"] for row in staged_rows}
        for cell in staged_cells:
            if (
                cell.get("organization_id") != preview.organization_id
                or cell.get("project_id") != project
                or cell.get("staged_row_id") not in row_ids
            ):
                raise ValueError("Staged cell scope does not match its staged row")
        for issue in issues:
            if (
                issue.get("organization_id") != preview.organization_id
                or issue.get("project_id") != project
                or issue.get("import_batch_id") != preview.import_batch_id
                or issue.get("transformation_run_id") != run_id
                or (issue.get("staged_row_id") and issue["staged_row_id"] not in row_ids)
            ):
                raise ValueError("Validation issue scope does not match the transformation")
        batch = {"organization_id": preview.organization_id, "batch_id": preview.import_batch_id}
        with self._transaction() as connection:
            with connection:
                run_number_row = connection.execute_native(
                    """select coalesce(max(run_number),0)+1 as run_number
                       from ingestion.transformation_runs
                       where organization_id=%(organization_id)s::uuid
                         and import_batch_id=%(batch_id)s::uuid""",
                    batch,
                ).fetchone()
                run_number = int(run_number_row["run_number"])
                changed = connection.execute_native(
                    """update ingestion.import_batches set status='validating',mapping_version_id=%(mapping_version_id)s::uuid,
                           input_profile_checksum=%(input_profile_checksum)s,
                           normalized_preview_checksum=%(normalized_preview_checksum)s,
                           validation_checksum=%(validation_checksum)s,status_changed_at=now(),updated_at=now()
                       where organization_id=%(organization_id)s::uuid and id=%(batch_id)s::uuid
                         and status in ('mapping','validation_failed','review_ready') returning id""",
                    {**batch, "mapping_version_id": preview.mapping_version_id,
                     "input_profile_checksum": preview.input_profile_checksum,
                     "normalized_preview_checksum": preview.normalized_preview_checksum,
                     "validation_checksum": preview.validation_checksum},
                ).fetchone()
                if not changed:
                    raise ValueError("Import batch is not in the mapping state")
                connection.execute_native(INSERT_TRANSFORMATION_RUN_SQL, {
                    "id": run_id, "organization_id": preview.organization_id,
                    "project_id": project, "import_batch_id": preview.import_batch_id,
                    "mapping_version_id": preview.mapping_version_id, "run_number": run_number,
                    "input_profile_checksum": preview.input_profile_checksum,
                    "normalized_preview_checksum": preview.normalized_preview_checksum,
                    "validation_checksum": preview.validation_checksum,
                    "transformation_version": TRANSFORMATION_VERSION,
                    "row_count": len(staged_rows), "created_by": actor,
                })
                for row in staged_rows:
                    connection.execute_native(INSERT_STAGED_ROW_SQL, row)
                for cell in staged_cells:
                    connection.execute_native(INSERT_STAGED_CELL_SQL, cell)
                for issue in issues:
                    connection.execute_native(INSERT_VALIDATION_ISSUE_SQL, issue)
                completed = connection.execute_native(
                    """update ingestion.transformation_runs set status='succeeded',completed_at=now()
                       where organization_id=%(organization_id)s::uuid and id=%(run_id)s::uuid
                         and status='running' returning id""",
                    {"organization_id": preview.organization_id, "run_id": run_id},
                ).fetchone()
                if not completed:
                    raise RuntimeError("Transformation run could not be completed")
                finalized = connection.execute_native(
                    """update ingestion.import_batches set status=%(status)s,status_changed_at=now(),updated_at=now()
                       where organization_id=%(organization_id)s::uuid and id=%(batch_id)s::uuid
                         and status='validating' returning id""",
                    {**batch, "status": "review_ready" if preview.valid else "validation_failed"},
                ).fetchone()
                if not finalized:
                    raise RuntimeError("Import batch could not leave the validating state")

    def submit_for_approval(self, organization_id: str, batch_id: str) -> None:
        params = {"organization_id": str(uuid.UUID(organization_id)), "batch_id": str(uuid.UUID(batch_id))}
        with self._transaction() as connection:
            with connection:
                row = connection.execute_native(
                    """update ingestion.import_batches set status='awaiting_approval',status_changed_at=now(),updated_at=now()
                       where organization_id=%(organization_id)s::uuid and id=%(batch_id)s::uuid
                         and status='review_ready' returning id""",
                    params,
                ).fetchone()
                if not row:
                    raise ValueError("Only a valid review-ready batch can be submitted")

    def decide(
        self,
        organization_id: str,
        batch_id: str,
        *,
        actor_id: str,
        decision: Literal["approved", "rejected"],
        reason: str | None = None,
    ) -> ApprovalBinding:
        organization = str(uuid.UUID(organization_id))
        batch = str(uuid.UUID(batch_id))
        actor = str(uuid.UUID(actor_id))
        with self._transaction() as connection:
            with connection:
                rows = connection.execute_native(
                    APPROVAL_CONTEXT_SQL, {"organization_id": organization, "batch_id": batch}
                ).fetchall()
                if len(rows) != 1:
                    raise ValueError("Approval requires exactly one persisted source and current run")
                context = dict(rows[0])
                role_row = connection.execute_native(
                    APPROVER_ROLE_SQL, {"organization_id": organization, "user_id": actor}
                ).fetchone()
                role = role_row and role_row["role"]
                if role not in {"owner", "admin"}:
                    raise PermissionError("Only an active organization owner or admin may decide")
                if context["status"] != "awaiting_approval" or context["run_status"] != "succeeded":
                    raise ValueError("Batch is not awaiting approval for a completed transformation")
                if decision == "approved" and context["has_errors"]:
                    raise ValueError("A batch with validation errors cannot be approved")
                expected_input_checksum = _expected_input_profile_checksum(context)
                if (
                    context["input_profile_checksum"] != expected_input_checksum
                    or context["run_input_profile_checksum"] != expected_input_checksum
                ):
                    raise ValueError(
                        "Approval source, mapping profile/version, or input checksum binding is invalid"
                    )
                approval_id = str(uuid.uuid4())
                binding = ApprovalBinding(
                    id=approval_id, organization_id=organization, batch_id=batch,
                    source_file_id=str(context["source_file_id"]),
                    mapping_profile_id=str(context["mapping_profile_id"]),
                    mapping_version_id=str(context["mapping_version_id"]),
                    mapping_version_no=int(context["version_no"]),
                    source_checksum=context["source_checksum"],
                    profile_checksum=context["mapping_checksum"],
                    input_profile_checksum=context["input_profile_checksum"],
                    validation_checksum=context["validation_checksum"],
                    normalized_preview_checksum=context["normalized_preview_checksum"],
                    requester_id=str(context["requester_user_id"]), approver_id=actor,
                    approver_role=role, decision=decision,
                    transformation_run_id=str(context["transformation_run_id"]),
                )
                connection.execute_native(INSERT_APPROVAL_SQL, {
                    "id": approval_id, "organization_id": organization,
                    "project_id": context["project_id"], "import_batch_id": batch,
                    "mapping_version_id": context["mapping_version_id"],
                    "transformation_run_id": context["transformation_run_id"],
                    "requester_user_id": context["requester_user_id"],
                    "approver_user_id": actor, "decision": decision,
                    "validation_checksum": context["validation_checksum"],
                    "normalized_preview_checksum": context["normalized_preview_checksum"],
                    "decision_reason": reason,
                }).fetchone()
        return binding

    def publication_request(self, organization_id: str, batch_id: str) -> PublicationRequest:
        params = {"organization_id": str(uuid.UUID(organization_id)), "batch_id": str(uuid.UUID(batch_id))}
        with self._transaction() as connection:
            with connection:
                rows = connection.execute_native(PUBLICATION_CONTEXT_SQL, params).fetchall()
        if len(rows) != 1:
            raise ValueError("Batch does not have one exact approved publication binding")
        row = dict(rows[0])
        expected_input_checksum = _expected_input_profile_checksum(row)
        if (
            row["input_profile_checksum"] != expected_input_checksum
            or row["run_input_profile_checksum"] != expected_input_checksum
            or row["validation_checksum"] != row["run_validation_checksum"]
            or row["normalized_preview_checksum"] != row["run_preview_checksum"]
            or row["batch_status"] not in {"approved", "publish_failed", "published"}
        ):
            raise ValueError("Approved publication binding no longer matches its exact evidence")
        binding = ApprovalBinding(
            id=str(row["approval_id"]), organization_id=str(row["organization_id"]),
            batch_id=str(row["import_batch_id"]), source_file_id=str(row["source_file_id"]),
            mapping_profile_id=str(row["mapping_profile_id"]),
            mapping_version_id=str(row["mapping_version_id"]),
            mapping_version_no=int(row["version_no"]), source_checksum=row["source_checksum"],
            profile_checksum=row["mapping_checksum"],
            input_profile_checksum=row["input_profile_checksum"],
            validation_checksum=row["validation_checksum"],
            normalized_preview_checksum=row["normalized_preview_checksum"],
            requester_id=str(row["requester_user_id"]), approver_id=str(row["approver_user_id"]),
            approver_role=row["approver_role"], decision=row["decision"],
            transformation_run_id=str(row["transformation_run_id"]),
        )
        return PublicationRequest.from_binding(binding)


class GovernedIngestionService:
    """Validate a local temporary XLSX and persist its complete review trail."""

    def __init__(self, adapter: GovernedIngestionAdapter):
        self.adapter = adapter

    def prepare(
        self,
        path: Path,
        *,
        organization_id: str,
        allowed_project_ids: tuple[str, ...],
        actor_id: str,
        mapping_profile_id: str,
        mapping_version_no: int,
        original_filename: str,
        mime_type: str | None,
        project_id: str | None = None,
    ) -> PreparedImport:
        organization = str(uuid.UUID(organization_id))
        actor = str(uuid.UUID(actor_id))
        allowed_projects = tuple(str(uuid.UUID(value)) for value in allowed_project_ids)
        project = str(uuid.UUID(project_id)) if project_id else None
        if project is not None and project not in allowed_projects:
            raise PermissionError("Project scope is not authorized for this request")
        record = self.adapter.load_mapping(organization, mapping_profile_id, mapping_version_no)
        from .catalog import validate_compiled_mapping
        validate_compiled_mapping(record.mapping)
        if record.source_kind != "excel":
            raise ValueError("Mapping profile is not an Excel source profile")
        batch_id, source_id, run_id = (str(uuid.uuid4()) for _ in range(3))
        source_sha = file_checksum(path)
        metadata = SourceFileMetadata(
            id=source_id, organization_id=organization, source_kind="excel",
            filename=Path(original_filename).name, checksum_sha256=source_sha,
            byte_size=path.stat().st_size, uploaded_by=actor,
        )
        preview = validate_workbook(
            path, record.mapping,
            ReferenceCatalog(frozenset({organization}), {value: organization for value in allowed_projects}),
            source_file_id=source_id, import_batch_id=batch_id,
            project_id=project,
        )
        idempotency_key = checksum({
            "organization_id": organization, "source_kind": "excel",
            "source_checksum": source_sha, "mapping_profile_id": mapping_profile_id,
            "mapping_version_id": record.mapping.mapping_version_id,
            "mapping_version_no": mapping_version_no,
        })
        self.adapter.register_source(
            metadata, import_batch_id=batch_id, project_id=project,
            mime_type=mime_type, idempotency_key=idempotency_key,
        )
        sheets, columns, rows, cells, issues = _persistence_records(
            path, preview, record.mapping, project, run_id
        )
        self.adapter.persist_source_profile(
            preview=preview, project_id=project, sheets=sheets, columns=columns
        )
        self.adapter.persist_validation(
            preview=preview, project_id=project, transformation_run_id=run_id,
            actor_id=actor, staged_rows=rows, staged_cells=cells, issues=issues,
        )
        return PreparedImport(preview, run_id)


def _persistence_records(
    path: Path,
    preview: ValidationPreview,
    mapping: OrganizationMapping,
    project_id: str | None,
    run_id: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Derive deterministic sheet/column/cell IDs without evaluating formulas."""
    accepted = {
        (row.lineage["sheet_name"], row.lineage["row_number"], row.entity_type): row
        for row in preview.rows
    }
    issue_by_location = {
        (issue.sheet_name, issue.row_number, issue.field): issue for issue in preview.issues
    }
    sheets_out: list[dict[str, Any]] = []
    columns_out: list[dict[str, Any]] = []
    rows_out: list[dict[str, Any]] = []
    cells_out: list[dict[str, Any]] = []
    row_ids: dict[tuple[str, int, str], str] = {}
    cell_ids: dict[tuple[str, int, str], str] = {}
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        mappings: dict[str, tuple[int, SheetMapping]] = {}
        for mapping_ordinal, item in enumerate(mapping.sheets):
            matches = matching_sheet_names(book.sheetnames, item)
            if len(matches) == 1:
                mappings[matches[0]] = (mapping_ordinal, item)
        for ordinal, sheet_name in enumerate(book.sheetnames):
            sheet = book[sheet_name]
            sheet_id = _uuid5(preview.source_file_id, f"sheet|{ordinal}|{sheet_name}")
            all_rows = list(sheet.iter_rows())
            header_cells = all_rows[0] if all_rows else ()
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header_cells]
            data_rows = [row for row in all_rows[1:] if any(cell.value is not None for cell in row)]
            sheets_out.append({
                "id": sheet_id, "organization_id": preview.organization_id,
                "project_id": project_id, "source_file_id": preview.source_file_id,
                "sheet_name": sheet_name, "sheet_ordinal": ordinal,
                "row_count": len(data_rows), "column_count": len(headers),
                "is_hidden": sheet.sheet_state != "visible",
                "profile": _json({"classification": preview.sheet_classifications.get(sheet_name, "unmapped")}),
            })
            column_ids: dict[int, str] = {}
            for column_index, header in enumerate(headers, start=1):
                column_id = _uuid5(sheet_id, f"column|{column_index - 1}|{header.casefold()}")
                column_ids[column_index] = column_id
                values = [row[column_index - 1] for row in data_rows if column_index <= len(row)]
                nonempty = [cell for cell in values if cell.value is not None]
                inferred = sorted({type(cell.value).__name__ for cell in nonempty})
                columns_out.append({
                    "id": column_id, "organization_id": preview.organization_id,
                    "project_id": project_id, "source_sheet_id": sheet_id,
                    "column_ordinal": column_index - 1,
                    "column_letter": get_column_letter(column_index), "raw_header": header or None,
                    "normalized_header": _normalize_header(header) or None,
                    "inferred_type": inferred[0] if len(inferred) == 1 else ("mixed" if inferred else None),
                    "null_count": len(values) - len(nonempty),
                    "formula_count": sum(cell.data_type == "f" for cell in values),
                    "profile": _json({}),
                })
            resolved_mapping = mappings.get(sheet_name)
            if resolved_mapping is None:
                continue
            mapping_ordinal, sheet_mapping = resolved_mapping
            target_by_header = {
                _normalize_header(source): target
                for target in sheet_mapping.columns
                for source in sheet_mapping.source_labels(target)
            }
            for cells in data_rows:
                row_number = cells[0].row
                key = (sheet_name, row_number, sheet_mapping.entity_type)
                normalized = accepted.get(key)
                raw_record = {
                    (headers[index - 1] or get_column_letter(index)): _cell_json(cell.value)
                    for index, cell in enumerate(cells, start=1)
                    if cell.value is not None
                }
                normalized_record = dict(normalized.values) if normalized else {}
                if normalized:
                    business_key = {name: normalized.values.get(name) for name in sheet_mapping.business_key}
                    business_key.setdefault("organization_id", preview.organization_id)
                else:
                    business_key = {"__invalid_row__": f"{sheet_name}!{row_number}"}
                staged_row_id = _uuid5(run_id, f"row|{sheet_id}|{row_number}|{sheet_mapping.entity_type}")
                row_ids[key] = staged_row_id
                rows_out.append({
                    "id": staged_row_id, "organization_id": preview.organization_id,
                    "project_id": project_id, "import_batch_id": preview.import_batch_id,
                    "source_file_id": preview.source_file_id, "source_sheet_id": sheet_id,
                    "transformation_run_id": run_id, "row_number": row_number,
                    "mapping_version_sheet_id": _mapping_sheet_id(
                        mapping.mapping_version_id,
                        mapping_ordinal,
                        sheet_mapping,
                    ),
                    "target_entity_type": sheet_mapping.entity_type,
                    "raw_record": _json(raw_record), "normalized_record": _json(normalized_record),
                    "business_key": _json(business_key), "business_key_checksum": checksum(business_key),
                    "raw_row_checksum": checksum(raw_record),
                    "normalized_row_checksum": checksum(normalized_record),
                    "status": "valid" if normalized else "invalid",
                })
                for column_index, cell in enumerate(cells, start=1):
                    if column_index not in column_ids:
                        continue
                    header = headers[column_index - 1]
                    target = target_by_header.get(_normalize_header(header))
                    cell_id = _uuid5(staged_row_id, f"cell|{column_ids[column_index]}")
                    if target:
                        cell_ids[(sheet_name, row_number, target)] = cell_id
                    cells_out.append({
                        "id": cell_id, "organization_id": preview.organization_id,
                        "project_id": project_id, "staged_row_id": staged_row_id,
                        "source_column_id": column_ids[column_index],
                        "cell_reference": f"{quote_sheetname(sheet_name)}!{cell.coordinate}",
                        "target_field": target,
                        "raw_value": _json(_cell_json(cell.value)),
                        "normalized_value": _json(
                            normalized_record.get(target) if target and normalized else None
                        ),
                        "formula_text": str(cell.value) if cell.data_type == "f" else None,
                    })
    finally:
        book.close()
    issues_out = []
    for issue in preview.issues:
        location = (issue.sheet_name, issue.row_number)
        row_id = next((value for (sheet, row, _), value in row_ids.items() if (sheet, row) == location), None)
        cell_id = cell_ids.get((issue.sheet_name, issue.row_number, issue.field))
        fingerprint = checksum({"transformation_run_id": run_id, "issue": issue.__dict__})
        issues_out.append({
            "id": _uuid5(run_id, f"issue|{fingerprint}"),
            "organization_id": preview.organization_id, "project_id": project_id,
            "import_batch_id": preview.import_batch_id, "transformation_run_id": run_id,
            "staged_row_id": row_id, "staged_cell_id": cell_id, "severity": "error",
            "issue_code": issue.code, "message": issue.message,
            "issue_fingerprint": fingerprint, "details": _json(issue.__dict__),
        })
    return sheets_out, columns_out, rows_out, cells_out, issues_out


def _cell_json(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return str(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)
