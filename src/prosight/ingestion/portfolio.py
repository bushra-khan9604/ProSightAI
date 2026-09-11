"""Canonical multi-project manpower and invoice workbook parsing."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from typing import Any, Callable

from openpyxl import Workbook, load_workbook

from ..portfolio_contract import INVOICE_DB_COLUMNS, MANPOWER_DB_COLUMNS


MANPOWER_HEADERS = [
    "Sr.no", "EMP Code", "Name", "Designation", "Department",
    "Mobilized Project", "Current Project", "Category", "Current Location",
    "Allocation", "Status", "Leave Balance", "Remarks",
]
SCHEDULE_HEADERS = [
    "Activity ID", "Activity Name", "Start", "Finish", "Original Duration",
]
INVOICE_HEADERS = [
    "S/N", "Prof. Date", "Draft/ Prof. INV no.", "Job No", "SAP PO No#",
    "PO Line#", "Contract", "COO", "ICV %", "ICV @ 5% (USD)",
    "Invoice Value Excl. VAT (USD)", "Invoice Value Excl. VAT (AED)",
    "Invoice Value Excl. Tax & ICV (USD)", "Invoice Value Excl. Tax & ICV (AED)",
    "Tax Invoice/ Inv Ref", "P.S Job Officer", "Client Job Officer",
    "Work Description", "AVL/ NON-AVL", "Year of Work Done", "Asset",
    "ICV Claim Year", "Last Disc. Month", "MM/YY", "INVOICE STATUS",
    "Submission Date/ Status Change Date", "Current Date",
    "Days Pending Approval/Action", "Reason for Rejection", "Aging (Ref)",
    "INVOICE Approval STATUS", "Invoice PAYMENT Status (Ref)", "Client Ref",
    "Payment Terms (Days)", "Days Pending for Remittance", "Current Date (Payment)",
    "Expected Remittance Date", "Outstanding Status", "Project", "Levels",
    "Tax Invoice Taken Outstanding", "Aging of Approval", "Risk Profile", "Remarks",
]
INVOICE_DATE_HEADERS = {
    "Prof. Date", "Last Disc. Month", "Submission Date/ Status Change Date",
    "Current Date", "Current Date (Payment)", "Expected Remittance Date",
}
INVOICE_NUMBER_HEADERS = {
    "ICV %", "ICV @ 5% (USD)", "Invoice Value Excl. VAT (USD)",
    "Invoice Value Excl. VAT (AED)", "Invoice Value Excl. Tax & ICV (USD)",
    "Invoice Value Excl. Tax & ICV (AED)", "Days Pending Approval/Action",
    "Payment Terms (Days)", "Days Pending for Remittance",
}
INVOICE_INTEGER_HEADERS = {
    "S/N", "Year of Work Done", "Days Pending Approval/Action",
    "Payment Terms (Days)", "Days Pending for Remittance",
}
INVOICE_IDENTIFIER_HEADERS = {
    "Draft/ Prof. INV no.", "Job No", "SAP PO No#", "PO Line#", "Contract",
    "COO", "Tax Invoice/ Inv Ref", "Asset", "ICV Claim Year", "Client Ref",
}

# Keep this mapping explicit: these names are both the normalized application
# contract and the typed PostgreSQL column names used by portfolio imports.
INVOICE_FIELD_MAP = {
    "S/N": "serial_number",
    "Prof. Date": "proforma_date",
    "Draft/ Prof. INV no.": "draft_invoice_number",
    "Job No": "job_number",
    "SAP PO No#": "sap_po_number",
    "PO Line#": "po_line_number",
    "Contract": "contract_number",
    "COO": "coo_number",
    "ICV %": "icv_percent",
    "ICV @ 5% (USD)": "icv_five_percent_usd",
    "Invoice Value Excl. VAT (USD)": "invoice_value_usd",
    "Invoice Value Excl. VAT (AED)": "invoice_value_aed",
    "Invoice Value Excl. Tax & ICV (USD)": "invoice_value_excl_tax_icv_usd",
    "Invoice Value Excl. Tax & ICV (AED)": "invoice_value_excl_tax_icv_aed",
    "Tax Invoice/ Inv Ref": "tax_invoice_reference",
    "P.S Job Officer": "ps_job_officer",
    "Client Job Officer": "client_job_officer",
    "Work Description": "work_description",
    "AVL/ NON-AVL": "avl_status",
    "Year of Work Done": "work_year",
    "Asset": "asset",
    "ICV Claim Year": "icv_claim_year",
    "Last Disc. Month": "last_discussion_month",
    "MM/YY": "mm_yy",
    "INVOICE STATUS": "status",
    "Submission Date/ Status Change Date": "submission_date",
    "Current Date": "approval_current_date",
    "Days Pending Approval/Action": "days_pending_approval",
    "Reason for Rejection": "reason_for_rejection",
    "Aging (Ref)": "aging_reference",
    "INVOICE Approval STATUS": "approval_status",
    "Invoice PAYMENT Status (Ref)": "payment_status",
    "Client Ref": "client_reference",
    "Payment Terms (Days)": "payment_terms_days",
    "Days Pending for Remittance": "days_pending_remittance",
    "Current Date (Payment)": "payment_current_date",
    "Expected Remittance Date": "expected_remittance_date",
    "Outstanding Status": "outstanding_status",
    "Project": "project",
    "Levels": "levels",
    "Tax Invoice Taken Outstanding": "tax_invoice_taken_outstanding",
    "Aging of Approval": "aging_of_approval",
    "Risk Profile": "risk_profile",
    "Remarks": "remarks",
}
MANPOWER_FIELD_MAP = {
    "Sr.no": "serial_number", "EMP Code": "emp_code", "Name": "name",
    "Designation": "designation", "Department": "department",
    "Mobilized Project": "mobilized_project", "Current Project": "current_project",
    "Category": "category", "Current Location": "current_location",
    "Allocation": "allocation", "Status": "status",
    "Leave Balance": "leave_balance", "Remarks": "remarks",
}



def _value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def _decimal(value: Any, label: str, row: int, required: bool = False) -> float | None:
    if value in (None, ""):
        if required:
            raise ValueError(f"{label} row {row}: value is required")
        return None
    try:
        return float(Decimal(str(value).replace(",", "").replace("%", "").strip()))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"{label} row {row}: invalid number {value!r}") from error


def _integer(value: Any, label: str, row: int) -> int | None:
    number = _decimal(value, label, row)
    if number is None:
        return None
    if not number.is_integer():
        raise ValueError(f"{label} row {row}: value must be a whole number")
    return int(number)


def _identifier(value: Any, number_format: str | None = None) -> str | None:
    """Keep spreadsheet identifiers as text without adding a trailing .0."""
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        if number_format and re.fullmatch(r"0+", number_format):
            return f"{value:0{len(number_format)}d}"
        return str(value)
    if isinstance(value, float) and value.is_integer():
        if number_format and re.fullmatch(r"0+", number_format):
            return f"{int(value):0{len(number_format)}d}"
        return str(int(value))
    return str(value).strip() or None


def _identifier_cell(sheet: Any, header: str, row: int) -> str | None:
    columns = {
        str(cell.value or "").strip(): cell.column
        for cell in sheet[1]
    }
    cell = sheet.cell(row=row, column=columns[header])
    return _identifier(cell.value, cell.number_format)


def _date(value: Any, label: str, row: int) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y", "%d-%b-%y", "%d %B %Y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"{label} row {row}: invalid date {value!r}")


def _headers(sheet: Any, expected: list[str]) -> None:
    actual = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(max_row=1))]
    missing = [header for header in expected if header not in actual]
    if missing:
        raise ValueError(f"{sheet.title} is missing columns: {', '.join(missing)}")
    unexpected = [header for header in actual if header not in expected]
    if unexpected:
        raise ValueError(f"{sheet.title} has unsupported columns: {', '.join(unexpected)}")
    if actual != expected:
        raise ValueError(f"{sheet.title} columns must match the downloaded template order")


def _records(sheet: Any) -> list[tuple[int, dict[str, Any]]]:
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    return [
        (number, {header: _value(value) for header, value in zip(headers, values)})
        for number, values in enumerate(rows, start=2)
        if any(value not in (None, "") for value in values)
    ]


def parse_portfolio_workbook(
    path: Path,
    resolve_project: Callable[[str], str | None],
    dataset: str = "combined",
    project_code: str | None = None,
) -> dict[str, Any]:
    """Validate the canonical workbook and return transaction-ready records."""
    if dataset not in {"combined", "manpower", "invoices", "schedule"}:
        raise ValueError("Unsupported portfolio import dataset")
    workbook = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    formula_workbook = (
        load_workbook(path, read_only=False, data_only=False, keep_links=False)
        if dataset in {"combined", "invoices"} else None
    )
    try:
        required_sheets = {
            "combined": {"Manpower", "Projects Invoices"},
            "manpower": {"Manpower"},
            "invoices": {"Projects Invoices"},
            "schedule": {"Project Schedule"},
        }[dataset]
        missing_sheets = required_sheets - set(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(f"Workbook is missing sheets: {', '.join(sorted(missing_sheets))}")
        manpower_sheet = workbook["Manpower"] if dataset in {"combined", "manpower"} else None
        invoice_sheet = workbook["Projects Invoices"] if dataset in {"combined", "invoices"} else None
        invoice_formula_sheet = (
            formula_workbook["Projects Invoices"]
            if formula_workbook and "Projects Invoices" in formula_workbook.sheetnames else None
        )
        schedule_sheet = workbook["Project Schedule"] if dataset == "schedule" else None
        if manpower_sheet:
            _headers(manpower_sheet, MANPOWER_HEADERS)
        if invoice_sheet:
            _headers(invoice_sheet, INVOICE_HEADERS)
        if schedule_sheet:
            _headers(schedule_sheet, SCHEDULE_HEADERS)
            if not project_code or not resolve_project(project_code):
                raise ValueError("Project Schedule requires an authorized project")

        manpower, employee_keys = [], set()
        for row_number, source in _records(manpower_sheet) if manpower_sheet else []:
            emp_code = _identifier_cell(manpower_sheet, "EMP Code", row_number) or ""
            if not emp_code:
                raise ValueError(f"Manpower row {row_number}: EMP Code is required")
            if emp_code.casefold() in employee_keys:
                raise ValueError(f"Manpower row {row_number}: duplicate EMP Code {emp_code}")
            employee_keys.add(emp_code.casefold())
            current = resolve_project(str(source.get("Current Project") or ""))
            mobilized_raw = str(source.get("Mobilized Project") or "")
            mobilized = resolve_project(mobilized_raw) if mobilized_raw else None
            if not current:
                raise ValueError(f"Manpower row {row_number}: unknown Current Project")
            if mobilized_raw and not mobilized:
                raise ValueError(f"Manpower row {row_number}: unknown Mobilized Project")
            record = {MANPOWER_FIELD_MAP[key]: source.get(key) for key in MANPOWER_HEADERS}
            record.update(
                serial_number=_integer(source.get("Sr.no"), "Sr.no", row_number),
                emp_code=emp_code, current_project_code=current,
                mobilized_project_code=mobilized,
                allocation=str(source.get("Allocation") or "").strip() or None,
                leave_balance=_decimal(source.get("Leave Balance"), "Leave Balance", row_number),
            )
            manpower.append(record)

        invoices, invoice_keys = [], set()
        for row_number, source in _records(invoice_sheet) if invoice_sheet else []:
            job_number = _identifier_cell(invoice_sheet, "Job No", row_number) or ""
            draft_number = _identifier_cell(
                invoice_sheet, "Draft/ Prof. INV no.", row_number
            ) or ""
            if not job_number or not draft_number:
                raise ValueError(
                    f"Projects Invoices row {row_number}: Job No and Draft/ Prof. INV no. are required"
                )
            invoice_key = (job_number.casefold(), draft_number.casefold())
            if invoice_key in invoice_keys:
                raise ValueError(
                    f"Projects Invoices row {row_number}: duplicate invoice key "
                    f"{job_number} / {draft_number}"
                )
            invoice_keys.add(invoice_key)
            job_project = resolve_project(job_number)
            named_project = resolve_project(str(source.get("Project") or ""))
            if job_project and named_project and job_project != named_project:
                raise ValueError(f"Projects Invoices row {row_number}: Job No and Project conflict")
            project_code = job_project or named_project
            if not project_code:
                raise ValueError(f"Projects Invoices row {row_number}: project could not be resolved")
            record = {INVOICE_FIELD_MAP[key]: source.get(key) for key in INVOICE_HEADERS}
            for header in INVOICE_DATE_HEADERS:
                record[INVOICE_FIELD_MAP[header]] = _date(
                    source.get(header), header, row_number
                )
            for header in INVOICE_NUMBER_HEADERS:
                record[INVOICE_FIELD_MAP[header]] = _decimal(
                    source.get(header), header, row_number
                )
            for header in INVOICE_INTEGER_HEADERS:
                record[INVOICE_FIELD_MAP[header]] = _integer(
                    source.get(header), header, row_number
                )
            for header in INVOICE_IDENTIFIER_HEADERS:
                record[INVOICE_FIELD_MAP[header]] = _identifier_cell(
                    invoice_sheet, header, row_number
                )
            if source.get("Invoice Value Excl. VAT (USD)") in (None, ""):
                if invoice_formula_sheet:
                    header_cells = {
                        str(cell.value or "").strip(): cell.column
                        for cell in invoice_formula_sheet[1]
                    }
                    formula_value = invoice_formula_sheet.cell(
                        row=row_number,
                        column=header_cells["Invoice Value Excl. VAT (USD)"],
                    ).value
                    if isinstance(formula_value, str) and formula_value.startswith("="):
                        raise ValueError(
                            f"Projects Invoices row {row_number}: Invoice Value Excl. VAT "
                            "(USD) formula has no cached value; recalculate and save the workbook"
                        )
            record.update(
                job_number=job_number, draft_invoice_number=draft_number,
                project_code=project_code,
                invoice_value_usd=_decimal(
                    source.get("Invoice Value Excl. VAT (USD)"),
                    "Invoice Value Excl. VAT (USD)", row_number, True,
                ),
                invoice_value_aed=_decimal(
                    source.get("Invoice Value Excl. VAT (AED)"),
                    "Invoice Value Excl. VAT (AED)", row_number,
                ),
            )
            invoices.append(record)

        schedule, activity_keys = [], set()
        for row_number, source in _records(schedule_sheet) if schedule_sheet else []:
            activity_id = str(source.get("Activity ID") or "").strip()
            activity_name = str(source.get("Activity Name") or "").strip()
            if not activity_id or not activity_name:
                raise ValueError(
                    f"Project Schedule row {row_number}: Activity ID and Activity Name are required"
                )
            key = activity_id.casefold()
            if key in activity_keys:
                raise ValueError(
                    f"Project Schedule row {row_number}: duplicate Activity ID {activity_id}"
                )
            activity_keys.add(key)
            start = _date(source.get("Start"), "Start", row_number)
            finish = _date(source.get("Finish"), "Finish", row_number)
            if not start or not finish:
                raise ValueError(f"Project Schedule row {row_number}: Start and Finish are required")
            if start > finish:
                raise ValueError(f"Project Schedule row {row_number}: Start must not be after Finish")
            duration = _decimal(
                source.get("Original Duration"), "Original Duration", row_number, True
            )
            if duration is None or duration < 0 or not duration.is_integer():
                raise ValueError(
                    f"Project Schedule row {row_number}: Original Duration must be a non-negative integer"
                )
            schedule.append({
                "project_code": project_code, "activity_id": activity_id,
                "activity_name": activity_name, "start": start, "finish": finish,
                "original_duration": int(duration),
            })

        return {
            "manpower": manpower,
            "invoices": invoices,
            "schedule": schedule,
            "counts": {"manpower": len(manpower), "invoices": len(invoices),
                       "schedule": len(schedule)},
        }
    finally:
        workbook.close()
        if formula_workbook:
            formula_workbook.close()


def generate_invoice_pivot(invoices: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], dict[str, float]] = {}
    for invoice in invoices:
        key = (str(invoice.get("levels") or ""), str(invoice.get("status") or ""))
        projects = grouped.setdefault(key, {})
        code = invoice["project_code"]
        projects[code] = round(projects.get(code, 0) + float(invoice["invoice_value_usd"]), 2)
    return [
        {
            "levels": levels, "status": status, "projects": projects,
            "grand_total": round(sum(projects.values()), 2),
        }
        for (levels, status), projects in sorted(grouped.items())
    ]


def create_portfolio_template(path: Path, dataset: str = "combined") -> Path:
    """Create a canonical combined or dataset-specific XLSX template."""
    if dataset not in {"combined", "manpower", "invoices", "schedule"}:
        raise ValueError("Unsupported portfolio template dataset")
    workbook = Workbook()
    first = workbook.active
    if dataset == "schedule":
        first.title = "Project Schedule"
        first.append(SCHEDULE_HEADERS)
        first.append(["A1001", "Receive PO", "30-Mar-26", "06-Apr-26", 6])
        first.append(["A1006", "Receive sketch from client", "06-Apr-26", "14-Apr-26", 7])
    elif dataset == "invoices":
        first.title = "Projects Invoices"
        first.append(INVOICE_HEADERS)
    else:
        first.title = "Manpower"
        first.append(MANPOWER_HEADERS)
        if dataset == "combined":
            invoices = workbook.create_sheet("Projects Invoices")
            invoices.append(INVOICE_HEADERS)
    workbook.save(path)
    return path
