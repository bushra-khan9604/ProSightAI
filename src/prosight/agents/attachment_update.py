"""Attachment and Update Agent: local analysis and durable, authorized execution.

Only server-owned tools produce writes. No document content or generated SQL is
sent to a model during validation. Existing change_requests persist operations.
"""
from __future__ import annotations

import copy
import hashlib
import json
import re
import shutil
import uuid
import zipfile
from contextlib import closing
from pathlib import Path

from ..ingestion.excel import preview_workbook, PROJECT_COLUMNS
from ..ingestion.pdf import extract_pdf_chunks, detect_reporting_date
from ..ingestion.portfolio import parse_portfolio_workbook
from ..ingestion.workforce import (
    WORKFORCE_DATASETS,
    WORKFORCE_TABLES,
    parse_workforce_workbook,
)
from ..project_locks import project_lock

ACTION = 'attachment_update'
ROLES = {'admin', 'project_manager', 'planning_engineer'}
TABLES = ('seed_migrations', 'projects', 'documents', 'change_requests',
          'manpower_assignments', 'project_invoices', 'project_schedule_activities',
          'employees', 'employee_training', 'attendance_records', 'payroll_records',
          'direct_allocations', 'subcontract_crews', 'deployment_forecasts',
          'portfolio_imports', 'portfolio_import_jobs', 'ingestion_jobs', 'notifications', 'audit_events')


def encoded(value):
    return json.dumps(value, sort_keys=True, default=str, separators=(',', ':'))


def digest(value):
    return hashlib.sha256(encoded(value).encode()).hexdigest()


class AttachmentUpdateAgent:
    name = 'Attachment and Update Agent'

    def __init__(self, repository, ingestion, storage_root):
        self.repository, self.ingestion = repository, ingestion
        self.storage_root = Path(storage_root).resolve()

    def _lock(self, db):
        if self.repository.backend == 'postgres':
            db.execute('LOCK TABLE seed_migrations IN EXCLUSIVE MODE')
            db.execute('LOCK TABLE '+','.join(TABLES[1:])+' IN EXCLUSIVE MODE')
        else:
            db.execute('BEGIN IMMEDIATE')

    def _bound(self, db):
        """Run existing validated repository tools inside our single transaction."""
        class Borrowed:
            def execute(self, *args): return db.execute(*args)
            def executemany(self, *args): return db.executemany(*args)
            def __enter__(self): return self
            def __exit__(self, *args): return False
            def close(self): pass
        repo = copy.copy(self.repository)
        repo.connect = lambda: Borrowed()
        repo._ensure = lambda: None
        repo.ensure_schema = lambda: None
        repo._reject_deleted_project = lambda ignored, code: self.repository._reject_deleted_project(db, code)
        return repo

    @staticmethod
    def _identity(user):
        if not user or not user.get('id') or user.get('role') not in ROLES:
            raise PermissionError('A signed-in project member is required')

    def _read(self, db, operation_id, user):
        self._identity(user)
        row = db.execute('SELECT * FROM change_requests WHERE id=? AND action=?', (operation_id, ACTION)).fetchone()
        if not row: raise KeyError('Attachment operation not found')
        operation = dict(row)
        operation['payload'] = json.loads(operation['payload'])
        operation['preview'] = json.loads(operation['preview'])
        if user['role'] != 'admin' and operation['payload']['owner_id'] != user['id']:
            raise PermissionError('This attachment belongs to another member')
        return operation

    def _public(self, operation, user):
        payload = operation['payload']
        preview = copy.deepcopy(operation['preview'])
        if user['role'] != 'admin':
            def mask(value):
                if isinstance(value, dict):
                    restricted={'billing_rate','cost_rate','cost_value','basic_aed','allowance_aed',
                                'gross_monthly_aed','ot_rate_aed','ot_pay_aed','gross_aed',
                                'employer_burden_aed','total_cost_aed','payroll_cost_aed'}
                    return {k: ('Restricted' if k in restricted else mask(v)) for k,v in value.items()}
                if isinstance(value, list): return [mask(v) for v in value]
                return value
            preview = mask(preview)
        return dict(id=operation['id'], status=operation['status'], project_code=operation['project_code'],
                    filename=payload['filename'], instruction=payload['instruction'], tool=payload['kind'],
                    preview=preview, preview_token=payload['preview_token'],
                    can_confirm=operation['status']=='awaiting_confirmation' and payload['owner_id']==user['id'],
                    can_approve=operation['status']=='pending' and user['role']=='admin',
                    can_retry=operation['status'] in {'failed','indexing'} and payload['kind']=='pdf' and user['role']=='admin',
                    external_processing=payload['external_processing'], result=payload.get('result'),
                    error=payload.get('error'), created_at=operation['created_at'])

    def list(self, user):
        self._identity(user)
        with closing(self.repository.connect()) as db:
            rows = db.execute('SELECT * FROM change_requests WHERE action=? ORDER BY created_at DESC', (ACTION,)).fetchall()
        result=[]
        for raw in rows:
            row=dict(raw); row['payload']=json.loads(row['payload']); row['preview']=json.loads(row['preview'])
            if user['role']=='admin' or row['payload']['owner_id']==user['id']:
                result.append(self._public(row,user))
        return result[:100]

    def get(self, operation_id, user):
        with closing(self.repository.connect()) as db:
            return self._public(self._read(db,operation_id,user),user)

    def _state(self, db, code, kind):
        """Conservative snapshot: any project/dataset change requires a new preview."""
        result={}
        for table,where in [('projects','code'),('documents','project_code')]:
            result[table]=sorted([dict(r) for r in db.execute('SELECT * FROM '+table+' WHERE '+where+'=?',(code,)).fetchall()],key=encoded)
        if kind in WORKFORCE_DATASETS:
            result['projects']=sorted([dict(r) for r in db.execute('SELECT * FROM projects').fetchall()],key=encoded)
            for table in WORKFORCE_TABLES.values():
                name=table[0]
                result[name]=sorted([dict(r) for r in db.execute('SELECT * FROM '+name).fetchall()],key=encoded)
        elif kind not in {'pdf','new_project','project'}:
            # Global natural keys can otherwise overwrite a row owned by another project.
            for table in ('manpower_assignments','project_invoices','project_schedule_activities'):
                result[table]=sorted([dict(r) for r in db.execute('SELECT * FROM '+table).fetchall()],key=encoded)
        return result

    def pdf_upload_tool(self, path, code, filename, replacement_id, db):
        document={'id':'preview','project_code':code,'filename':filename,'approval_status':'pending'}
        chunks=extract_pdf_chunks(path,document)
        replacement=None
        if replacement_id:
            row=db.execute("SELECT * FROM documents WHERE id=? AND project_code=? AND kind='pdf' AND approval_status='approved'",(replacement_id,code)).fetchone()
            if not row: raise ValueError('Select an approved PDF from this project to replace')
            replacement={'id':row['id'],'filename':row['filename']}
        return {'reporting_date':detect_reporting_date(path),'replacement_id':replacement_id}, {
            'summary':'PDF evidence: admin approval is required before external embedding and publication.',
            'pages_with_text':len({c['metadata']['page_number'] for c in chunks}), 'chunks':len(chunks),
            'excerpt':chunks[0]['text'][:800], 'replaces':replacement,
            'warnings':['Scanned pages without extractable text are not indexed. Verify the source document.'],
        }

    def excel_analysis_tool(self, path, code, kind, db):
        from ..project_models import ProjectDraft
        if kind in {'new_project','project'}:
            parsed=preview_workbook(path,code)
            if parsed['mapping_required']: raise ValueError('Use the defined Projects workbook schema; automatic mapping is not applied')
            if len(parsed['projects'])!=1: raise ValueError('Upload one project per workbook for an explicit preview')
            project=parsed['projects'][0]
            project.update(ProjectDraft(**{k:project.get(k) for k in PROJECT_COLUMNS}).storage_record())
            if kind=='project' and project['code']!=code: raise ValueError('Workbook project does not match the selected project')
            if kind=='new_project': code=project['code']
            before=db.execute('SELECT payload FROM projects WHERE code=?',(code,)).fetchone()
            if kind=='new_project' and before: raise ValueError('That project already exists; select Update project instead')
            return {'projects':[project]}, {'summary':'Review the complete project record. Related sheets replace the corresponding project lists.',
                'rows':[{'operation':'update' if before else 'insert','key':code,'before':json.loads(before['payload']) if before else None,'after':project}],
                'warnings':parsed['warnings']}, code
        if kind in WORKFORCE_DATASETS:
            repo=self._bound(db)
            def resolve_employee(value):
                row=db.execute('SELECT employee_id FROM employees WHERE lower(employee_id)=lower(?)',(str(value).strip(),)).fetchone()
                return row['employee_id'] if row else None
            parsed=parse_workforce_workbook(path,repo.resolve_project_reference,resolve_employee,kind)
            rows=[];total_rows=0;preview_limit=200
            immutable_references={
                'training':('employee_id',),
                'attendance':('employee_id','project_id'),
                'payroll':('employee_id','project_id'),
            }
            for collection,(table,keys) in WORKFORCE_TABLES.items():
                existing={
                    tuple(str(row[key]) for key in keys):json.loads(row['data_json'])
                    for row in db.execute('SELECT '+','.join(keys)+',data_json FROM '+table).fetchall()
                }
                for item in parsed.get(collection,[]):
                    total_rows+=1
                    key=tuple(str(item[value]) for value in keys);before=existing.get(key)
                    if before and any(
                        before.get(field)!=item.get(field)
                        for field in immutable_references.get(collection,())
                    ):
                        raise ValueError(
                            f"{collection.replace('_',' ').title()} key {' / '.join(key)} "
                            "is already assigned to a different employee or project"
                        )
                    if len(rows)>=preview_limit:continue
                    rows.append({
                        'dataset':collection,
                        'key':' / '.join(key),
                        'operation':'update' if before else 'insert',
                        'before':before,
                        'after':item,
                    })
            if not total_rows:raise ValueError('No workforce or deployment rows were found')
            return parsed, {
                'summary':f'Validated and mapped {total_rows} workforce/deployment rows. Review the schema mappings, counts, and row sample before approval.',
                'rows':rows,'warnings':parsed.get('warnings',[]),
                'mappings':parsed.get('semantic_mappings',{}),'counts':parsed.get('counts',{}),
                'row_count':total_rows,'preview_limit':preview_limit,
            }, 'PORTFOLIO'
        parsed=parse_portfolio_workbook(path,self._bound(db).resolve_project_reference,dataset=kind,project_code=code)
        rows=[]
        for collection,table,keys in [('manpower','manpower_assignments',('emp_code',)),
                                     ('invoices','project_invoices',('job_number','draft_invoice_number')),
                                     ('schedule','project_schedule_activities',('project_code','activity_id'))]:
            for item in parsed.get(collection,[]):
                target=item.get('project_code') or item.get('current_project_code')
                if target!=code: raise ValueError('Every row must belong to the selected project; split multi-project workbooks')
                if item.get('mobilized_project_code') not in {None,'',code}:
                    raise ValueError('Cross-project mobilization requires a separate reviewed workflow')
                before=db.execute('SELECT * FROM '+table+' WHERE '+' AND '.join(k+'=?' for k in keys),tuple(item[k] for k in keys)).fetchone()
                if before:
                    existing=dict(before)
                    if (existing.get('project_code') or existing.get('current_project_code'))!=code:
                        raise ValueError('A row key is already assigned to another project')
                    old=json.loads(existing['data_json']) if 'data_json' in existing else existing
                else: old=None
                rows.append({'dataset':collection,'key':' / '.join(str(item[k]) for k in keys),'operation':'update' if before else 'insert','before':old,'after':item})
        if not rows: raise ValueError('No data rows were found')
        return parsed, {'summary':'Review every proposed insert and update. Unlisted rows will remain unchanged.',
                        'rows':rows,'warnings':parsed.get('warnings',[]),'mappings':parsed.get('semantic_mappings',{})}, code

    def prepare(self, source, filename, instruction, code, kind, user, external_processing=False, replacement_id=None):
        self._identity(user)
        if kind not in {'pdf','new_project','project','manpower','invoices','schedule','combined',*WORKFORCE_DATASETS}:
            raise ValueError('Select a supported attachment action')
        if len(instruction.strip())<5 or len(instruction)>4000:
            raise ValueError('Describe the intended upload or update in 5–4000 characters')
        if re.search(r'\b(delete|drop|truncate|execute sql)\b',instruction,re.I):
            raise ValueError('This agent handles uploads and reviewed updates, not deletion or arbitrary SQL')
        if kind=='new_project' and user['role'] not in {'admin','project_manager'}:
            raise PermissionError('Only Admin and Project Manager can submit new projects')
        if kind in {'employees_training','attendance_payroll'} and user['role'] not in {'admin','project_manager'}:
            raise PermissionError('Only Admin and Project Manager can submit employee or payroll workbooks')
        suffix=Path(filename).suffix.lower()
        if Path(filename).name!=filename or suffix!=('.pdf' if kind=='pdf' else '.xlsx'):
            raise ValueError('Use a PDF for evidence or an XLSX workbook for structured updates')
        if not 0<source.stat().st_size<=20*1024*1024: raise ValueError('Attachments must be between 1 byte and 20 MB')
        if kind=='pdf' and not source.read_bytes().startswith(b'%PDF-'): raise ValueError('Invalid PDF signature')
        if kind!='pdf':
            with zipfile.ZipFile(source) as archive:
                if len(archive.infolist())>2000 or sum(i.file_size for i in archive.infolist())>80*1024*1024:
                    raise ValueError('Workbook expanded size exceeds the safe limit')
            from openpyxl import load_workbook
            book=load_workbook(source,read_only=True,data_only=False,keep_links=False)
            try:
                limit=400000 if kind in WORKFORCE_DATASETS else 200000
                if sum((s.max_row or 0)*(s.max_column or 0) for s in book)>limit:
                    raise ValueError('Workbook exceeds the cell limit')
                if kind not in WORKFORCE_DATASETS:
                    for sheet in book:
                        for row in sheet:
                            if any(cell.data_type=='f' for cell in row):
                                raise ValueError('Export a values-only workbook; formulas are not accepted for database updates')
            finally: book.close()
        checksum=hashlib.sha256(source.read_bytes()).hexdigest()
        self.repository._ensure()
        with closing(self.repository.connect()) as db:
            with db:
                self._lock(db)
                if kind!='new_project' and kind not in WORKFORCE_DATASETS and not db.execute('SELECT code FROM projects WHERE code=?',(code,)).fetchone():
                    raise ValueError('Select an existing project')
                if kind=='pdf': prepared,preview=self.pdf_upload_tool(source,code,filename,replacement_id,db)
                else:
                    prepared,preview,code=self.excel_analysis_tool(source,code,kind,db)
                    for project in prepared.get('projects',[]):
                        project['sources']=[value.replace(source.name,filename) for value in project.get('sources',[])]
                self.repository._reject_deleted_project(db,code)
                signature=digest([checksum,code,kind,replacement_id,bool(external_processing)])
                for raw in db.execute('SELECT * FROM change_requests WHERE action=?',(ACTION,)).fetchall():
                    old=dict(raw);old['payload']=json.loads(old['payload']);old['preview']=json.loads(old['preview'])
                    if old['payload'].get('signature')==signature and old['payload']['owner_id']==user['id'] and old['status'] not in {'rejected','stale'}:
                        return self._public(old,user)
                operation_id=str(uuid.uuid4())
                root=self.storage_root/code
                if not root.resolve().is_relative_to(self.storage_root): raise ValueError('Invalid project code')
                root.mkdir(parents=True,exist_ok=True)
                path=root/(operation_id+suffix)
                shutil.copyfile(source,path)
                snapshot=digest(self._state(db,code,kind))
                payload=dict(owner_id=user['id'],filename=filename,instruction=instruction.strip(),kind=kind,
                             stored_path=str(path),checksum=checksum,signature=signature,prepared=prepared,
                             snapshot=snapshot,external_processing=bool(external_processing))
                payload['preview_token']=digest([preview,checksum,snapshot])
                try:
                    db.execute('INSERT INTO change_requests VALUES (?,?,?,?,?,?,?,?,?,?)',
                               (operation_id,ACTION,code,encoded(payload),encoded(preview),'awaiting_confirmation',user['role'],None,self.repository._now(),None))
                except Exception:
                    path.unlink(missing_ok=True);raise
        return self.get(operation_id,user)

    def _save(self,db,operation,status):
        operation['status']=status
        db.execute('UPDATE change_requests SET payload=?,status=? WHERE id=?',
                   (encoded(operation['payload']),status,operation['id']))

    def decide(self, operation_id, user, token, decision):
        self._identity(user)
        if decision not in {'confirm','approve','reject','retry'}: raise ValueError('Invalid attachment decision')
        schedule=None
        with closing(self.repository.connect()) as db:
            with db:
                self._lock(db)
                op=self._read(db,operation_id,user);p=op['payload'];status=op['status']
                if token!=p['preview_token']: raise ValueError('The preview has changed; reload it before continuing')
                if decision in {'approve','retry'} and user['role']!='admin': raise PermissionError('Admin approval is required')
                if decision=='confirm' and p['owner_id']!=user['id']: raise PermissionError('Only the uploader can confirm this preview')
                if decision=='reject':
                    if status not in {'awaiting_confirmation','pending'}: raise ValueError('This operation can no longer be rejected')
                    self._save(db,op,'rejected')
                elif status in {'completed','indexing'} and decision!='retry':
                    return self._public(op,user)
                elif decision=='retry':
                    if p['kind']!='pdf' or status not in {'failed','indexing'} or not p.get('result'):
                        raise ValueError('Only an approved PDF indexing operation can be retried')
                    p.pop('error',None)
                    self._save(db,op,'indexing');schedule=op
                else:
                    expected='awaiting_confirmation' if decision=='confirm' else 'pending'
                    if status!=expected: raise ValueError('This operation is not awaiting that decision')
                    path=Path(p['stored_path']).resolve()
                    if not path.is_relative_to(self.storage_root) or not path.is_file(): raise ValueError('Original attachment is unavailable')
                    if hashlib.sha256(path.read_bytes()).hexdigest()!=p['checksum'] or digest(self._state(db,op['project_code'],p['kind']))!=p['snapshot']:
                        self._save(db,op,'stale')
                        return self._public(op,user)
                    if decision=='confirm':
                        p['confirmed_by']=user['id'];p['confirmed_at']=self.repository._now()
                    if decision=='confirm' and user['role']!='admin':
                        self._save(db,op,'pending')
                        self.repository._insert_notification(db,'admin','approval_required',op['project_code'],operation_id,
                            'Attachment approval required',p['filename']+' is confirmed and ready for Admin review.')
                    else:
                        p['approved_by']=user['id'];p['approved_at']=self.repository._now()
                        db.execute('UPDATE change_requests SET decided_by=?,decided_at=? WHERE id=?',(user['id'],self.repository._now(),operation_id))
                        repo=self._bound(db)
                        if p['kind']=='pdf':
                            if not p['external_processing']: raise ValueError('External embedding permission is required; submit a permitted copy or keep analysis local')
                            if not self.ingestion: raise ValueError('PDF indexing is unavailable; configure the embedding service before approval')
                            doc=repo.create_document(op['project_code'],p['filename'],'pdf',p['checksum'],p['stored_path'])
                            date=p['prepared'].get('reporting_date') or self.repository._now()[:10]
                            db.execute("UPDATE documents SET approval_status='approved',status='approved',index_status='queued',reporting_date=?,effective_date=?,date_status=?,approved_by=?,approved_at=? WHERE id=?",
                                (date,date,'detected' if p['prepared'].get('reporting_date') else 'fallback',user['id'],self.repository._now(),doc['id']))
                            job=repo.create_job(doc['id'])
                            p['result']={'document_id':doc['id'],'job_id':job['id']}
                            self._save(db,op,'indexing');schedule=op
                        elif p['kind'] in {'project','new_project'}:
                            repo._apply_change(db,{'action':'excel_import','payload':p['prepared'],'project_code':op['project_code']})
                            p['result']={'project_code':op['project_code']};self._save(db,op,'completed')
                        else:
                            record=repo.create_portfolio_import(p['filename'],p['checksum'],p['stored_path'],user['role'],p['kind'],op['project_code'])
                            repo.apply_portfolio_import(record['id'],p['prepared'])
                            p['result']={'import_id':record['id']};self._save(db,op,'completed')
                        repo._insert_audit(db,user['role'],'attachment_applied','attachment',operation_id,None,
                                           {'project_code':op['project_code'],'kind':p['kind'],'actor_id':user['id']},self.repository._now())
                if op["status"] != "pending":
                    db.execute("UPDATE notifications SET read_at=? WHERE change_request_id=?",(self.repository._now(),operation_id))
        if schedule:
            try: self.ingestion.executor.submit(self._publish_pdf,schedule)
            except Exception: self._fail_pdf(operation_id,'Indexing could not start. Admin can retry.')
        return self.get(operation_id,user)

    def _fail_pdf(self,operation_id,message):
        with closing(self.repository.connect()) as db:
            with db:
                row=db.execute('SELECT payload FROM change_requests WHERE id=?',(operation_id,)).fetchone()
                if row:
                    payload=json.loads(row['payload']);payload['error']=message
                    db.execute("UPDATE change_requests SET status='failed',payload=? WHERE id=?",(encoded(payload),operation_id))

    def _publish_pdf(self,operation):
        p=operation['payload'];code=operation['project_code'];result=p['result']
        try:
            with project_lock(self.repository,code):
                document=self.repository.get_document(result['document_id'])
                if not document: return
                if document['index_status']!='ready':
                    self.ingestion._index_pdf(result['job_id'],document)
                document=self.repository.get_document(result['document_id'])
                if not document or document['index_status']!='ready':
                    self._fail_pdf(operation['id'],'Indexing failed. Admin can retry without applying the upload again.');return
                with closing(self.repository.connect()) as db:
                    with db:
                        old=p['prepared'].get('replacement_id')
                        if old:
                            db.execute("UPDATE documents SET status='superseded',approval_status='superseded' WHERE id=? AND project_code=?",(old,code))
                        db.execute("UPDATE change_requests SET status='completed' WHERE id=? AND status='indexing'",(operation['id'],))
        except Exception:
            self._fail_pdf(operation['id'],'Indexing could not finish. Admin can retry.')
