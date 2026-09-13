"""Export immutable structured draft revisions without model-generated numbers."""
from io import BytesIO
import json
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, PageBreak


def analysis_rows(result):
    """Flatten calculation output into readable scalar rows without truncating JSON cells."""
    rows=[]
    def visit(section,path,value):
        if isinstance(value,dict):
            for key,item in value.items():visit(section,path+[key.replace('_',' ')],item)
        elif isinstance(value,list):
            for index,item in enumerate(value,1):visit(section,path+[str(index)],item)
            if not value:rows.append({'section':section,'metric':' / '.join(path),'value':'No records'})
        else:
            text=value if value is not None else 'Unavailable'
            # Preserve unusually long evidence text over separate workbook cells.
            pieces=[text[i:i+12000] for i in range(0,len(text),12000)] if isinstance(text,str) and len(text)>12000 else [text]
            for piece in pieces:rows.append({'section':section,'metric':' / '.join(path),'value':piece})
    for section,value in result.items():visit(section.replace('_',' '),[],value)
    return rows


def render(version, format):
    stream=BytesIO()
    notice="Synthetic demonstration data — not an actual project record." if version.get("synthetic") else "Proposed project controls; authority depends on recorded approval status."
    title=f"{version['project_code']} · {version['status']} · {version['reporting_date']}"
    if format=="xlsx":
        wb=Workbook();readme=wb.active;readme.title="ReadMe"
        for pair in [("Project",version['project_code']),("Version",version['id']),("Status",version['status']),("Reporting date",version['reporting_date']),("Notice",notice)]:readme.append(pair)
        for name,rows in version['content'].items():
            if not rows: continue
            sheet=wb.create_sheet(name[:31]);headers=list(dict.fromkeys(k for r in rows for k in r));sheet.append(headers)
            for row in rows:
                values=[]
                for h in headers:
                    value=row.get(h)
                    if isinstance(value,(dict,list)):value=json.dumps(value)
                    if isinstance(value,str) and value.startswith(('=','+','-','@')):value="'"+value
                    values.append(value)
                sheet.append(values)
            sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
        for sheet in wb:
            sheet.oddFooter.center.text=notice
            for cell in sheet[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill('solid',fgColor='17324D')
            for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=24
            for row in sheet:
                for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
                sheet.row_dimensions[row[0].row].height=min(409,15*max(2,max((len(str(c.value or ''))+25)//26 for c in row)))
        wb.save(stream)
        return stream.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    styles=getSampleStyleSheet()
    story=[Paragraph(escape(title),styles['Title']),Paragraph(escape(notice),styles['Normal']),Spacer(1,12)]
    sections=("Requirements","Planning Documents","Schedule","Cost Baseline","Measurement Plan","Assignments","Procurement","Risks","Billing","Cash Flow","Proposed Recovery","Analysis Findings")
    preferred={
        'Schedule':['activity_id','name','forecast_start','forecast_finish','duration_wd','resource_id'],
        'Cost Baseline':['activity_id','quantity','unit','labor_budget_usd','equipment_budget_usd','budget_usd'],
        'Measurement Plan':['activity_id','baseline_quantity','unit','earning_method','budget_weight'],
        'Assignments':['activity_id','resource_id','start_date','finish_date','people'],
        'Procurement':['package_id','description','planned_order_date','required_on_site','owner','status'],
        'Risks':['risk_id','description','probability','impact_cost_usd','owner','mitigation'],
        'Analysis Findings':['section','metric','value']}
    styles['BodyText'].fontSize=8;styles['BodyText'].leading=10
    for name in sections:
        rows=version['content'].get(name,[])
        if not rows:continue
        story.extend([Paragraph(escape(name),styles['Heading2']),Spacer(1,6)])
        if name in preferred:
            headers=preferred[name]
            values=[[Paragraph(escape(h.replace('_',' ')),styles['BodyText']) for h in headers]]
            values.extend([Paragraph(escape(f'{row.get(h):,.2f}' if isinstance(row.get(h),float) else str(row.get(h) if row.get(h) is not None else 'Unavailable')),styles['BodyText']) for h in headers] for row in rows[:500])
            table=LongTable(values,colWidths=[(landscape(A4)[0]-72)/len(headers)]*len(headers),repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCE5EF')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#B7C5D3')),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
            story.extend([table,Spacer(1,12)])
        else:
            for row in rows[:500]:
                text="<br/>".join(f"<b>{escape(k.replace('_',' '))}:</b> {escape(str(v))}" for k,v in row.items() if v is not None)
                story.extend([Paragraph(text,styles['Normal']),Spacer(1,8)])
        if len(rows)>500:story.append(Paragraph("Further records are available in the XLSX export.",styles['Normal']))
    def page(canvas,doc):
        canvas.setFont('Helvetica',7);canvas.drawString(36,22,notice.replace('—','-'))
        canvas.drawRightString(landscape(A4)[0]-36,22,f"{version['id'][:8]} | Page {doc.page}")
    SimpleDocTemplate(stream,pagesize=landscape(A4),rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=40).build(story,onFirstPage=page,onLaterPages=page)
    return stream.getvalue(),"application/pdf"
