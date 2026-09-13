"""Bounded XLSX adapters for controls schema 1.0 and existing upload contracts."""
from __future__ import annotations

import copy
import hashlib
import json
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook

from ..ingestion.excel import preview_workbook, CANONICAL_SHEETS
from ..ingestion.portfolio import parse_portfolio_workbook
from .calculations import number, schedule, evm

SCHEMA = {t["sheet"]: t for t in json.loads(Path(__file__).with_name("workbook_schema.json").read_text())["tables"]}
INFO = {"ReadMe", "Dataset Notes"}
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100000


def json_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def inspect_archive(path):
    if Path(path).stat().st_size > MAX_BYTES:
        raise ValueError("File exceeds 20 MB")
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 150*1024*1024:
                raise ValueError("Decompressed workbook exceeds safe limits")
            if any(e.flag_bits & 1 or e.filename.endswith("vbaProject.bin") for e in entries):
                raise ValueError("Encrypted or macro-enabled files are unsupported")
    except BadZipFile as exc:
        raise ValueError("Invalid XLSX file") from exc


def parse(path, project_code, resolve_project):
    inspect_archive(path)
    tables, metadata, sources = {}, {}, []
    with_values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    with_formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        names = set(with_values.sheetnames)
        total = 0
        for sheet in with_values:
            formulas = with_formulas[sheet.title]
            for row, formula_row in zip(sheet.iter_rows(), formulas.iter_rows()):
                total += 1
                if total > MAX_ROWS or len(row) > 100:
                    raise ValueError("Workbook exceeds row or column limits")
                for cell, formula in zip(row, formula_row):
                    if formula.data_type == "f" and cell.value is None:
                        raise ValueError(f"{sheet.title}!{cell.coordinate}: formula needs a saved calculated value")
                    if cell.data_type == "e":
                        raise ValueError(f"{sheet.title}!{cell.coordinate}: Excel error")
                    if isinstance(cell.value, str) and len(cell.value) > 16000:
                        raise ValueError("Cell text exceeds 16,000 characters")
            if sheet.title in INFO:
                text = " ".join(str(c) for r in sheet.iter_rows(values_only=True) for c in r if c is not None)
                lowered = text.lower()
                if any(term in lowered for term in ("evaluation_only", "reference answer", "reference schedule", "evaluation reference", "reference-only", "evaluation-only")):
                    raise ValueError("Evaluation reference answers cannot be ingested")
                metadata["synthetic"] = metadata.get("synthetic", False) or "synthetic demonstration data" in lowered
                if sheet.title == "ReadMe":
                    pairs = {r[0]: r[1] for r in sheet.iter_rows(values_only=True) if len(r)>1}
                    if str(pairs.get('Classification','')).casefold()=='evaluation' or 'excluded from agent inputs' in lowered:
                        raise ValueError('Evaluation reference answers cannot be ingested')
                    code = pairs.get("Project")
                    if code and code not in {project_code, "PORTFOLIO", "Portfolio"}:
                        raise ValueError(f"Workbook belongs to {code}, not {project_code}")
                else:
                    rows = list(sheet.iter_rows(values_only=True))
                    if len(rows)>1:
                        notes = dict(zip(rows[0], rows[1]))
                        if notes.get("project_code") not in (None, project_code):
                            raise ValueError("Dataset Notes project mismatch")
                        metadata.update({k: notes[k] for k in ("reporting_date",) if notes.get(k)})
        if "Projects" in names:
            unknown = names - CANONICAL_SHEETS - INFO
            if unknown:
                raise ValueError(f"Unsupported sheets: {sorted(unknown)}")
            projects = preview_workbook(Path(path), project_code, max_rows=MAX_ROWS)["projects"]
            if not projects or any(p["code"] != project_code for p in projects):
                raise ValueError("Project Master must contain only the selected project")
            tables["Project Master"] = projects
        elif "Project Schedule" in names or "Projects Invoices" in names or "Manpower" in names:
            kind = "schedule" if "Project Schedule" in names else "invoices" if "Projects Invoices" in names else "manpower"
            expected = {"schedule": "Project Schedule", "invoices": "Projects Invoices", "manpower": "Manpower"}[kind]
            if names - INFO - {expected}:
                raise ValueError("Unsupported sheets in legacy workbook")
            result = parse_portfolio_workbook(Path(path), resolve_project, kind, project_code)
            rows = result[kind]
            if any(r.get("project_code", r.get("current_project_code")) != project_code for r in rows):
                raise ValueError("Workbook includes another project")
            tables[{"schedule": "Basic Schedule", "invoices": "Invoices", "manpower": "Current Manpower"}[kind]] = rows
        else:
            unknown = names - INFO - set(SCHEMA)
            if unknown:
                raise ValueError(f"Unsupported sheets: {sorted(unknown)}")
            if not names - INFO:
                raise ValueError("No supported data sheets")
            for name in sorted(names - INFO):
                table = SCHEMA[name]
                rows = with_values[name].iter_rows(values_only=True)
                headers = [str(v or "").strip() for v in next(rows, ())]
                expected = [f["name"] for f in table["fields"]]
                if headers != expected:
                    raise ValueError(f"{name}: expected schema 1.0 headers {expected}")
                records, seen = [], set()
                for index, values in enumerate(rows, 2):
                    if not any(v is not None for v in values):
                        continue
                    record = dict(zip(headers, map(json_value, values)))
                    for field in table["fields"]:
                        value = record[field["name"]]
                        if value is None:
                            if not field["nullable"] and field["name"] not in {"revised_finish", "actual_start", "actual_finish", "actual_order_date", "actual_delivery", "paid_date", "resolution_date"}:
                                raise ValueError(f"{name} row {index}: {field['name']} is required")
                        elif field["type"] == "number":
                            record[field["name"]] = None if name in {"Performance", "Overview"} and value == "n.a." else number(value)
                        elif field["type"] == "date":
                            date.fromisoformat(str(value))
                        elif field["type"] == "boolean" and not isinstance(value, bool):
                            raise ValueError(f"{name} row {index}: boolean required")
                    code = record.get("project_code", project_code)
                    if name == 'Source Snapshot' and record.get('code') != project_code:
                        raise ValueError('Source snapshot belongs to another project')
                    if name != "Demand" and code != project_code:
                        raise ValueError(f"{name} row {index}: wrong project")
                    if name == "Demand" and resolve_project(code) != code:
                        raise ValueError("Shared demand references an unauthorized project")
                    key = tuple(record.get(k, project_code if k == "project_code" else None) for k in table["primary_key"])
                    if key in seen:
                        raise ValueError(f"{name} row {index}: duplicate key {key}")
                    seen.add(key)
                    records.append(record)
                    sources.append(dict(sheet=name, row=index, key=list(key)))
                tables[name] = records
        return {"tables": tables, "metadata": metadata, "sources": sources,
                "checksum": hashlib.sha256(Path(path).read_bytes()).hexdigest(), "filename": Path(path).name}
    finally:
        with_values.close()
        with_formulas.close()


def merge_files(parsed_files, parent=None):
    """Replace supplied sections, inherit omitted sections, reject cross-file conflicts."""
    supplied = {}
    for file in parsed_files:
        for name, rows in file["tables"].items():
            if name in supplied and supplied[name] != rows:
                raise ValueError(f"Conflicting definitions for {name}; reconcile files before submission")
            supplied[name] = rows
    merged = {**copy.deepcopy(parent or {}), **copy.deepcopy(supplied)}
    # Legacy schedule updates propose forecast changes while retaining the baseline.
    # Supplying both formats still requires exact agreement during validation.
    if parent and 'Basic Schedule' in supplied and 'Schedule' not in supplied and parent.get('Schedule'):
        basic = {r['activity_id']: r for r in supplied['Basic Schedule']}
        if len(basic) != len(supplied['Basic Schedule']) or set(basic) != {r['activity_id'] for r in parent['Schedule']}:
            raise ValueError('Basic schedule changes must preserve activity identities; use Project Controls for structural edits')
        for activity in merged['Schedule']:
            proposed = basic[activity['activity_id']]
            if proposed['start'] > proposed['finish']:
                raise ValueError('Schedule start cannot follow finish')
            activity['forecast_start'] = proposed['start']
            activity['forecast_finish'] = proposed['finish']
    return merged


def validate(tables, project, reporting_date):
    """Cross-sheet invariants; incomplete sections become explicit readiness findings."""
    warnings = []
    date.fromisoformat(reporting_date)
    activities = tables.get("Schedule", [])
    activity_ids = {a["activity_id"] for a in activities}
    if len(activity_ids) != len(activities):
        raise ValueError('Duplicate schedule activity IDs')
    for activity in activities:
        duration=number(activity['duration_wd']);remaining=number(activity['remaining_duration_wd'])
        if duration<0 or remaining<0 or remaining>duration or int(duration)!=duration or int(remaining)!=remaining:
            raise ValueError('Activity durations must be nonnegative integer working days; remaining duration cannot exceed original duration')
        if activity.get('actual_finish') and (not activity.get('actual_start') or activity['actual_finish']<activity['actual_start'] or remaining):
            raise ValueError('Completed activities require coherent actual dates and zero remaining duration')
    milestones = tables.get("Milestones", [])
    boundary = {}
    for m in milestones:
        key = m["activity_id"]
        if key not in activity_ids:
            boundary.setdefault(key, {"activity_id": key, "name": m["name"], "duration_wd": 0})
    nodes = activities + list(boundary.values())
    keys = {a["activity_id"] for a in nodes}
    if activities:
        from collections import defaultdict, deque
        successors=defaultdict(list);degree={key:0 for key in keys}
        for relationship in tables.get('Relationships',[]):
            if relationship['type']!='FS':raise ValueError('Only finish-to-start relationships are supported')
            pred,succ=relationship['predecessor_id'],relationship['successor_id']
            if pred not in keys or succ not in keys:raise ValueError('Broken schedule relationship')
            successors[pred].append(succ);degree[succ]+=1
        ready=deque(k for k,d in degree.items() if d==0);visited=0
        while ready:
            key=ready.popleft();visited+=1
            for successor in successors[key]:
                degree[successor]-=1
                if degree[successor]==0:ready.append(successor)
        if visited!=len(keys):raise ValueError('Schedule contains a cycle')
    resources={r['resource_id'] for r in tables.get('Resources',[])}
    costs={r['cost_code'] for r in tables.get('Cost Baseline',[])}
    for name, rows in tables.items():
        if name in {"Basic Schedule", "Project Master", "Current Manpower", "Invoices", "Milestones", "Source Snapshot"}:
            continue
        for row in rows:
            if name in SCHEMA:
                for field in SCHEMA[name]['fields']:
                    value=row.get(field['name'])
                    if value is not None and field['type']=='number':number(value)
                    if value is not None and field['type']=='date':date.fromisoformat(str(value))
            if row.get('resource_id') and resources and name in {'Schedule','Assignments'} and row['resource_id'] not in resources:
                raise ValueError(f'{name}: unknown resource')
            if row.get('cost_code') and costs and row['cost_code'] not in costs:
                raise ValueError(f'{name}: unknown cost code')
            if row.get("activity_id") and name != "Schedule" and activity_ids and row["activity_id"] not in keys:
                raise ValueError(f"{name}: unknown activity {row['activity_id']}")
            for field in ("actual_start", "actual_finish", "actual_order_date", "actual_delivery"):
                if row.get(field) and row[field] > reporting_date:
                    raise ValueError(f"{name}: {field} is after the reporting cutoff")
    if tables.get("Basic Schedule") and activities:
        basic = {a["activity_id"]: a for a in tables["Basic Schedule"]}
        if set(basic) != activity_ids:
            raise ValueError("Basic and detailed schedules define different activities")
        for a in activities:
            b = basic[a["activity_id"]]
            if b["start"] != a["forecast_start"] or b["finish"] != a["forecast_finish"]:
                raise ValueError(f"Basic/detailed schedule dates conflict for {a['activity_id']}")
    if project["status"] == "future":
        if tables.get("Actual Costs") or tables.get("Measurements") or tables.get("Current Manpower") or any(a.get("actual_start") or a.get("actual_finish") for a in activities):
            raise ValueError("Future projects cannot contain execution actuals")
    if project["status"] == "completed" and tables.get("Current Manpower"):
        raise ValueError("Completed-project assignments cannot replace current manpower")
    for row in tables.get("Actual Costs", []):
        if row["date"] > reporting_date:
            raise ValueError("Actual cost is after reporting cutoff")
    for row in tables.get("Measurements", []):
        if row["period_end"] > reporting_date:
            raise ValueError("Measurement is after reporting cutoff")
    from .calculations import risk_exposure
    risk_exposure(tables.get('Risks',[]))
    for budget in tables.get('Cost Baseline',[]):
        calculated=sum(number(budget[k]) for k in ('labor_budget_usd','equipment_budget_usd','material_subcontract_usd'))
        if abs(calculated-number(budget['budget_usd']))>.05:
            raise ValueError('Cost baseline components do not reconcile to the activity budget')
    if tables.get('Budget Periods') and tables.get('Cost Baseline'):
        planned=sum(number(p['planned_value_usd']) for p in tables['Budget Periods'])
        budget=sum(number(b['budget_usd']) for b in tables['Cost Baseline'])
        if abs(planned-budget)>max(.05,budget*.000001):raise ValueError('Time-phased budgets do not reconcile to BAC')
    for bill in tables.get('Billing',[]):
        if abs(bill['gross_usd']-bill['retention_usd']-bill['net_due_usd'])>.05:raise ValueError('Billing gross, retention and net do not reconcile')
    for benchmark in tables.get('Benchmarks',[]):
        if benchmark['actual_labor_hours']<=0 or benchmark['quantity']<=0:raise ValueError('Benchmarks require positive quantities and hours')
        if abs(benchmark['quantity']/benchmark['actual_labor_hours']-benchmark['quantity_per_labor_hour'])>.00001:raise ValueError('Benchmark productivity does not reproduce source quantity/hours')
        if abs(benchmark['actual_cost_usd']/benchmark['quantity']-benchmark['unit_cost_usd'])>.01:raise ValueError('Benchmark unit cost does not reproduce source cost/quantity')
        transactions=[r for r in tables.get('Actual Costs',[]) if r['activity_id']==benchmark['activity_id']]
        if transactions and abs(sum(r['cost_usd'] for r in transactions)-benchmark['actual_cost_usd'])>.05:raise ValueError('Benchmark cost differs from actual transactions')
    tables["WBS"] = [{"wbs_id": key} for key in sorted({r["wbs_id"] for r in activities+tables.get("Scope Quantities", [])})]
    scope = tables.get("Scope Quantities", [])
    if scope and activities:
        for a in activities:
            matches = [s for s in scope if s["wbs_id"] == a["wbs_id"] and s["work_item"].casefold() == a["name"].casefold()]
            if len(matches) != 1 or matches[0]["unit"] != a["unit"] or abs(matches[0]["quantity"]-a["quantity"]) > .001:
                raise ValueError(f"Ambiguous or inconsistent scope mapping for {a['activity_id']}")
    readiness = {"schedule": bool(activities and tables.get("Calendars") and tables.get("Relationships")),
                 "evm": bool(tables.get("Cost Baseline") and tables.get("Budget Periods")) and (project["status"] == "future" or bool(tables.get("Measurements") and tables.get("Actual Costs"))),
                 "planning": bool(scope and tables.get("Resources") and tables.get("Commercial") and tables.get("Calendars")),
                 "cash_flow": bool(tables.get("Billing") and (tables.get("Actual Costs") or (project['status']=='future' and tables.get('Budget Periods')))),
                 "resources": bool((tables.get("Demand") and tables.get("Capacity")) or (tables.get('Assignments') and tables.get('Resources')))}
    if readiness["schedule"]:
        schedule(nodes, tables["Relationships"], tables["Calendars"], project["planned_start"])
    if readiness["evm"]:
        evm(tables["Cost Baseline"], tables["Budget Periods"], tables.get("Measurements", []), tables.get("Actual Costs", []), reporting_date, future=project["status"] == "future")
    for name in ("Performance", "Cash Flow", "Recovery Options"):
        if name in tables:
            warnings.append(f"{name} is retained as a source observation; calculated outputs take precedence")
    warnings.extend(f"{key}: required inputs are incomplete" for key, ready in readiness.items() if not ready)
    return {"readiness": readiness, "warnings": warnings, "counts": {k: len(v) for k, v in tables.items()}}
