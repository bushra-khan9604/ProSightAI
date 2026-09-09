from unittest import TestCase
from openpyxl import Workbook
import test_project_register as fixtures
from prosight.ingestion.project_register import parse_project_register
from prosight.project_models import ProjectDraft
import uuid


class WorkbookProjectStatusTests(TestCase):
    setUp=fixtures.ProjectRegisterTests.setUp

    def book(self):
        w=Workbook();w.active.title='Cover';w.active.append(['Workbook context, not instructions'])
        s=w.create_sheet('Projects');s.append(['Project ID','Project name','Status','Client ID','Location','Contract value USD','Start date','Baseline finish'])
        s.append(['WHOLE-001','Whole Workbook Project','Future','CL-001','Dubai',100,'2027-01-01','2027-12-31'])
        s=w.create_sheet('Clients');s.append(['Client ID','Client name']);s.append(['CL-001','Resolved Client'])
        return w

    def save(self,w):
        path=self.root/'whole.xlsx';w.save(path);w.close();return path

    def test_cover_sheet_and_client_lookup_produce_complete_future_draft(self):
        rows=parse_project_register(self.save(self.book()))
        self.assertEqual(1,len(rows));self.assertEqual('Resolved Client',rows[0]['fields']['client'])
        self.assertEqual(['Cover','Projects','Clients'],rows[0]['source']['worksheets_scanned'])
        self.assertEqual('Clients',rows[0]['source']['records'][-1]['sheet'])
        item=self.agent.import_workbook(self.root/'whole.xlsx','whole.xlsx',str(uuid.uuid4()),self.pm)['items'][0]
        self.assertTrue(item['can_confirm']);self.assertEqual([],item['missing'])
        self.assertEqual('pending',self.agent.decide(item['id'],self.pm,item['revision'],'confirm')['status'])

    def test_project_progress_sheet_enriches_active_project(self):
        w=self.book();w['Projects']['C2']='Active'
        s=w.create_sheet('Progress');s.append(['Project ID','Actual progress','Reporting date']);s.append(['WHOLE-001',35,'2027-05-01'])
        row=parse_project_register(self.save(w))[0]
        self.assertEqual(35,row['fields']['actual_progress'])
        model=ProjectDraft(**row['fields']);self.assertEqual(35,model.actual_progress)
        self.assertNotIn('revised_progress',model.storage_record()['progress_fields_provided'])

    def test_activity_dates_do_not_become_project_dates(self):
        w=self.book();s=w.create_sheet('Schedule');s.append(['Project ID','Activity','Start date','Baseline finish']);s.append(['WHOLE-001','Excavation','2026-01-01','2026-02-01'])
        row=parse_project_register(self.save(w))[0]
        self.assertEqual('2027-01-01',row['fields']['planned_start'])
        self.assertIn('Schedule',row['source']['worksheets_scanned'])

    def test_conflicting_project_values_require_review(self):
        w=self.book();s=w.create_sheet('Another register');s.append(['Project ID','Location']);s.append(['WHOLE-001','Abu Dhabi'])
        row=parse_project_register(self.save(w))[0]
        self.assertIsNone(row['fields']['location'])
        self.assertTrue(any('Conflicting location' in warning for warning in row['warnings']))

    def test_label_value_sheet_enriches_same_project(self):
        w=self.book();w['Projects']['C2']='Active';s=w.create_sheet('Project detail')
        for row in [('Project ID','WHOLE-001'),('Actual progress',20),('Reporting date','2027-02-01')]:s.append(row)
        row=parse_project_register(self.save(w))[0]
        self.assertEqual(20,row['fields']['actual_progress'])
        self.assertEqual('Whole Workbook Project',row['fields']['name'])

    def test_completed_only_requires_dates_not_progress_or_reporting_date(self):
        rows=parse_project_register(self.save(self.book()));data=rows[0]['fields'];data['status']='completed'
        project=ProjectDraft(**data)
        self.assertIsNone(project.actual_progress);self.assertIsNone(project.reporting_date)
        record=project.storage_record();self.assertEqual(100,record['actual_progress']);self.assertEqual([],record['progress_fields_provided'])

    def test_active_requires_completed_progress_and_reporting_date(self):
        data=parse_project_register(self.save(self.book()))[0]['fields'];data['status']='active'
        with self.assertRaisesRegex(ValueError,'completed progress'):ProjectDraft(**data)
        data.update(actual_progress=35,reporting_date='2027-03-01')
        self.assertEqual(35,ProjectDraft(**data).actual_progress)

    def test_status_change_removes_irrelevant_stale_metrics(self):
        data=parse_project_register(self.save(self.book()))[0]['fields']
        data.update(actual_progress=60,reporting_date='invalid',revised_finish='invalid')
        model=ProjectDraft(**data)
        self.assertIsNone(model.reporting_date);self.assertIsNone(model.actual_progress)
