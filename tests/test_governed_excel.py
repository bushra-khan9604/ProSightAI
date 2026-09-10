"""Offline tests for governed Excel validation, approval, and publication."""

from __future__ import annotations

import tempfile
import unittest
import uuid
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock

from openpyxl import Workbook

from prosight.ingestion.governed_excel import (
    FieldSpec, MappingRegistry, OrganizationMapping, ReferenceCatalog,
    SheetMapping, SourceFileMetadata, convert_value, source_dedupe_key,
    validate_workbook,
)
from prosight.ingestion.publication import (
    AtomicPublicationAdapter, BATCH_TRANSITIONS, PUBLISH_SQL, PublicationRequest, bind_approval,
    validate_batch_transition,
)


class GovernedExcelTests(unittest.TestCase):
    def setUp(self):
        self.organization_id = str(uuid.uuid4())
        self.project_id = str(uuid.uuid4())
        self.batch_id = str(uuid.uuid4())
        self.source_file_id = str(uuid.uuid4())
        self.mapping_profile_id = str(uuid.uuid4())
        self.mapping_version_id = str(uuid.uuid4())
        self.profile = OrganizationMapping(
            self.organization_id, self.mapping_profile_id, self.mapping_version_id, 1,
            (SheetMapping(
                "invoice", "Invoices",
                {"project_id": "Project ID", "invoice_number": "Invoice", "invoice_date": "Date",
                 "currency_code": "Currency", "amount": "Amount"},
                {"project_id": FieldSpec("uuid", True), "invoice_number": FieldSpec("text", True),
                 "invoice_date": FieldSpec("date", True), "currency_code": FieldSpec("currency", True),
                 "amount": FieldSpec("decimal", True, 2)},
                ("project_id", "invoice_number"),
            ),),
        )
        self.references = ReferenceCatalog(
            frozenset({self.organization_id}), {self.project_id: self.organization_id}
        )

    def workbook(self, rows):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "input.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Invoices"
        sheet.append(["Project ID", "Invoice", "Date", "Currency", "Amount"])
        for row in rows:
            sheet.append(row)
        book.save(path)
        return path

    def validate(self, path, **kwargs):
        return validate_workbook(
            path, self.profile, self.references, source_file_id=self.source_file_id,
            import_batch_id=self.batch_id, **kwargs,
        )

    def test_exact_types_lineage_and_reconciliation(self):
        preview = self.validate(self.workbook([
            [self.project_id, "INV-1", "2026-09-01", "aed", "10.25"],
            [self.project_id, "INV-2", "2026-09-02", "AED", 20],
        ]), expected_totals={("invoice", "amount"): Decimal("30.25")})
        self.assertTrue(preview.valid)
        self.assertEqual("10.25", preview.rows[0].values["amount"])
        self.assertEqual("AED", preview.rows[0].values["currency_code"])
        self.assertEqual(self.source_file_id, preview.rows[0].lineage["source_file_id"])
        self.assertEqual(self.batch_id, preview.import_batch_id)
        self.assertEqual(self.source_file_id, preview.source_file_id)
        self.assertEqual(self.mapping_profile_id, preview.mapping_profile_id)
        self.assertEqual(self.profile.mapping_checksum, preview.profile_checksum)
        self.assertEqual("invoice", preview.sheet_classifications["Invoices"])

    def test_formulas_are_data_and_are_never_accepted_as_values(self):
        preview = self.validate(self.workbook([
            [self.project_id, "INV-1", "2026-09-01", "AED", "=5+5"],
        ]))
        self.assertFalse(preview.valid)
        self.assertEqual(("Invoices!E2",), preview.formula_cells)
        self.assertIn("formula", {issue.code for issue in preview.issues})

    def test_invalid_currency_date_reference_and_decimal_are_deterministic(self):
        other_project = str(uuid.uuid4())
        preview = self.validate(self.workbook([
            [other_project, "INV-1", "09/01/2026", "ZZZ", "1.234"],
        ]))
        messages = " ".join(issue.message for issue in preview.issues)
        self.assertIn("ISO date", messages)
        self.assertIn("ISO 4217", messages)
        self.assertIn("at most 2", messages)
        self.assertIn("same organization", messages)

    def test_business_key_duplicate_and_row_limit_are_rejected(self):
        row = [self.project_id, "INV-1", "2026-09-01", "AED", 1]
        preview = self.validate(self.workbook([row, row]))
        self.assertIn("duplicate_business_key", {issue.code for issue in preview.issues})
        with self.assertRaisesRegex(ValueError, "1-row limit"):
            self.validate(self.workbook([row, [self.project_id, "INV-2", "2026-09-02", "AED", 2]]), max_rows=1)

    def test_mapping_versions_are_immutable_and_suggestions_use_same_validator(self):
        registry = MappingRegistry()
        registry.register(self.profile)
        registry.register(self.profile)
        changed = OrganizationMapping(
            self.organization_id, self.mapping_profile_id, str(uuid.uuid4()), 1,
            self.profile.sheets,
        )
        with self.assertRaisesRegex(ValueError, "immutable"):
            registry.register(changed)
        bad_mapping = SheetMapping(
            "invoice", "Invoices", {**self.profile.sheets[0].columns, "amount": "Missing"},
            self.profile.sheets[0].fields, self.profile.sheets[0].business_key,
        )
        suggested = OrganizationMapping(
            self.organization_id, self.mapping_profile_id, str(uuid.uuid4()), 2,
            (bad_mapping,),
        )
        preview = validate_workbook(
            self.workbook([[self.project_id, "INV-1", "2026-09-01", "AED", 1]]),
            suggested, self.references, source_file_id=self.source_file_id,
            import_batch_id=self.batch_id,
        )
        self.assertIn("missing_column", {issue.code for issue in preview.issues})

    def test_multiple_profiles_can_each_have_version_one_and_ids_bind_checksums(self):
        registry = MappingRegistry()
        registry.register(self.profile)
        other_profile = OrganizationMapping(
            self.organization_id, str(uuid.uuid4()), str(uuid.uuid4()), 1,
            self.profile.sheets,
        )
        registry.register(other_profile)
        self.assertNotEqual(self.profile.mapping_checksum, other_profile.mapping_checksum)
        same_profile_new_version_id = OrganizationMapping(
            self.organization_id, self.mapping_profile_id, str(uuid.uuid4()), 1,
            self.profile.sheets,
        )
        with self.assertRaisesRegex(ValueError, "immutable"):
            registry.register(same_profile_new_version_id)

    def test_source_file_dedupe_is_org_and_kind_scoped(self):
        digest = "a" * 64
        self.assertEqual(source_dedupe_key(self.organization_id, "xlsx", digest),
                         source_dedupe_key(self.organization_id, "XLSX", digest))
        self.assertNotEqual(source_dedupe_key(self.organization_id, "xlsx", digest),
                            source_dedupe_key(str(uuid.uuid4()), "xlsx", digest))
        metadata = SourceFileMetadata(
            self.source_file_id, self.organization_id, "xlsx", "source.xlsx",
            digest, 100, str(uuid.uuid4()),
        )
        self.assertEqual(source_dedupe_key(self.organization_id, "xlsx", digest), metadata.dedupe_key)

    def test_conversion_never_uses_float_for_money_arithmetic(self):
        self.assertEqual(Decimal("0.10"), convert_value(0.1, FieldSpec("decimal", True, 2)))
        with self.assertRaises(ValueError):
            convert_value(float("nan"), FieldSpec("decimal", True, 2))

    def test_empty_workbook_is_not_approvable(self):
        preview = self.validate(self.workbook([]))
        self.assertFalse(preview.valid)
        self.assertIn("no_records", {issue.code for issue in preview.issues})


class ApprovalAndPublicationTests(unittest.TestCase):
    def test_lifecycle_is_monotonic_and_cannot_bypass_approval(self):
        expected = {
            "uploaded": {"profiling", "cancelled"},
            "profiling": {"mapping", "validation_failed", "cancelled"},
            "mapping": {"validating", "validation_failed", "cancelled"},
            "validating": {"review_ready", "validation_failed", "cancelled"},
            "review_ready": {"awaiting_approval", "validating", "cancelled"},
            "awaiting_approval": {"approved", "rejected", "validating", "cancelled"},
            "approved": {"publishing", "cancelled"},
            "publishing": {"published", "publish_failed"},
            "publish_failed": {"approved", "publishing", "cancelled"},
            "validation_failed": {"profiling", "mapping", "validating", "cancelled"},
        }
        self.assertEqual(expected, BATCH_TRANSITIONS)
        for current, requested_states in expected.items():
            for requested in requested_states:
                validate_batch_transition(current, requested)
        with self.assertRaises(ValueError):
            validate_batch_transition("validating", "published")

    def test_approval_binds_exact_preview_and_atomic_adapter_makes_one_call(self):
        organization_id, project_id = str(uuid.uuid4()), str(uuid.uuid4())
        batch_id, source_id = str(uuid.uuid4()), str(uuid.uuid4())
        profile_id, version_id = str(uuid.uuid4()), str(uuid.uuid4())
        profile = OrganizationMapping(organization_id, profile_id, version_id, 1, (SheetMapping(
            "risk", "Risks", {"project_id": "Project", "risk_number": "Risk", "title": "Title"},
            {"project_id": FieldSpec("uuid", True), "risk_number": FieldSpec("text", True), "title": FieldSpec("text", True)},
            ("project_id", "risk_number"),
        ),))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "risk.xlsx"
            book = Workbook(); sheet = book.active; sheet.title = "Risks"
            sheet.append(["Project", "Risk", "Title"]); sheet.append([project_id, "R-1", "Delay"]); book.save(path)
            preview = validate_workbook(path, profile, ReferenceCatalog(frozenset({organization_id}), {project_id: organization_id}),
                                        source_file_id=source_id, import_batch_id=batch_id)
        binding = bind_approval(
            preview, approval_id=str(uuid.uuid4()), batch_id=batch_id,
            source_file_id=source_id, mapping_profile_id=profile_id,
            mapping_version_id=version_id, mapping_version_no=1,
            requester_id=str(uuid.uuid4()), approver_id=str(uuid.uuid4()),
            approver_role="admin", decision="approved",
        )
        request = PublicationRequest.from_binding(binding)
        connection = Mock()
        connection.__enter__ = Mock(return_value=connection); connection.__exit__ = Mock(return_value=False)
        connection.execute_native.return_value.fetchone.side_effect = [
            {"status": "published"}, {"status": "already_published"},
        ]
        adapter = AtomicPublicationAdapter(Mock(return_value=connection))
        self.assertEqual("published", adapter.publish(request)["status"])
        self.assertEqual("already_published", adapter.publish(request)["status"])
        self.assertEqual(2, connection.execute_native.call_count)
        self.assertIn("ingestion.publish_import_batch", PUBLISH_SQL)
        self.assertNotIn("construction.", PUBLISH_SQL)
        with self.assertRaises(PermissionError):
            bind_approval(
                preview, approval_id=str(uuid.uuid4()), batch_id=batch_id,
                source_file_id=source_id, mapping_profile_id=profile_id,
                mapping_version_id=version_id, mapping_version_no=1,
                requester_id=str(uuid.uuid4()), approver_id=str(uuid.uuid4()),
                approver_role="manager", decision="approved",  # type: ignore[arg-type]
            )

    def test_empty_rejected_preview_still_rejects_batch_source_and_profile_mismatch(self):
        organization_id, project_id = str(uuid.uuid4()), str(uuid.uuid4())
        batch_id, source_id = str(uuid.uuid4()), str(uuid.uuid4())
        profile_id, version_id = str(uuid.uuid4()), str(uuid.uuid4())
        profile = OrganizationMapping(organization_id, profile_id, version_id, 1, (SheetMapping(
            "risk", "Risks", {"project_id": "Project", "risk_number": "Risk"},
            {"project_id": FieldSpec("uuid", True), "risk_number": FieldSpec("text", True)},
            ("project_id", "risk_number"),
        ),))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty.xlsx"
            book = Workbook(); sheet = book.active; sheet.title = "Risks"
            sheet.append(["Project", "Risk"]); book.save(path)
            preview = validate_workbook(
                path, profile,
                ReferenceCatalog(frozenset({organization_id}), {project_id: organization_id}),
                source_file_id=source_id, import_batch_id=batch_id,
            )
        base = dict(
            preview=preview, approval_id=str(uuid.uuid4()), batch_id=batch_id,
            source_file_id=source_id, mapping_profile_id=profile_id,
            mapping_version_id=version_id, mapping_version_no=1,
            requester_id=str(uuid.uuid4()), approver_id=str(uuid.uuid4()),
            approver_role="admin", decision="rejected",
        )
        binding = bind_approval(**base)
        self.assertEqual(batch_id, binding.batch_id)
        for changed, message in (
            ({"batch_id": str(uuid.uuid4())}, "batch"),
            ({"source_file_id": str(uuid.uuid4())}, "source file"),
            ({"mapping_profile_id": str(uuid.uuid4())}, "mapping profile"),
        ):
            with self.assertRaisesRegex(ValueError, message):
                bind_approval(**{**base, **changed})


if __name__ == "__main__":
    unittest.main()
