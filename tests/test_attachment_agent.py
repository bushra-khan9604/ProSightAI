import json
from pathlib import Path
from contextlib import closing
import tempfile
import unittest
from unittest.mock import Mock,patch
from openpyxl import Workbook
from prosight.repository import ProjectRepository,DEFAULT_DATA
from prosight.agents.attachment_update import AttachmentUpdateAgent
from prosight.ingestion.excel import PROJECT_COLUMNS
from prosight.ingestion.portfolio import SCHEDULE_HEADERS

class AttachmentAgentTests(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
  self.repo=ProjectRepository(self.root/'db.sqlite');self.repo.initialize(DEFAULT_DATA)
  self.code=self.repo.list_projects(user_role='admin')[0]['code']
  self.admin={'id':'admin1','role':'admin'};self.pm={'id':'pm1','role':'project_manager'};self.engineer={'id':'eng1','role':'planning_engineer'}
  self.agent=AttachmentUpdateAgent(self.repo,None,self.root/'uploads')
 def workbook(self,kind='schedule'):
  path=self.root/'source.xlsx';book=Workbook();sheet=book.active
  if kind=='schedule':
   sheet.title='Project Schedule';sheet.append(SCHEDULE_HEADERS);sheet.append(['A1','Excavation','2026-01-01','2026-02-01',31])
  else:
   sheet.title='Projects';sheet.append(PROJECT_COLUMNS)
   source=self.repo.list_projects(user_role='admin')[0];source['code']='NEW-ATTACH-001';source['name']='Attachment created project'
   sheet.append([source.get(k) for k in PROJECT_COLUMNS])
  book.save(path);book.close();return path
 def prepare(self,user=None,kind='schedule'):
  return self.agent.prepare(self.workbook('project' if kind=='new_project' else kind),'source.xlsx','Update this project using the attached data',self.code,kind,user or self.pm)
 def test_preview_then_uploader_confirmation_then_admin_applies_once(self):
  op=self.prepare();self.assertEqual('awaiting_confirmation',op['status'])
  self.assertEqual([],self.repo.list_project_schedule(self.code))
  with self.assertRaises(ValueError):self.agent.decide(op['id'],self.admin,op['preview_token'],'approve')
  result=self.agent.decide(op['id'],self.pm,op['preview_token'],'confirm');self.assertEqual('pending',result['status'])
  self.assertEqual([],self.repo.list_project_schedule(self.code))
  result=self.agent.decide(op['id'],self.admin,op['preview_token'],'approve');self.assertEqual('completed',result['status'])
  self.assertEqual(1,len(self.repo.list_project_schedule(self.code)))
  self.agent.decide(op['id'],self.admin,op['preview_token'],'approve')
  with closing(self.repo.connect()) as db:self.assertEqual(1,db.execute('SELECT count(*) FROM portfolio_imports').fetchone()[0])
 def test_owner_and_role_boundaries(self):
  op=self.prepare()
  with self.assertRaises(PermissionError):self.agent.get(op['id'],{'id':'pm2','role':'project_manager'})
  with self.assertRaises(PermissionError):self.agent.decide(op['id'],self.engineer,op['preview_token'],'confirm')
  with self.assertRaises(PermissionError):self.agent.decide(op['id'],self.pm,op['preview_token'],'approve')
  with self.assertRaises(PermissionError):self.prepare(self.engineer,'new_project')
 def test_admin_can_confirm_own_new_project(self):
  op=self.prepare(self.admin,'new_project');self.assertEqual('NEW-ATTACH-001',op['project_code'])
  self.agent.decide(op['id'],self.admin,op['preview_token'],'confirm')
  self.assertIsNotNone(self.repo.find_project('NEW-ATTACH-001','admin'))
 def test_stale_database_requires_fresh_preview(self):
  op=self.prepare()
  with closing(self.repo.connect()) as db:
   with db:db.execute('UPDATE projects SET name=? WHERE code=?',('Changed concurrently',self.code))
  result=self.agent.decide(op['id'],self.pm,op['preview_token'],'confirm')
  self.assertEqual('stale',result['status']);self.assertEqual([],self.repo.list_project_schedule(self.code))
 def test_changed_file_is_not_applied(self):
  op=self.prepare();path=next((self.root/'uploads'/self.code).glob('*.xlsx'));path.write_bytes(b'changed')
  self.assertEqual('stale',self.agent.decide(op['id'],self.pm,op['preview_token'],'confirm')['status'])
 def test_duplicate_upload_returns_same_operation(self):
  source=self.workbook();args=(source,'source.xlsx','Update schedule from this file',self.code,'schedule',self.pm)
  first=self.agent.prepare(*args);second=self.agent.prepare(*args)
  self.assertEqual(first['id'],second['id'])
 def test_transaction_failure_leaves_confirmation_and_no_import(self):
  op=self.prepare(self.admin)
  with patch.object(ProjectRepository,'apply_portfolio_import',side_effect=ValueError('test failure')):
   with self.assertRaises(ValueError):self.agent.decide(op['id'],self.admin,op['preview_token'],'confirm')
  self.assertEqual('awaiting_confirmation',self.agent.get(op['id'],self.admin)['status'])
  with closing(self.repo.connect()) as db:self.assertEqual(0,db.execute('SELECT count(*) FROM portfolio_imports').fetchone()[0])
 def test_formula_rejected_before_preview(self):
  source=self.workbook();from openpyxl import load_workbook
  book=load_workbook(source);book.active['B2']='=1+1';book.save(source);book.close()
  with self.assertRaisesRegex(ValueError,'formulas'):
   self.agent.prepare(source,'source.xlsx','Update schedule from workbook',self.code,'schedule',self.pm)
 def test_pdf_requires_approval_and_external_permission(self):
  source=self.root/'file.pdf';source.write_bytes(b'%PDF-test')
  with patch('prosight.agents.attachment_update.extract_pdf_chunks',return_value=[{'text':'test evidence','metadata':{'page_number':1}}]),patch('prosight.agents.attachment_update.detect_reporting_date',return_value='2026-01-01'):
   op=self.agent.prepare(source,'file.pdf','Add project PDF evidence',self.code,'pdf',self.admin,False)
  with self.assertRaisesRegex(ValueError,'External'):
   self.agent.decide(op['id'],self.admin,op['preview_token'],'confirm')
  self.assertEqual([],self.repo.list_documents(self.code))
 def test_pdf_approval_registers_and_schedules_once(self):
  ingestion=Mock();self.agent.ingestion=ingestion
  source=self.root/'file.pdf';source.write_bytes(b'%PDF-test')
  with patch('prosight.agents.attachment_update.extract_pdf_chunks',return_value=[{'text':'test evidence','metadata':{'page_number':1}}]),patch('prosight.agents.attachment_update.detect_reporting_date',return_value='2026-01-01'):
   op=self.agent.prepare(source,'file.pdf','Add project PDF evidence',self.code,'pdf',self.pm,True)
  self.agent.decide(op['id'],self.pm,op['preview_token'],'confirm');ingestion.executor.submit.assert_not_called()
  self.agent.decide(op['id'],self.admin,op['preview_token'],'approve');self.agent.decide(op['id'],self.admin,op['preview_token'],'approve')
  ingestion.executor.submit.assert_called_once();self.assertEqual(1,len(self.repo.list_documents(self.code)))

 def test_pdf_real_extraction_replacement_and_failed_retry(self):
  from reportlab.pdfgen import canvas
  from prosight.ingestion.manager import IngestionManager
  source=self.root/'real.pdf';page=canvas.Canvas(str(source));page.drawString(40,700,'Project report 2026-09-08. Excavation is complete.');page.save()
  rag=Mock();manager=IngestionManager(self.repo,rag,self.root/'uploads');self.addCleanup(manager.executor.shutdown,wait=True)
  # Deterministic synchronous dispatch exercises the same complete publication flow.
  manager.executor.shutdown(wait=True)
  manager.executor=Mock();manager.executor.submit.side_effect=lambda f,*args:f(*args)
  self.agent.ingestion=manager
  old=self.repo.create_document(self.code,'old.pdf','pdf','old-checksum',str(source))
  with closing(self.repo.connect()) as db:
   with db:db.execute("UPDATE documents SET approval_status='approved',status='ready',index_status='ready' WHERE id=?",(old['id'],))
  op=self.agent.prepare(source,'real.pdf','Replace the previous report with this PDF',self.code,'pdf',self.admin,True,old['id'])
  rag.add_chunks.side_effect=RuntimeError('provider offline')
  result=self.agent.decide(op['id'],self.admin,op['preview_token'],'confirm')
  self.assertEqual('failed',result['status']);self.assertEqual('approved',self.repo.get_document(old['id'])['approval_status'])
  rag.add_chunks.side_effect=None
  result=self.agent.decide(op['id'],self.admin,op['preview_token'],'retry')
  self.assertEqual('completed',result['status']);self.assertEqual('superseded',self.repo.get_document(old['id'])['approval_status'])
  self.assertEqual('ready',self.repo.get_document(result['result']['document_id'])['index_status'])

 def test_http_identity_and_saved_preview_endpoint(self):
  from fastapi.testclient import TestClient
  from prosight.api import create_app
  with patch('prosight.api.RAGStore',side_effect=RuntimeError('offline')):
   app=create_app(self.repo)
  app.state.runtime.attachments=self.agent
  # Router holds the original agent; use its isolated root too.
  with TestClient(app) as client:
   self.assertEqual(401,client.get('/api/attachments').status_code)
   op=self.prepare()
   with patch.object(self.repo,'get_session_user',return_value=self.pm):
    response=client.get('/api/attachments/'+op['id']);self.assertEqual(200,response.status_code)
    self.assertNotIn('stored_path',response.json())
    denied=client.post('/api/attachments/'+op['id']+'/decision',json={'decision':'approve','preview_token':op['preview_token']})
    self.assertEqual(403,denied.status_code)
   with patch.object(self.repo,'get_session_user',return_value={'id':'another','role':'project_manager'}):
    self.assertEqual(403,client.get('/api/attachments/'+op['id']).status_code)
