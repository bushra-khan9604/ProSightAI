"""Tests for platform catalog and organization-owned label overlays."""

from __future__ import annotations

import tempfile
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook

from prosight.ingestion.catalog import (
    OrganizationSheetProfile,
    compile_sheet_profile,
    default_project_sheet,
    load_catalog,
    organization_profile_json_schema,
    validate_compiled_mapping,
)
from prosight.ingestion.governed_excel import (
    OrganizationMapping,
    ReferenceCatalog,
    validate_workbook,
)


class IngestionCatalogTests(unittest.TestCase):
    def setUp(self):
        self.organization_id = str(uuid.uuid4())
        self.mapping = OrganizationMapping(
            self.organization_id,
            str(uuid.uuid4()),
            str(uuid.uuid4()),
            1,
            (compile_sheet_profile(OrganizationSheetProfile(
                entity_type="projects",
                sheet_name="Job Summary",
                sheet_aliases=["Our Projects"],
                columns={
                    "code": ["Job Ref", "Project No."],
                    "name": ["Job Description", "Project Title"],
                    "contract_value": ["Award Amount"],
                },
            )),),
            catalog_version=load_catalog().catalog_version,
        )

    def workbook(self, title="Our Projects", headers=None, values=None) -> Path:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "projects.xlsx"
        book = Workbook(); sheet = book.active; sheet.title = title
        sheet.append(headers or ["Project No", "Project Title"])
        sheet.append(values or ["JOB-001", "Harbour Offices"])
        book.save(path)
        return path

    def test_catalog_matches_all_database_publication_entities(self):
        catalog = load_catalog()
        self.assertEqual(1, catalog.catalog_version)
        self.assertEqual({
            "projects", "employees", "activities", "invoices", "risks", "rfis",
            "nonconformance_reports", "safety_incidents", "daily_reports",
        }, set(catalog.entities))
        self.assertEqual("object", organization_profile_json_schema()["type"])

    def test_organization_aliases_map_to_canonical_columns(self):
        preview = validate_workbook(
            self.workbook(), self.mapping,
            ReferenceCatalog(frozenset({self.organization_id}), {}),
            source_file_id=str(uuid.uuid4()), import_batch_id=str(uuid.uuid4()),
        )
        self.assertTrue(preview.valid, preview.issues)
        self.assertEqual("JOB-001", preview.rows[0].values["code"])
        self.assertEqual("Harbour Offices", preview.rows[0].values["name"])
        self.assertNotIn("contract_value", preview.rows[0].values)
        self.assertEqual("Our Projects", preview.rows[0].lineage["sheet_name"])

    def test_organization_cannot_map_scope_or_unknown_database_columns(self):
        for field_name in ("organization_id", "created_by", "unknown_column"):
            with self.subTest(field_name=field_name), self.assertRaisesRegex(
                ValueError, "not an organization-configurable import field"
            ):
                compile_sheet_profile(OrganizationSheetProfile(
                    entity_type="projects", sheet_name="Projects",
                    columns={"code": "Code", "name": "Name", field_name: "Unsafe"},
                ))

    def test_ambiguous_sheet_aliases_are_rejected_without_guessing(self):
        path = self.workbook(title="Job Summary")
        book = Workbook(); first = book.active; first.title = "Job Summary"
        first.append(["Job Ref", "Job Description"]); first.append(["A-1", "One"])
        second = book.create_sheet("Our Projects")
        second.append(["Job Ref", "Job Description"]); second.append(["A-2", "Two"])
        book.save(path)
        preview = validate_workbook(
            path, self.mapping, ReferenceCatalog(frozenset({self.organization_id}), {}),
            source_file_id=str(uuid.uuid4()), import_batch_id=str(uuid.uuid4()),
        )
        self.assertIn("ambiguous_sheet", {issue.code for issue in preview.issues})

    def test_project_scope_is_injected_instead_of_read_from_workbook(self):
        project_id = str(uuid.uuid4())
        sheet = compile_sheet_profile(OrganizationSheetProfile(
            entity_type="risks", sheet_name="Risk Register",
            columns={
                "risk_number": "Risk Ref", "title": "Risk Title",
                "description": "Details",
            },
        ))
        mapping = OrganizationMapping(
            self.organization_id, str(uuid.uuid4()), str(uuid.uuid4()), 1, (sheet,),
            catalog_version=1,
        )
        path = self.workbook(
            title="Risk Register",
            headers=["Risk Ref", "Risk Title", "Details"],
            values=["R-01", "Late steel", "Supplier delay"],
        )
        preview = validate_workbook(
            path, mapping,
            ReferenceCatalog(frozenset({self.organization_id}), {project_id: self.organization_id}),
            source_file_id=str(uuid.uuid4()), import_batch_id=str(uuid.uuid4()),
            project_id=project_id,
        )
        self.assertTrue(preview.valid, preview.issues)
        self.assertEqual(project_id, preview.rows[0].values["project_id"])

    def test_default_project_profile_is_compiled_from_catalog(self):
        sheet = compile_sheet_profile(default_project_sheet())
        self.assertEqual("projects", sheet.entity_type)
        self.assertIn("Project ID", sheet.source_labels("code"))
        self.assertEqual(("code",), sheet.business_key)
        mapping = OrganizationMapping(
            self.organization_id, str(uuid.uuid4()), str(uuid.uuid4()), 1, (sheet,),
            catalog_version=1,
        )
        preview = validate_workbook(
            self.workbook(title="Project Register", headers=["Project ID", "Project Name"],
                          values=["PRJ-100", "Metro Station"]),
            mapping, ReferenceCatalog(frozenset({self.organization_id}), {}),
            source_file_id=str(uuid.uuid4()), import_batch_id=str(uuid.uuid4()),
        )
        self.assertTrue(preview.valid, preview.issues)
        self.assertEqual({"organization_id", "code", "name"}, set(preview.rows[0].values))

    def test_pre_catalog_and_type_overrides_are_not_usable(self):
        legacy = OrganizationMapping(
            self.organization_id, str(uuid.uuid4()), str(uuid.uuid4()), 1,
            self.mapping.sheets,
        )
        with self.assertRaisesRegex(ValueError, "active construction catalog"):
            validate_compiled_mapping(legacy)


if __name__ == "__main__":
    unittest.main()
