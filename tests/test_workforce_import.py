"""Governed workbook mapping and approval tests for workforce datasets."""

import json
import tempfile
import unittest
from unittest.mock import patch
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook

from prosight.agents.attachment_update import AttachmentUpdateAgent
from prosight.ingestion.workforce import parse_workforce_workbook
from prosight.repository import DEFAULT_DATA, ProjectRepository


class WorkforceImportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.repository = ProjectRepository(self.root / "prosight.db")
        self.repository.initialize(DEFAULT_DATA)
        self.project = self.repository.list_projects(user_role="admin")[0]["code"]
        self.agent = AttachmentUpdateAgent(self.repository, None, self.root / "uploads")
        self.manager = {"id": "pm-workforce", "role": "project_manager"}
        self.admin = {"id": "admin-workforce", "role": "admin"}

    def workbook(self, dataset):
        path = self.root / f"{dataset}.xlsx"
        book = Workbook()
        sheet = book.active
        if dataset == "employees_training":
            sheet.title = "Employees"
            sheet.append(["Employees and training reference layout"])
            sheet.append([])
            sheet.append([])
            sheet.append([])
            sheet.append([
                "Employee ID", "Synthetic name", "Role", "Department",
                "Home project ID", "Employment type", "Join date", "Basic AED",
                "Allowance AED", "Gross monthly AED", "Status", "Data origin",
            ])
            sheet.append(["EMP-T01", "Synthetic One", "Engineer", "Delivery", self.project,
                          "Employee", "2026-01-10", 1000, 250, "=H6+I6", "Active", "Test"])
            training = book.create_sheet("Training")
            training.append(["Training reference layout"])
            training.append([]); training.append([]); training.append([])
            training.append(["Training ID", "Employee ID", "Course", "Completed date",
                             "Expiry date", "Status", "Evidence ID"])
            training.append(["TRN-T01", "EMP-T01", "Safety", "2026-02-01",
                             "2027-02-01", "Valid", "EVD-T01"])
        elif dataset == "attendance_payroll":
            sheet.title = "Attendance"
            sheet.append(["Attendance and payroll reference layout"])
            sheet.append([]); sheet.append([]); sheet.append([])
            sheet.append(["Timesheet ID", "Employee ID", "Project ID", "Work date",
                          "Attendance status", "Regular hours", "OT hours", "Total hours"])
            sheet.append(["TIME-T01", "EMP-T01", self.project, "2026-03-02",
                          "Present", 8, 2, "=F6+G6"])
            payroll = book.create_sheet("Payroll")
            payroll.append(["Payroll reference layout"])
            payroll.append([]); payroll.append([]); payroll.append([])
            payroll.append(["Payroll ID", "Employee ID", "Project ID", "Month", "Basic AED",
                            "Allowance AED", "OT hours", "OT rate AED", "OT pay AED", "Gross AED",
                            "Employer burden AED", "Total cost AED", "Payment date", "Days worked",
                            "Paid leave days"])
            payroll.append(["PAY-T01", "EMP-T01", self.project, "2026-03-01", 1000, 250,
                            2, 50, "=G6*H6", "=E6+F6", 100, "=I6+J6+K6",
                            "2026-03-31", 22, 1])
        else:
            sheet.title = "Direct allocation"
            sheet.append(["Manpower and deployment reference layout"])
            sheet.append([]); sheet.append([]); sheet.append([])
            sheet.append(["Employee ID", "Project ID", "Month", "FTE allocation",
                          "Regular hours", "OT hours", "Payroll cost AED", "Basis"])
            sheet.append(["EMP-T01", self.project, "2026-04-01", 1, 176, 4, 1450, "Monthly"])
            crews = book.create_sheet("Subcontract crews")
            crews.append(["Crew reference layout"])
            crews.append([]); crews.append([]); crews.append([])
            crews.append(["Subcontract ID", "Project ID", "Vendor ID", "Month", "Trade",
                          "Planned workers", "Actual workers", "Worker variance",
                          "Available labor hours", "Definition"])
            crews.append(["SUB-T01", self.project, "VEN-T01", "2026-04-01", "Electrical",
                          10, 12, "=G6-F6", "=G6*208", "Monthly crew"])
            forecast = book.create_sheet("Forecast deployment")
            forecast.append(["Forecast reference layout"])
            forecast.append([]); forecast.append([]); forecast.append([])
            forecast.append(["Project ID", "Month", "Direct headcount",
                             "Subcontract workers", "Basis"])
            forecast.append([self.project, "2026-05-01", 15, 12, "Approved plan"])
        book.save(path)
        book.close()
        return path

    def approve(self, dataset):
        operation = self.agent.prepare(
            self.workbook(dataset), f"{dataset}.xlsx",
            "Add the mapped workforce records after review", "", dataset, self.manager,
        )
        self.assertEqual("PORTFOLIO", operation["project_code"])
        self.assertEqual("awaiting_confirmation", operation["status"])
        operation = self.agent.decide(
            operation["id"], self.manager, operation["preview_token"], "confirm"
        )
        self.assertEqual("pending", operation["status"])
        return self.agent.decide(
            operation["id"], self.admin, operation["preview_token"], "approve"
        )

    def test_formula_columns_are_recomputed_and_mapping_is_reviewable(self):
        parsed = parse_workforce_workbook(
            self.workbook("employees_training"), lambda value: self.project,
            lambda value: None, "employees_training",
        )
        self.assertEqual(1250, parsed["employees"][0]["gross_monthly_aed"])
        self.assertEqual("EmployeeRecord", parsed["semantic_mappings"]["Employees"]["schema"])
        self.assertIn("gross_monthly_aed", parsed["warnings"][0])

    def test_head_office_assignments_apply_without_creating_a_project(self):
        original_projects = self.repository.list_projects(user_role='admin')
        self.project = 'Head Office'
        for dataset in ('employees_training', 'attendance_payroll', 'manpower_deployment'):
            self.assertEqual('completed', self.approve(dataset)['status'])
        with closing(self.repository.connect()) as db:
            self.assertEqual('HO', db.execute('SELECT home_project_id FROM employees').fetchone()[0])
            for table in ('attendance_records', 'payroll_records', 'direct_allocations',
                          'subcontract_crews', 'deployment_forecasts'):
                self.assertEqual('HO', db.execute(f'SELECT project_id FROM {table}').fetchone()[0])
        self.assertEqual(original_projects, self.repository.list_projects(user_role='admin'))
        self.assertIsNone(self.repository.resolve_project_reference('HO'))

    def test_head_office_aliases_are_canonical_and_unknown_projects_still_fail(self):
        for alias in ('HO', 'ho', 'H.O.', 'head-office'):
            self.project = alias
            parsed = parse_workforce_workbook(self.workbook('employees_training'),
                                             lambda value: None, lambda value: None, 'employees_training')
            self.assertEqual('HO', parsed['employees'][0]['home_project_id'])
        self.project = 'UNKNOWN'
        with self.assertRaisesRegex(ValueError, 'unknown project'):
            parse_workforce_workbook(self.workbook('employees_training'),
                                     lambda value: None, lambda value: None, 'employees_training')

    def test_all_three_workbook_groups_apply_only_after_approval(self):
        first = self.agent.prepare(
            self.workbook("employees_training"), "employees.xlsx",
            "Add employees and training after review", "", "employees_training", self.manager,
        )
        with closing(self.repository.connect()) as db:
            self.assertEqual(0, db.execute("SELECT COUNT(*) FROM employees").fetchone()[0])
        self.agent.decide(first["id"], self.manager, first["preview_token"], "confirm")
        with closing(self.repository.connect()) as db:
            self.assertEqual(0, db.execute("SELECT COUNT(*) FROM employees").fetchone()[0])
        result = self.agent.decide(first["id"], self.admin, first["preview_token"], "approve")
        self.assertEqual("completed", result["status"])

        self.assertEqual("completed", self.approve("attendance_payroll")["status"])
        self.assertEqual("completed", self.approve("manpower_deployment")["status"])
        with closing(self.repository.connect()) as db:
            counts = {
                table: db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                for table in (
                    "employees", "employee_training", "attendance_records", "payroll_records",
                    "direct_allocations", "subcontract_crews", "deployment_forecasts",
                )
            }
            self.assertEqual({table: 1 for table in counts}, counts)
            attendance = db.execute("SELECT total_hours FROM attendance_records").fetchone()
            payroll = db.execute("SELECT ot_pay_aed,gross_aed,total_cost_aed FROM payroll_records").fetchone()
            crew = db.execute("SELECT worker_variance,available_labor_hours FROM subcontract_crews").fetchone()
        self.assertEqual(10, attendance["total_hours"])
        self.assertEqual((100, 1250, 1450), tuple(payroll))
        self.assertEqual((2, 2496), tuple(crew))

    def test_unknown_employee_stops_preview_without_database_changes(self):
        with self.assertRaisesRegex(ValueError, "unknown employee"):
            self.agent.prepare(
                self.workbook("attendance_payroll"), "attendance.xlsx",
                "Add attendance after review", "", "attendance_payroll", self.manager,
            )
        with closing(self.repository.connect()) as db:
            self.assertEqual(0, db.execute("SELECT COUNT(*) FROM attendance_records").fetchone()[0])

    def test_formula_outside_an_approved_calculated_column_is_rejected(self):
        path = self.workbook("employees_training")
        from openpyxl import load_workbook
        book = load_workbook(path)
        book["Employees"]["C6"] = "=1+1"
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, "formulas are not allowed in Role"):
            self.agent.prepare(
                path, "employees.xlsx", "Add employees after review", "",
                "employees_training", self.manager,
            )

    def test_employee_and_payroll_workbooks_are_role_restricted_and_masked(self):
        path = self.workbook("employees_training")
        with self.assertRaises(PermissionError):
            self.agent.prepare(
                path, "employees.xlsx", "Add employees after review", "",
                "employees_training", {"id": "planner", "role": "planning_engineer"},
            )
        operation = self.agent.prepare(
            path, "employees.xlsx", "Add employees after review", "",
            "employees_training", self.manager,
        )
        after = operation["preview"]["rows"][0]["after"]
        self.assertEqual("Restricted", after["basic_aed"])
        self.assertEqual("Restricted", after["gross_monthly_aed"])

    def test_checked_in_schema_documents_are_valid_json(self):
        schema_root = Path(__file__).parents[1] / "src" / "prosight" / "schemas"
        for filename in (
            "employees-training.schema.json", "attendance-payroll.schema.json",
            "manpower-deployment.schema.json",
        ):
            document = json.loads((schema_root / filename).read_text(encoding="utf-8"))
            self.assertEqual("object", document["type"])
            self.assertIn("x-prosight-import-kind", document)

    def test_single_renamed_employee_sheet_with_alias_and_optional_columns(self):
        from openpyxl import load_workbook
        path = self.workbook('employees_training')
        book = load_workbook(path)
        del book['Training']
        sheet = book['Employees']; sheet.title = 'Staff register'
        sheet['B5'] = 'Employee name'
        sheet.delete_cols(12)
        book.save(path); book.close()
        parsed = parse_workforce_workbook(path, lambda value: self.project, lambda value: None, 'employees_training')
        self.assertEqual(1, parsed['counts']['employees'])
        self.assertEqual([], parsed['training'])
        self.assertEqual('Staff register', parsed['semantic_mappings']['Employees']['source_sheet'])

    def test_json_profile_maps_headerless_columns_and_rejects_invalid_configuration(self):
        from openpyxl import load_workbook
        from prosight.ingestion.workforce import SPECS
        path = self.workbook('employees_training')
        book = load_workbook(path)
        del book['Training']
        sheet = book['Employees']; sheet.title = 'People'
        sheet.delete_rows(1, 5)
        book.save(path); book.close()
        profile = json.loads((Path(__file__).parents[1] / 'src/prosight/schemas/workforce-mapping.json').read_text())
        rule = profile['sheets']['Employees']
        rule.update(aliases=['People'], data_start_row=1)
        for column, field in enumerate(SPECS['Employees'].fields.values(), 1):
            rule['fields'][field]['column'] = column
        config = self.root / 'mapping.json'
        config.write_text(json.dumps(profile))
        with patch.dict('os.environ', {'PROSIGHT_WORKFORCE_MAPPING_PATH': str(config)}):
            parsed = parse_workforce_workbook(path, lambda value: self.project, lambda value: None, 'employees_training')
            self.assertEqual('EMP-T01', parsed['employees'][0]['employee_id'])
            self.assertEqual(1, parsed['employees'][0]['source_row'])
            rule['fields']['synthetic_name']['column'] = 1
            config.write_text(json.dumps(profile))
            with self.assertRaisesRegex(ValueError, 'same column'):
                parse_workforce_workbook(path, lambda value: self.project, lambda value: None, 'employees_training')

    def test_reject_and_repeated_approval_do_not_duplicate_records(self):
        operation = self.agent.prepare(self.workbook('employees_training'), 'employees.xlsx',
                                       'Add employees after review', '', 'employees_training', self.admin)
        self.agent.decide(operation['id'], self.admin, operation['preview_token'], 'reject')
        with closing(self.repository.connect()) as db:
            self.assertEqual(0, db.execute('SELECT COUNT(*) FROM employees').fetchone()[0])
        completed = self.approve('employees_training')
        self.agent.decide(completed['id'], self.admin, completed['preview_token'], 'approve')
        with closing(self.repository.connect()) as db:
            self.assertEqual(1, db.execute('SELECT COUNT(*) FROM employees').fetchone()[0])


if __name__ == "__main__":
    unittest.main()
