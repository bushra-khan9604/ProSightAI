import copy
import json
import tempfile
import unittest
import asyncio
from types import SimpleNamespace
from unittest.mock import patch
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from prosight.controls.calculations import Calendar, schedule, evm, resource_conflicts, cash_flow, risk_exposure
from prosight.controls.workbooks import parse, merge_files, validate
from prosight.controls.specialists import PlannerAgent, PlanningProposal, PlannedWork, analyze, recalculate_draft


class CalculationTests(unittest.TestCase):
    def setUp(self):
        self.calendar = [{"calendar_id": "C", "working_days": "Mon,Tue,Wed,Thu,Fri", "exceptions": "2026-09-15"}]
        self.activities = [{"activity_id": k, "duration_wd": d, "calendar_id": "C"} for k,d in [("A",2),("B",3),("C",1),("D",2)]]
        self.links = [{"predecessor_id": p,"successor_id": s,"type":"FS","baseline_lag_wd":0} for p,s in [("A","B"),("A","C"),("B","D"),("C","D")]]

    def test_calendar_and_float(self):
        result = schedule(self.activities,self.links,self.calendar,"2026-09-14")
        rows = {r["activity_id"]:r for r in result["activities"]}
        self.assertEqual(result["finish"],"2026-09-23")
        self.assertEqual(rows["C"]["float_wd"],2)
        self.assertEqual(result["critical_path"],["A","B","D"])

    def test_noncritical_recovery_does_not_shorten_project(self):
        result=schedule(self.activities,self.links,self.calendar,"2026-09-14",reductions={"C":1})
        self.assertEqual(result["finish"],"2026-09-23")

    def test_cycles_and_unsupported_relationships(self):
        for link in ({"predecessor_id":"D","successor_id":"A","type":"FS"},{"predecessor_id":"A","successor_id":"D","type":"SS"}):
            with self.assertRaises(ValueError): schedule(self.activities,self.links+[link],self.calendar,"2026-09-14")

    def test_evm_independent_quantity_arithmetic(self):
        budgets=[{"activity_id":"A","unit":"m","quantity":100,"budget_usd":1000}]
        periods=[{"period_end":"2026-09-12","planned_value_usd":500}]
        measurements=[{"activity_id":"A","unit":"m","period_end":"2026-09-12","quantity_this_period":40,"earned_value_usd":99999}]
        costs=[{"date":"2026-09-12","cost_usd":800}]
        result=evm(budgets,periods,measurements,costs,"2026-09-12")
        self.assertEqual(result["ev_usd"],400)
        self.assertEqual(result["cpi"],.5)
        self.assertEqual(result["spi"],.8)
        self.assertEqual(result["eac_usd"],2000)
        self.assertIsNone(evm(budgets,periods,[],[],"2026-09-12",future=True)["cpi"])

    def test_resources_missing_is_not_zero(self):
        demand=[{"resource_id":"R","start_date":"2026-09-14","finish_date":"2026-09-18","requested_teams":2}]
        cap=[{"resource_id":"R","start_date":"2026-09-14","finish_date":"2026-09-18","available_teams":1}]
        self.assertEqual(resource_conflicts(demand,cap)[0]["shortage"],1)
        self.assertIsNone(resource_conflicts(demand,[])[0]["capacity"])

    def test_cash_and_risk(self):
        result=cash_flow([{"due_date":"2026-10-12","net_due_usd":95}], [{"date":"2026-09-01","cost_usd":70}],30)
        self.assertEqual(result["funding_requirement_usd"],70)
        self.assertEqual(result["periods"][-1]["closing_cash_usd"],25)
        self.assertEqual(risk_exposure([{"probability":.25,"impact_cost_usd":100}])["open_expected_cost_usd"],25)

    def test_conflicting_files_and_inheritance(self):
        with self.assertRaises(ValueError): merge_files([{"tables":{"A":[1]}},{"tables":{"A":[2]}}])
        self.assertEqual(merge_files([{"tables":{"A":[2]}}],{"A":[1],"B":[3]}),{"A":[2],"B":[3]})

    def test_uncached_formula_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'bad.xlsx'
            wb=Workbook();wb.active.title='ReadMe';wb.active.append(['value','=1+1']);wb.save(path)
            with self.assertRaisesRegex(ValueError,"saved calculated value"): parse(path,'P',lambda x:x)


class GeneratedFixtureTests(unittest.TestCase):
    def test_generated_workbooks(self):
        root=Path(__file__).resolve().parents[1]/'outputs/project-dataset-20260912/projects'
        if not root.exists(): self.skipTest('Generated demonstration pack is not present')
        from prosight.ingestion.excel import preview_workbook
        projects={folder.name:preview_workbook(folder/'upload_now/Project Master.xlsx',folder.name)['projects'][0] for folder in root.iterdir()}
        def resolve(term):
            return next((code for code,p in projects.items() if term in (code,p['name'])),None)
        for folder in root.iterdir():
            with self.subTest(project=folder.name):
                parsed=[parse(path,folder.name,resolve) for path in folder.rglob('*.xlsx')]
                data=merge_files(parsed)
                project=projects[folder.name]
                cutoff=(data.get('Overview')or[{}])[0].get('data_date',project['reporting_date'])
                result=validate(data,project,cutoff)
                self.assertTrue(result['readiness']['planning'])
                version={'id':'11111111-1111-1111-1111-111111111111','project_code':folder.name,'status':'active','reporting_date':cutoff,'content':data,'synthetic':True}
                if project['status']!='future':
                    self.assertTrue(result['readiness']['schedule'])
                    analysis=analyze(version,project)
                    self.assertEqual(analysis['schedule']['forecast_finish'],max(a['forecast_finish'] for a in data['Schedule']))
                    expected=next(p for p in data['Performance'] if p['period_end']==cutoff)
                    # Quantities in the demonstration workbooks are rounded; source EV
                    # uses the generator's unrounded fractions. Keep the calculation
                    # quantity-based and allow at most one USD of aggregate rounding.
                    self.assertAlmostEqual(analysis['evm']['ev_usd'],expected['ev_usd'],delta=1)
                    self.assertAlmostEqual(analysis['evm']['ac_usd'],expected['ac_usd'],places=1)
                else:
                    resources=data['Resources']
                    previous={};work=[]
                    for scope in data['Scope Quantities']:
                        resource=next(r for r in resources if r['trade'].casefold() in scope['work_item'].casefold())
                        predecessors=[previous[scope['wbs_id']]] if scope['wbs_id'] in previous else []
                        work.append(PlannedWork(scope_id=scope['scope_id'],resource_id=resource['resource_id'],predecessor_scope_ids=predecessors,rationale='Sequential phase within supplied work front',procurement_lead_days=60,procurement_basis='Proposed 60-day lead allowance; supplier confirmation required'))
                        previous[scope['wbs_id']]=scope['scope_id']
                    proposal=PlanningProposal(requirements=[],work=work,assumptions=['Sequential phases within each work front'],unresolved_conflicts=[],kickoff_agenda=['Confirm scope, owners and milestones'])
                    class Client:
                        async def __aenter__(self):return self
                        async def __aexit__(self,*args):pass
                        @property
                        def responses(self):return self
                        async def parse(self,**kwargs):return SimpleNamespace(output_parsed=proposal)
                    with patch('openai.AsyncOpenAI',return_value=Client()),patch('prosight.controls.specialists.get_settings',return_value=SimpleNamespace(ai_provider='openai',openai_api_key='test',openai_model='test')):
                        draft,notice=asyncio.run(PlannerAgent().draft(version,project,'Create a project plan',[]))
                    self.assertIsNotNone(draft,notice)
                    validate(draft,project,cutoff)
                    self.assertEqual(len(draft['Schedule']),120)
                    self.assertEqual(draft['Actual Costs'],[])
                    self.assertAlmostEqual(sum(b['budget_usd'] for b in draft['Cost Baseline']),project['contract_value_usd']*.82,places=2)
                    self.assertAlmostEqual(sum(r['budget_weight'] for r in draft['Measurement Plan']),1,places=5)
                    before=sum(b['budget_usd'] for b in draft['Cost Baseline'])
                    draft['Cost Baseline'][0]['material_subcontract_usd']+=100
                    recalculate_draft(draft,project)
                    self.assertAlmostEqual(sum(b['budget_usd'] for b in draft['Cost Baseline']),before+100,places=2)
                    from prosight.controls.exports import render
                    from pypdf import PdfReader
                    pdf,mime=render({**version,'status':'draft','content':draft},'pdf')
                    self.assertIn('application/pdf',mime)
                    reader=PdfReader(BytesIO(pdf))
                    self.assertTrue(all('Synthetic demonstration data' in page.extract_text() for page in reader.pages))
                    xlsx,_=render({**version,'status':'draft','content':draft},'xlsx')
                    from openpyxl import load_workbook
                    book=load_workbook(BytesIO(xlsx),data_only=True)
                    self.assertEqual(book['Schedule'].max_row,121)
                    book.close()


class SpecialistBoundaryTests(unittest.TestCase):
    def test_active_roster_replaces_legacy_and_partial_sections_fall_back(self):
        from prosight.supabase_repository import SupabaseProjectRepository
        repository=SupabaseProjectRepository.__new__(SupabaseProjectRepository)
        repository.user=SimpleNamespace(select_all=lambda *args,**kwargs:[{'project_code':'P','content':{'Current Manpower':[{'current_project_code':'P','emp_code':'NEW'}]}},{'project_code':'Q','content':{}}])
        legacy=[{'current_project_code':'P','emp_code':'OLD'},{'current_project_code':'Q','emp_code':'KEEP'}]
        rows=repository._controls_rows('Current Manpower',legacy,project_field='current_project_code')
        self.assertEqual({r['emp_code'] for r in rows},{'NEW','KEEP'})

    def test_basic_schedule_update_preserves_original_baseline(self):
        parent={'Schedule':[{'activity_id':'A','baseline_finish':'2026-09-01','forecast_start':'2026-09-01','forecast_finish':'2026-09-02'}]}
        merged=merge_files([{'tables':{'Basic Schedule':[{'activity_id':'A','start':'2026-09-01','finish':'2026-09-03'}]}}],parent)
        self.assertEqual(merged['Schedule'][0]['baseline_finish'],'2026-09-01')
        self.assertEqual(merged['Schedule'][0]['forecast_finish'],'2026-09-03')
        self.assertEqual(parent['Schedule'][0]['forecast_finish'],'2026-09-02')

    def test_evaluation_metadata_blocks_otherwise_valid_workbooks(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'renamed.xlsx'
            book=Workbook();book.active.title='ReadMe';book.active.append(['Classification','evaluation']);book.save(path);book.close()
            with self.assertRaisesRegex(ValueError,'Evaluation reference answers'):parse(path,'P',lambda x:x)

    def test_routing_and_verified_citations(self):
        from prosight.agents.orchestrator import MultiAgentOrchestrator
        from prosight.contracts import WriterInput,DatabaseEvidence,EvidenceItem
        orchestrator=MultiAgentOrchestrator(SimpleNamespace())
        self.assertEqual(orchestrator.plan('Analyze EVM').agents,['analyst','writer'])
        self.assertEqual(orchestrator.plan('Create a schedule').agents,['planner','writer'])
        self.assertEqual(orchestrator.plan('Analyze delays and prepare a recovery plan').agents,['analyst','planner','writer'])
        self.assertNotIn('planner',orchestrator.plan('Summarize the scope document','P').agents)
        packet=WriterInput(query='Q',database=DatabaseEvidence(evidence=[EvidenceItem(text='Fact',citation='Report.pdf, page 2')]))
        answer=orchestrator._answer('Report.pdf, page 2; Invented.pdf, page 99',[],packet)
        self.assertEqual(answer.citations,['Report.pdf, page 2'])

    def test_worker_rechecks_membership(self):
        from prosight.controls.jobs import WorkerGateway
        class Service:
            def select(self,table,**kw):return [{'id':'U','role':'planning_engineer'}] if table=='profiles' else []
            def select_all(self,table,**kw):return []
        gateway=WorkerGateway(Service(),{'requested_by':'U','project_code':'P','id':'J'})
        with self.assertRaises(PermissionError):gateway.select('controls_versions',project_code='eq.P')
        with self.assertRaises(PermissionError):gateway.select('projects')

    def test_model_unavailable_does_not_create_a_fake_plan(self):
        with patch('prosight.controls.specialists.get_settings',return_value=SimpleNamespace(ai_provider='local',openai_api_key='')):
            draft,notice=asyncio.run(PlannerAgent().draft({}, {}, 'plan', []))
            self.assertIsNone(draft)
            self.assertIn('could not complete',notice)


if __name__=='__main__': unittest.main()
